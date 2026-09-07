#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Validate the generated UAT pack (workbook + Thai companion doc).

    python -X utf8 requirement_spec/5_uat/_build/validate_uat_pack.py

Writes the full per-check report to ``_build/validate_uat_pack.out.txt`` and
prints ONLY a compact one-line-per-check PASS/FAIL summary to stdout.
Exit code 0 when every check passes, 1 otherwise.

The ten checks (1-9 verbatim from the T5 brief; 10 added 2026-09-07):

  1. 4 sheets present with the expected names
  2. header row is row 2 on '2. Test Cases'; data starts row 3 (no 'TC-EX' row)
  3. exactly 54 data rows, ids UAT-01..UAT-54 contiguous, no duplicates
  4. every Summary range ends at the real last data row (parse the formula
     strings; fail on any literal '$64')
  5. the 8 module names in the Summary breakdown exactly equal the distinct
     modules used in '2. Test Cases', and the per-module totals sum to 54
  6. every DV list present on the columns that need one, covering the real
     used range
  7. the production URL appears in sheet 1 and the staging hostname
     'cman-budget-web-stg' appears NOWHERE in the workbook
  8. UAT_Test_Script.md exists and its 'UAT-xx -> SIT' table has 54 rows
  9. coverage check — every required SIT id and all of U-01..U-25 still appear
     in the traceability data.  This is the proof that the de-duplication pass
     from 68 cases down to 54 did not quietly drop coverage.
 10. the 'Run Order / ลำดับการรัน' column is a permutation of 1..54 AND every
     case's Run Order is strictly greater than the Run Order of every id in
     its 'runs_after' array.  This is the proof that the round is actually
     runnable in the order the workbook prints: Wave order is not execution
     order, and a tester who reaches the irreversible UAT-40 Approve early
     permanently loses UAT-41, UAT-42 step 5 and UAT-44 — on production,
     inside the round.  Checked against the case JSON, never trusted because
     the generator wrote it.

Nothing here is allowed to "pass by accident": every check that walks a
collection also asserts the collection is non-empty, and check 9 reads the
per-case rows of the traceability table only — never the prose summary
sentence underneath it, which enumerates every id and would make the check
vacuous.  Check 10 likewise fails when the graph it walks carries no edges.
All ten checks were fault-injected against throwaway copies on 2026-09-07;
each one caught its own fault.

House rules: run under ``python -X utf8`` (the Windows console is cp1252) and
give every open() an explicit encoding="utf-8".
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# paths
# --------------------------------------------------------------------------
HERE = Path(__file__).resolve().parent            # requirement_spec/5_uat/_build
UAT_DIR = HERE.parent                             # requirement_spec/5_uat
REPO_ROOT = UAT_DIR.parent.parent                 # repo root

XLSX = UAT_DIR / "Budget_UAT_Test_Cases_07.09.2026.xlsx"
MD = UAT_DIR / "UAT_Test_Script.md"
CASES_DIR = REPO_ROOT / ".scratch" / "uat-prep" / "cases"
REPORT = HERE / "validate_uat_pack.out.txt"

# --------------------------------------------------------------------------
# expected facts
# --------------------------------------------------------------------------
SHEET_INFO = "1. Info & Instructions"
SHEET_CASES = "2. Test Cases"
SHEET_SUMMARY = "3. Summary"
SHEET_DEFECTS = "4. Defects"
EXPECTED_SHEETS = [SHEET_INFO, SHEET_CASES, SHEET_SUMMARY, SHEET_DEFECTS]

EXPECTED_CASE_COUNT = 54
HEADER_ROW = 2
FIRST_DATA_ROW = 3

# column order on '2. Test Cases'.  A bilingual header such as
# "Preconditions / เงื่อนไขก่อนทดสอบ" is matched on the part before " / ".
EXPECTED_CASE_HEADERS = [
    "Wave", "Run Order", "No.", "Test Case ID", "Module", "Department", "PIC",
    "Test Scenario", "Priority", "Preconditions", "Test Steps", "Test Data",
    "Expected Result", "Actual Result", "Status", "Severity", "Defect ID",
    "Tester", "Test Date", "Remarks",
]

# 1-based column indexes derived from the expected layout above, so a future
# column insertion is a one-line change here instead of a hunt for literals.
# Check 2 independently asserts the sheet's real header row equals that layout.
CASE_COL_IDX = {h: i + 1 for i, h in enumerate(EXPECTED_CASE_HEADERS)}
IDX_CASE_ID = CASE_COL_IDX["Test Case ID"]
IDX_NO = CASE_COL_IDX["No."]
IDX_RUN_ORDER = CASE_COL_IDX["Run Order"]

EXPECTED_DEFECT_HEADERS = [
    "Defect ID", "UAT Case ID", "Department", "Title", "Severity",
    "Reported by", "Reported date", "Owner", "Status", "Fixed in",
    "Retest date", "Retest result", "Notes",
]

# DV lists keyed by the header label of the column that must carry them
CASE_DV_BY_HEADER = {
    "Department": "Data & Analytic,Solution Delivery,ทั้งสองฝ่าย,-",
    "Priority": "High,Medium,Low",
    "Status": "Not Run,Pass,Fail,Blocked,N/A",
    "Severity": "Critical,High,Medium,Low,-",
}
DEFECT_DV_BY_HEADER = {
    # ONE severity vocabulary pack-wide since 2026-09-07: the defect log used
    # to read Blocker,Major,Minor while the case sheet read Critical..Low, and
    # companion 5.4 told the tester to carry one onto the other.  Still an
    # exact-match assertion, just against the unified list.
    "Severity": "Critical,High,Medium,Low",
    "Status": "Open,In Progress,Fixed,Retested,Closed,Wont Fix",
    "Retest result": "Pass,Fail",
}

EXPECTED_MODULES = [
    "Login / Authentication",
    "Budget Entry (Filler)",
    "Special GL Subform",
    "Attachments",
    "Submission / Workflow",
    "Approval",
    "Email Alert / Notification",
    "Business Rules & UI",
]

PROD_URL = (
    "https://cman-budget-web-prd.kindstone-f34836dd.southeastasia."
    "azurecontainerapps.io/"
)
STAGING_HOSTNAME = "cman-budget-web-stg"

REQUIRED_SIT_IDS = [
    "TC-001", "TC-002", "TC-004", "TC-005", "TC-006", "TC-007", "TC-008",
    "TC-010", "TC-011", "TC-012", "TC-013", "TC-015", "TC-016", "TC-019",
    "TC-020", "TC-021", "TC-022", "TC-023", "TC-024", "TC-025", "TC-026",
    "TC-028", "TC-029", "TC-035", "TC-036", "TC-037", "TC-038", "TC-039",
    "TC-043", "TC-044", "TC-045", "TC-046", "TC-047", "TC-048", "TC-049",
    "TC-050", "TC-051", "TC-052", "TC-053", "TC-054", "TC-055", "TC-056",
    "TC-058", "TC-059", "TC-060", "TC-061",
]
REQUIRED_NEW_IDS = ["U-%02d" % n for n in range(1, 26)]

BEGIN_MARK = "<!-- BEGIN GENERATED TRACEABILITY (build_uat_test_cases.py) -->"
END_MARK = "<!-- END GENERATED TRACEABILITY -->"

WAVE_FILES = ["W1", "W2", "W3", "W4", "W5", "W6", "W7", "W8"]

# range like  '2. Test Cases'!$C$3:$C$56
RANGE_RE = re.compile(
    r"'(?P<sheet>[^']+)'!\$(?P<c1>[A-Z]{1,3})\$(?P<r1>\d+)"
    r"(?::\$(?P<c2>[A-Z]{1,3})\$(?P<r2>\d+))?"
)


# --------------------------------------------------------------------------
# report plumbing
# --------------------------------------------------------------------------
class Report:
    """Collects the long-form report; stdout gets one line per check."""

    def __init__(self):
        self.lines = []
        self.results = []          # (n, title, ok, reason)

    def detail(self, text=""):
        self.lines.append(text)

    def check(self, n, title, ok, reason=""):
        self.results.append((n, title, bool(ok), reason))
        self.detail("")
        self.detail("  RESULT: %s%s" % ("PASS" if ok else "FAIL",
                                        ("  — " + reason) if reason else ""))
        self.detail("-" * 78)

    def head(self, n, title):
        self.detail("=" * 78)
        self.detail("CHECK %d — %s" % (n, title))
        self.detail("=" * 78)


def norm_header(value):
    """'Preconditions / เงื่อนไขก่อนทดสอบ' -> 'Preconditions'."""
    if value is None:
        return ""
    return str(value).split(" / ")[0].strip() if " / " in str(value) else str(value).strip()


def col_letter_by_header(ws, header_row, wanted):
    """Column letter whose normalised header equals `wanted`, else None."""
    for c in range(1, ws.max_column + 1):
        if norm_header(ws.cell(header_row, c).value) == wanted:
            return ws.cell(header_row, c).column_letter
    return None


# --------------------------------------------------------------------------
# checks
# --------------------------------------------------------------------------
def check_1_sheets(rep, wb):
    rep.head(1, "4 sheets present with the expected names")
    actual = list(wb.sheetnames)
    rep.detail("  expected : %r" % (EXPECTED_SHEETS,))
    rep.detail("  actual   : %r" % (actual,))
    problems = []
    for name in EXPECTED_SHEETS:
        if name not in actual:
            problems.append("missing sheet %r" % name)
    extra = [n for n in actual if n not in EXPECTED_SHEETS]
    if extra:
        problems.append("unexpected sheet(s) %r" % extra)
    if not problems and actual != EXPECTED_SHEETS:
        problems.append("sheet order is %r, expected %r" % (actual, EXPECTED_SHEETS))
    rep.check(1, "4 sheets present with the expected names",
              not problems, "; ".join(problems))
    return not problems


def check_2_header(rep, wb):
    rep.head(2, "header row is row 2 on '2. Test Cases'; data starts row 3 (no TC-EX)")
    ws = wb[SHEET_CASES]
    problems = []

    headers = [norm_header(ws.cell(HEADER_ROW, c).value)
               for c in range(1, ws.max_column + 1)]
    rep.detail("  row %d headers (normalised): %r" % (HEADER_ROW, headers))
    if headers != EXPECTED_CASE_HEADERS:
        problems.append("row 2 headers != expected; got %r" % (headers,))

    title = ws.cell(1, 1).value
    rep.detail("  row 1 A1 (title row): %r" % (title,))
    if not isinstance(title, str) or not title.strip():
        problems.append("row 1 is not a non-empty title row")
    row1 = [norm_header(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    if row1 == EXPECTED_CASE_HEADERS:
        problems.append("row 1 also looks like a header row — header must be row 2 only")

    id_letter = get_column_letter(IDX_CASE_ID)
    first_id = ws.cell(FIRST_DATA_ROW, IDX_CASE_ID).value
    rep.detail("  first data cell %s%d: %r"
               % (id_letter, FIRST_DATA_ROW, first_id))
    if first_id != "UAT-01":
        problems.append("%s%d is %r, expected 'UAT-01'"
                        % (id_letter, FIRST_DATA_ROW, first_id))

    # The SIT example row is a DATA row, so it can only exist on '2. Test Cases'.
    # Scope the assertion to that sheet: a mention on another sheet is prose.
    hits = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "TC-EX" in cell.value:
                hits.append("%s!%s" % (ws.title, cell.coordinate))
    rep.detail("  'TC-EX' occurrences on '%s': %d %s"
               % (SHEET_CASES, len(hits), hits if hits else ""))
    if hits:
        problems.append("SIT example row 'TC-EX' still present at %s" % ", ".join(hits))

    # Elsewhere in the workbook 'TC-EX' can only be prose (the Summary note says
    # the file has no example row).  Reported for visibility, never a failure.
    elsewhere = []
    for sheet in wb.worksheets:
        if sheet.title == SHEET_CASES:
            continue
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "TC-EX" in cell.value:
                    elsewhere.append("%s!%s" % (sheet.title, cell.coordinate))
    rep.detail("  INFO (not a failure) — 'TC-EX' mentioned as prose on other "
               "sheets: %r" % elsewhere)

    rep.check(2, "header row is row 2; data starts row 3; no TC-EX row",
              not problems, "; ".join(problems))
    return not problems


def check_3_rows(rep, wb):
    rep.head(3, "exactly 54 data rows, ids UAT-01..UAT-54 contiguous, no "
                "duplicates")
    ws = wb[SHEET_CASES]
    problems = []

    ids = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        v = ws.cell(r, IDX_CASE_ID).value
        if v is None or str(v).strip() == "":
            continue
        ids.append((r, str(v).strip()))
    rep.detail("  data rows found : %d (rows %s..%s)"
               % (len(ids), ids[0][0] if ids else "-", ids[-1][0] if ids else "-"))
    rep.detail("  first / last id : %s / %s"
               % (ids[0][1] if ids else "-", ids[-1][1] if ids else "-"))

    if len(ids) != EXPECTED_CASE_COUNT:
        problems.append("%d data rows, expected %d" % (len(ids), EXPECTED_CASE_COUNT))

    expected_ids = ["UAT-%02d" % n for n in range(1, EXPECTED_CASE_COUNT + 1)]
    actual_ids = [i for _r, i in ids]
    if actual_ids != expected_ids:
        missing = [i for i in expected_ids if i not in actual_ids]
        unexpected = [i for i in actual_ids if i not in expected_ids]
        if missing:
            problems.append("missing ids: %s" % ", ".join(missing))
        if unexpected:
            problems.append("unexpected ids: %s" % ", ".join(unexpected))
        if not missing and not unexpected:
            problems.append("ids present but out of order")

    dupes = sorted({i for i in actual_ids if actual_ids.count(i) > 1})
    if dupes:
        problems.append("duplicate ids: %s" % ", ".join(dupes))
    rep.detail("  duplicates      : %s" % (dupes if dupes else "none"))

    # rows must be contiguous — no blank row inside the block
    rows = [r for r, _i in ids]
    if rows and rows != list(range(rows[0], rows[0] + len(rows))):
        problems.append("data rows are not contiguous: %r" % rows)

    # the No. column must run 1..54 alongside the ids
    nos = [ws.cell(r, IDX_NO).value for r, _i in ids]
    if nos != list(range(1, len(ids) + 1)):
        problems.append("'No.' column is not 1..%d" % len(ids))

    # NOTE: the 'Run Order' column is NOT checked here — it is CHECK 10, which
    # owns the permutation and the runs_after ordering proof on its own.

    rep.check(3, "54 rows UAT-01..UAT-54 contiguous, no duplicates",
              not problems, "; ".join(problems))
    return (not problems), (rows[-1] if rows else None)


def check_4_summary_ranges(rep, wb, last_data_row):
    rep.head(4, "every Summary range ends at the real last data row (no literal $64)")
    ws = wb[SHEET_SUMMARY]
    problems = []
    seen = 0
    literal_64 = []

    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or not v.startswith("="):
                continue
            if "$64" in v:
                literal_64.append("%s: %s" % (cell.coordinate, v))
            for m in RANGE_RE.finditer(v):
                if m.group("sheet") != SHEET_CASES:
                    continue
                seen += 1
                r1 = int(m.group("r1"))
                r2 = int(m.group("r2")) if m.group("r2") else r1
                if r1 != FIRST_DATA_ROW or r2 != last_data_row:
                    problems.append(
                        "%s references $%s$%d:$%s$%d, expected rows %d..%d"
                        % (cell.coordinate, m.group("c1"), r1,
                           m.group("c2") or m.group("c1"), r2,
                           FIRST_DATA_ROW, last_data_row))

    rep.detail("  real last data row on '%s' : %s" % (SHEET_CASES, last_data_row))
    rep.detail("  '%s' ranges found in Summary formulas : %d" % (SHEET_CASES, seen))
    rep.detail("  formulas containing the literal '$64' : %d" % len(literal_64))
    for line in literal_64:
        rep.detail("      %s" % line)
    if literal_64:
        problems.append("%d formula(s) still carry the hard-coded SIT row '$64'"
                        % len(literal_64))
    if seen == 0:
        problems.append("no '%s' range found in any Summary formula — "
                        "the check would pass vacuously" % SHEET_CASES)

    # ranges that stay inside the Summary sheet (the SUM total rows)
    intra = 0
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("=SUM("):
                intra += 1
    rep.detail("  intra-sheet SUM total formulas        : %d" % intra)
    if intra == 0:
        problems.append("no SUM total row found on the Summary sheet")

    uniq = sorted(set(problems))
    rep.check(4, "every Summary range ends at the real last data row",
              not uniq, "; ".join(uniq[:6]))
    return not uniq


def check_5_modules(rep, wb, last_data_row):
    rep.head(5, "Summary module breakdown == distinct modules used, totals sum to 54")
    cases = wb[SHEET_CASES]
    summary = wb[SHEET_SUMMARY]
    problems = []

    mod_col = col_letter_by_header(cases, HEADER_ROW, "Module")
    wave_col = col_letter_by_header(cases, HEADER_ROW, "Wave")
    dept_col = col_letter_by_header(cases, HEADER_ROW, "Department")
    if not mod_col:
        rep.check(5, "Summary module breakdown", False, "no 'Module' column found")
        return False

    used_modules = []
    for r in range(FIRST_DATA_ROW, last_data_row + 1):
        v = cases["%s%d" % (mod_col, r)].value
        if v is not None and str(v).strip():
            used_modules.append(str(v).strip())
    distinct_used = sorted(set(used_modules))
    rep.detail("  distinct modules used in '%s' (%d): %r"
               % (SHEET_CASES, len(distinct_used), distinct_used))

    # module labels listed in the Summary breakdown: the column that holds the
    # 'Module' header, the rows under it up to (not including) the total row.
    labels = []
    hdr_cell = None
    for row in summary.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == "Module":
                hdr_cell = cell
                break
        if hdr_cell:
            break
    if hdr_cell is None:
        rep.check(5, "Summary module breakdown", False,
                  "no 'Module' header found on the Summary sheet")
        return False
    r = hdr_cell.row + 1
    while r <= summary.max_row:
        v = summary.cell(r, hdr_cell.column).value
        if v is None or not str(v).strip():
            break
        text = str(v).strip()
        if text.startswith("รวม") or text.lower().startswith("total"):
            break
        labels.append(text)
        r += 1
    rep.detail("  module labels in Summary breakdown (%d): %r" % (len(labels), labels))

    if sorted(set(labels)) != distinct_used:
        problems.append("Summary module list %r != distinct modules used %r"
                        % (sorted(set(labels)), distinct_used))
    if len(labels) != len(set(labels)):
        problems.append("duplicate module labels in the Summary breakdown")
    if sorted(set(labels)) != sorted(EXPECTED_MODULES):
        problems.append("Summary module list is not the 8 modules the brief fixes")

    totals = {m: used_modules.count(m) for m in labels}
    rep.detail("  per-module totals: %r" % (totals,))
    total = sum(totals.values())
    rep.detail("  sum of per-module totals: %d (expected %d)"
               % (total, EXPECTED_CASE_COUNT))
    if total != EXPECTED_CASE_COUNT:
        problems.append("per-module totals sum to %d, expected %d"
                        % (total, EXPECTED_CASE_COUNT))
    empty = [m for m, n in totals.items() if n == 0]
    if empty:
        problems.append("module(s) listed with zero cases: %s" % ", ".join(empty))

    # the Wave and Department breakdowns must add up to the same 54
    for label, col in (("Wave", wave_col), ("Department", dept_col)):
        if not col:
            problems.append("no '%s' column on '%s'" % (label, SHEET_CASES))
            continue
        vals = [str(cases["%s%d" % (col, r2)].value).strip()
                for r2 in range(FIRST_DATA_ROW, last_data_row + 1)
                if cases["%s%d" % (col, r2)].value is not None]
        rep.detail("  %s breakdown source count: %d rows, %d distinct"
                   % (label, len(vals), len(set(vals))))
        if len(vals) != EXPECTED_CASE_COUNT:
            problems.append("%s column has %d filled rows, expected %d"
                            % (label, len(vals), EXPECTED_CASE_COUNT))

    rep.check(5, "8 modules match the sheet and totals sum to 54",
              not problems, "; ".join(problems))
    return not problems


def check_6_dv(rep, wb, last_data_row):
    rep.head(6, "every DV list present on the columns that need one, over the real range")
    problems = []

    def dv_map(ws):
        out = {}
        for dv in ws.data_validations.dataValidation:
            out.setdefault(str(dv.formula1 or "").strip('"'), []).append(dv)
        return out

    # --- '2. Test Cases'
    cases = wb[SHEET_CASES]
    rep.detail("  sheet '%s' — data rows %d..%d" % (SHEET_CASES, FIRST_DATA_ROW,
                                                    last_data_row))
    for header, expected_list in CASE_DV_BY_HEADER.items():
        col = col_letter_by_header(cases, HEADER_ROW, header)
        if not col:
            problems.append("'%s': no column with header %r" % (SHEET_CASES, header))
            continue
        found = None
        for dv in cases.data_validations.dataValidation:
            f1 = str(dv.formula1 or "").strip('"')
            if f1 == expected_list:
                covered = all("%s%d" % (col, r) in dv.sqref
                              for r in range(FIRST_DATA_ROW, last_data_row + 1))
                if covered:
                    found = dv
                    break
        rep.detail("    %-12s col %s  list=%r  -> %s"
                   % (header, col, expected_list, "OK" if found else "MISSING"))
        if found is None:
            actual = [(str(dv.formula1 or "").strip('"'), str(dv.sqref))
                      for dv in cases.data_validations.dataValidation]
            problems.append("'%s'.%s (%s): no DV list %r covering %s%d:%s%d "
                            "(present: %r)"
                            % (SHEET_CASES, col, header, expected_list, col,
                               FIRST_DATA_ROW, col, last_data_row, actual))

    # --- '4. Defects'
    defects = wb[SHEET_DEFECTS]
    dh = [norm_header(defects.cell(HEADER_ROW, c).value)
          for c in range(1, defects.max_column + 1)]
    rep.detail("  sheet '%s' — headers %r" % (SHEET_DEFECTS, dh))
    if dh != EXPECTED_DEFECT_HEADERS:
        problems.append("'%s' headers %r != expected %r"
                        % (SHEET_DEFECTS, dh, EXPECTED_DEFECT_HEADERS))

    id_col = col_letter_by_header(defects, HEADER_ROW, "Defect ID")
    last_defect_row = HEADER_ROW
    if id_col:
        for r in range(FIRST_DATA_ROW, defects.max_row + 1):
            v = defects["%s%d" % (id_col, r)].value
            if v is None or not str(v).strip().startswith("D-"):
                break
            last_defect_row = r
    n_defect_rows = last_defect_row - HEADER_ROW
    rep.detail("  '%s' pre-filled numbered rows: %d (rows %d..%d)"
               % (SHEET_DEFECTS, n_defect_rows, FIRST_DATA_ROW, last_defect_row))
    if n_defect_rows != 40:
        problems.append("'%s' has %d pre-filled numbered rows, expected 40"
                        % (SHEET_DEFECTS, n_defect_rows))

    for header, expected_list in DEFECT_DV_BY_HEADER.items():
        col = col_letter_by_header(defects, HEADER_ROW, header)
        if not col:
            problems.append("'%s': no column with header %r" % (SHEET_DEFECTS, header))
            continue
        found = None
        for dv in defects.data_validations.dataValidation:
            f1 = str(dv.formula1 or "").strip('"')
            if f1 == expected_list:
                covered = all("%s%d" % (col, r) in dv.sqref
                              for r in range(FIRST_DATA_ROW, last_defect_row + 1))
                if covered:
                    found = dv
                    break
        rep.detail("    %-14s col %s  list=%r  -> %s"
                   % (header, col, expected_list, "OK" if found else "MISSING"))
        if found is None:
            actual = [(str(dv.formula1 or "").strip('"'), str(dv.sqref))
                      for dv in defects.data_validations.dataValidation]
            problems.append("'%s'.%s (%s): no DV list %r covering %s%d:%s%d "
                            "(present: %r)"
                            % (SHEET_DEFECTS, col, header, expected_list, col,
                               FIRST_DATA_ROW, col, last_defect_row, actual))

    rep.check(6, "every required DV list present and covering the real range",
              not problems, "; ".join(problems)[:400])
    return not problems


def check_7_urls(rep, wb):
    rep.head(7, "production URL on sheet 1; staging hostname nowhere in the workbook")
    problems = []

    info = wb[SHEET_INFO]
    prod_cells = []
    for row in info.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and PROD_URL in cell.value:
                prod_cells.append(cell.coordinate)
    rep.detail("  production URL: %s" % PROD_URL)
    rep.detail("  found on '%s' at: %r" % (SHEET_INFO, prod_cells))
    if not prod_cells:
        problems.append("production URL not found on '%s'" % SHEET_INFO)

    stg_cells = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and STAGING_HOSTNAME in cell.value:
                    stg_cells.append("%s!%s" % (sheet.title, cell.coordinate))
    rep.detail("  staging hostname %r found in cells: %r"
               % (STAGING_HOSTNAME, stg_cells))
    if stg_cells:
        problems.append("staging hostname present in cells: %s" % ", ".join(stg_cells))

    # belt and braces: scan every part of the .xlsx zip, not just live cells
    raw_hits = []
    with zipfile.ZipFile(XLSX) as zf:
        for name in zf.namelist():
            data = zf.read(name)
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                continue
            if STAGING_HOSTNAME in text:
                raw_hits.append(name)
    rep.detail("  raw .xlsx zip parts containing %r: %r"
               % (STAGING_HOSTNAME, raw_hits))
    if raw_hits:
        problems.append("staging hostname present in raw workbook parts: %s"
                        % ", ".join(raw_hits))

    rep.check(7, "prod URL on sheet 1, staging hostname absent",
              not problems, "; ".join(problems))
    return not problems


def read_traceability_rows():
    """Rows of the generated 'UAT-xx -> SIT' table in UAT_Test_Script.md.

    Returns (rows, error).  Only the per-case table rows are returned — the
    prose coverage summary underneath the table lists every id and would make
    check 9 vacuous if it were included.
    """
    if not MD.exists():
        return None, "%s does not exist" % MD.name
    with open(MD, encoding="utf-8") as fh:
        text = fh.read()
    if BEGIN_MARK not in text or END_MARK not in text:
        return None, "generated-traceability markers missing in %s" % MD.name
    block = text.split(BEGIN_MARK, 1)[1].split(END_MARK, 1)[0]
    rows = []
    for line in block.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if not cells or not re.fullmatch(r"UAT-\d{2}", cells[0]):
            continue
        rows.append(cells)
    return rows, None


def check_8_markdown(rep):
    rep.head(8, "UAT_Test_Script.md exists and its 'UAT-xx -> SIT' table has 54 rows")
    problems = []
    rep.detail("  markdown file: %s" % MD)
    rows, err = read_traceability_rows()
    if err:
        rep.detail("  ERROR: %s" % err)
        rep.check(8, "UAT_Test_Script.md traceability table has 54 rows", False, err)
        return False, None

    rep.detail("  traceability table rows: %d (expected %d)"
               % (len(rows), EXPECTED_CASE_COUNT))
    if len(rows) != EXPECTED_CASE_COUNT:
        problems.append("traceability table has %d rows, expected %d"
                        % (len(rows), EXPECTED_CASE_COUNT))

    ids = [r[0] for r in rows]
    expected_ids = ["UAT-%02d" % n for n in range(1, EXPECTED_CASE_COUNT + 1)]
    if ids != expected_ids:
        missing = [i for i in expected_ids if i not in ids]
        if missing:
            problems.append("rows missing for: %s" % ", ".join(missing))
        else:
            problems.append("traceability rows are out of order or duplicated")
    rep.detail("  first / last row id: %s / %s"
               % (ids[0] if ids else "-", ids[-1] if ids else "-"))

    # the source column must actually carry something for every row
    blank_src = [r[0] for r in rows if len(r) < 4 or not r[3].strip()]
    rep.detail("  rows with an empty source cell: %r" % blank_src)
    if blank_src:
        problems.append("empty source cell on: %s" % ", ".join(blank_src))

    rep.check(8, "UAT_Test_Script.md traceability table has 54 rows",
              not problems, "; ".join(problems))
    return (not problems), rows


def load_case_json():
    import json
    cases = []
    for name in WAVE_FILES:
        path = CASES_DIR / ("%s.json" % name)
        if not path.exists():
            return None, "missing case file %s" % path
        with open(path, encoding="utf-8") as fh:
            cases += json.load(fh)
    return cases, None


def check_9_coverage(rep, trace_rows):
    rep.head(9, "coverage — every required SIT id and all of U-01..U-25 still present")
    problems = []

    rep.detail("  This is the proof that the de-duplication pass from 68 cases")
    rep.detail("  down to 54 did not quietly drop coverage.")
    rep.detail("  required SIT ids : %d" % len(REQUIRED_SIT_IDS))
    rep.detail("  required U-xx ids: %d (U-01..U-25)" % len(REQUIRED_NEW_IDS))
    rep.detail("")

    # ---- source A: the per-case rows of the generated traceability table
    if trace_rows is None:
        problems.append("traceability table unreadable — cannot prove coverage")
        md_sit = md_new = set()
        md_by_id = {}
    else:
        md_by_id = {}
        md_sit, md_new = set(), set()
        for row in trace_rows:
            src = row[3] if len(row) > 3 else ""
            found = set(re.findall(r"\bTC-\d{3}\b", src)) | set(
                re.findall(r"\bU-\d{2}\b", src))
            md_by_id[row[0]] = found
            md_sit |= {i for i in found if i.startswith("TC-")}
            md_new |= {i for i in found if i.startswith("U-")}
        rep.detail("  source A — per-case rows of the traceability table in %s"
                   % MD.name)
        rep.detail("      distinct SIT ids cited : %d" % len(md_sit))
        rep.detail("      distinct U-xx ids cited: %d" % len(md_new))

    # ---- source B: the case JSON the workbook and the table are built from
    cases, err = load_case_json()
    if err:
        problems.append(err)
        js_sit = js_new = set()
    else:
        js_sit = {r for c in cases for r in (c.get("sit_refs") or [])
                  if str(r).startswith("TC-")}
        js_new = {r for c in cases for r in (c.get("new_refs") or [])
                  if re.fullmatch(r"U-\d{2}", str(r))}
        rep.detail("  source B — sit_refs / new_refs in .scratch/uat-prep/cases/*.json")
        rep.detail("      cases loaded           : %d" % len(cases))
        rep.detail("      distinct SIT ids cited : %d" % len(js_sit))
        rep.detail("      distinct U-xx ids cited: %d" % len(js_new))
    rep.detail("")

    missing_sit_md = [i for i in REQUIRED_SIT_IDS if i not in md_sit]
    missing_sit_js = [i for i in REQUIRED_SIT_IDS if i not in js_sit]
    missing_new_md = [i for i in REQUIRED_NEW_IDS if i not in md_new]
    missing_new_js = [i for i in REQUIRED_NEW_IDS if i not in js_new]

    for label, missing in (
        ("SIT ids missing from the traceability table", missing_sit_md),
        ("SIT ids missing from the case JSON", missing_sit_js),
        ("U-xx ids missing from the traceability table", missing_new_md),
        ("U-xx ids missing from the case JSON", missing_new_js),
    ):
        rep.detail("  %-46s : %s" % (label, ", ".join(missing) if missing else "none"))
        if missing:
            problems.append("%s: %s" % (label, ", ".join(missing)))

    rep.detail("")
    rep.detail("  where each required SIT id landed:")
    if md_by_id:
        owner = {}
        for uat_id, refs in md_by_id.items():
            for ref in refs:
                owner.setdefault(ref, []).append(uat_id)
        for sit_id in REQUIRED_SIT_IDS:
            rep.detail("      %-8s -> %s" % (sit_id,
                                             ", ".join(owner.get(sit_id, [])) or "MISSING"))
        rep.detail("  where each required U-xx id landed:")
        for new_id in REQUIRED_NEW_IDS:
            rep.detail("      %-8s -> %s" % (new_id,
                                             ", ".join(owner.get(new_id, [])) or "MISSING"))

        extra_sit = sorted(md_sit - set(REQUIRED_SIT_IDS))
        rep.detail("")
        rep.detail("  INFO (not a failure) — SIT ids carried beyond the required list: %s"
                   % (", ".join(extra_sit) if extra_sit else "none"))

    rep.check(9, "no coverage dropped by the 68 -> 54 de-duplication",
              not problems, "; ".join(problems))
    return not problems


def check_10_run_order(rep, wb):
    """The Run Order column has to be a runnable sequence, not a decoration.

    Wave order is NOT execution order — nine cases must run outside their own
    wave.  A tester who obeys Wave order reaches UAT-40, lands the department
    in ``Approved`` (irreversible without a direct database edit) and
    permanently loses UAT-41, UAT-42 step 5 and UAT-44, on production, inside
    the round.  That is the defect this column was added to prevent, so the
    column is PROVEN against the ``runs_after`` graph in the case JSON, never
    trusted because ``build_uat_test_cases.py`` happened to write it.

    Two assertions, both required:
      (a) the printed column is a permutation of 1..54 — every position used
          exactly once, so the sequence has no hole, no repeat and no gap;
      (b) for every edge ``dep -> case`` declared in ``runs_after``, the
          case's Run Order is STRICTLY greater than the dependency's.
    """
    rep.head(10, "Run Order is a permutation of 1..%d honouring every "
                 "'runs_after' edge" % EXPECTED_CASE_COUNT)
    ws = wb[SHEET_CASES]
    problems = []

    # ---- read the column exactly as the tester will read it, off the sheet
    run_by_id = {}
    bad_type = []
    dupe_ids = []
    rows_seen = 0
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        raw_id = ws.cell(r, IDX_CASE_ID).value
        if raw_id is None or str(raw_id).strip() == "":
            continue
        case_id = str(raw_id).strip()
        rows_seen += 1
        value = ws.cell(r, IDX_RUN_ORDER).value
        if isinstance(value, bool) or not isinstance(value, int):
            bad_type.append("%s=%r" % (case_id, value))
        elif case_id in run_by_id:
            dupe_ids.append(case_id)
        else:
            run_by_id[case_id] = value

    rep.detail("  column          : %s  (header '%s')"
               % (get_column_letter(IDX_RUN_ORDER),
                  EXPECTED_CASE_HEADERS[IDX_RUN_ORDER - 1]))
    rep.detail("  data rows read  : %d" % rows_seen)
    rep.detail("  integer values  : %d    non-integer: %d    duplicate ids: %d"
               % (len(run_by_id), len(bad_type), len(dupe_ids)))

    if rows_seen == 0:
        problems.append("no data rows on '%s' — this check would pass vacuously"
                        % SHEET_CASES)
    elif rows_seen != EXPECTED_CASE_COUNT:
        problems.append("%d data rows carry an id, expected %d"
                        % (rows_seen, EXPECTED_CASE_COUNT))
    if bad_type:
        problems.append("Run Order is not an integer on: %s"
                        % ", ".join(bad_type[:8]))
    if dupe_ids:
        problems.append("duplicate case id(s) while reading Run Order: %s"
                        % ", ".join(sorted(set(dupe_ids))))

    # ---- (a) permutation of 1..N
    values = sorted(run_by_id.values())
    expected_perm = list(range(1, EXPECTED_CASE_COUNT + 1))
    if values == expected_perm:
        rep.detail("  permutation     : OK — 1..%d, each position used once"
                   % EXPECTED_CASE_COUNT)
    else:
        absent = [n for n in expected_perm if n not in values]
        repeated = sorted({n for n in values if values.count(n) > 1})
        outside = sorted({n for n in values if n < 1 or n > EXPECTED_CASE_COUNT})
        bits = []
        if absent:
            bits.append("positions never used: %s" % ", ".join(map(str, absent[:10])))
        if repeated:
            bits.append("positions used twice: %s" % ", ".join(map(str, repeated[:10])))
        if outside:
            bits.append("positions outside 1..%d: %s"
                        % (EXPECTED_CASE_COUNT, ", ".join(map(str, outside[:10]))))
        if not bits:
            bits.append("got %r" % values)
        rep.detail("  permutation     : BROKEN — %s" % "; ".join(bits))
        problems.append("Run Order is not a permutation of 1..%d (%s)"
                        % (EXPECTED_CASE_COUNT, "; ".join(bits)))

    # ---- (b) every runs_after edge respected, strictly
    cases_json, err = load_case_json()
    if err:
        problems.append("cannot verify Run Order against the case JSON: %s" % err)
        rep.check(10, "Run Order is a runnable 1..%d sequence, every "
                      "'runs_after' edge respected" % EXPECTED_CASE_COUNT,
                  False, "; ".join(problems))
        return False

    json_ids = {str(c.get("id")) for c in cases_json}
    only_json = sorted(json_ids - set(run_by_id))
    only_sheet = sorted(set(run_by_id) - json_ids)
    rep.detail("  cases in JSON   : %d    ids only in JSON: %s    only on sheet: %s"
               % (len(cases_json),
                  ", ".join(only_json) if only_json else "none",
                  ", ".join(only_sheet) if only_sheet else "none"))
    if only_json:
        problems.append("case(s) in the JSON with no Run Order on the sheet: %s"
                        % ", ".join(only_json))
    if only_sheet:
        problems.append("row(s) on the sheet with no matching case in the JSON: %s"
                        % ", ".join(only_sheet))

    edges = 0
    violations = []
    no_dep_field = []
    for case in cases_json:
        case_id = str(case.get("id"))
        deps = case.get("runs_after")
        if deps is None:
            no_dep_field.append(case_id)
            continue
        if not isinstance(deps, list):
            problems.append("%s 'runs_after' is %s, expected a list"
                            % (case_id, type(deps).__name__))
            continue
        for dep in deps:
            edges += 1
            dep = str(dep)
            if dep == case_id:
                violations.append("%s lists itself in runs_after" % case_id)
            elif dep not in run_by_id:
                violations.append("%s runs_after unknown id %s" % (case_id, dep))
            elif case_id not in run_by_id:
                violations.append("%s has no Run Order on the sheet" % case_id)
            elif run_by_id[dep] >= run_by_id[case_id]:
                violations.append(
                    "%s (order %d) must run AFTER %s (order %d)"
                    % (case_id, run_by_id[case_id], dep, run_by_id[dep]))

    rep.detail("  runs_after edges: %d over %d cases" % (edges, len(cases_json)))
    rep.detail("  ordering breaks : %s"
               % ("; ".join(violations) if violations else "none"))
    if no_dep_field:
        problems.append("case(s) with no 'runs_after' array: %s"
                        % ", ".join(no_dep_field))
    if edges == 0:
        problems.append("no runs_after edge found in the case JSON — this check "
                        "would pass vacuously")
    if violations:
        problems.append("Run Order violates runs_after: %s"
                        % "; ".join(violations[:6]))

    # ---- the printed sequence, and the irreversible anchor
    if run_by_id:
        seq = [cid for cid, _n in sorted(run_by_id.items(), key=lambda kv: kv[1])]
        rep.detail("")
        rep.detail("  printed sequence:")
        for start in range(0, len(seq), 9):
            rep.detail("      %s" % " -> ".join(seq[start:start + 9]))
        if "UAT-40" in run_by_id:
            before_40 = [c for c in cases_json
                         if "UAT-40" in (c.get("runs_after") or [])]
            rep.detail("")
            rep.detail("  UAT-40 (irreversible 3rd-step Approve) sits at position "
                       "%d of %d; %d case(s) declare a dependency ON it"
                       % (run_by_id["UAT-40"], len(run_by_id), len(before_40)))
            deps_of_40 = next((c.get("runs_after") or [] for c in cases_json
                               if str(c.get("id")) == "UAT-40"), [])
            late = [d for d in deps_of_40
                    if d in run_by_id and run_by_id[d] > run_by_id["UAT-40"]]
            rep.detail("  cases that MUST finish before UAT-40: %d; still ordered "
                       "after it: %s"
                       % (len(deps_of_40), ", ".join(late) if late else "none"))

    rep.check(10, "Run Order is a runnable 1..%d sequence, every 'runs_after' "
                  "edge respected" % EXPECTED_CASE_COUNT,
              not problems, "; ".join(problems))
    return not problems


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------
def run_check(rep, n, title, fn, *args, default=None):
    """Run one check; turn any exception into that check's FAIL.

    A check that raises used to kill the whole run (a renamed sheet made
    check 6 raise KeyError and the validator printed nothing at all).  A pack
    broken badly enough to crash a check is exactly the pack that most needs a
    verdict, so the crash IS the failure — recorded, named, and the remaining
    checks still run.
    """
    try:
        return fn(*args)
    except Exception as exc:                      # noqa: BLE001 - deliberate
        rep.detail("")
        rep.detail("  EXCEPTION: %s: %s" % (type(exc).__name__, exc))
        rep.check(n, title, False,
                  "check crashed: %s: %s" % (type(exc).__name__, exc))
        return default


CHECK_TITLES = {
    1: "4 sheets present with the expected names",
    2: "header row is row 2; data starts row 3; no TC-EX row",
    3: "54 rows UAT-01..UAT-54 contiguous, no duplicates",
    4: "every Summary range ends at the real last data row",
    5: "8 modules match the sheet and totals sum to 54",
    6: "every required DV list present and covering the real range",
    7: "prod URL on sheet 1, staging hostname absent",
    8: "UAT_Test_Script.md traceability table has 54 rows",
    9: "no coverage dropped by the 68 -> 54 de-duplication",
    10: "Run Order is a runnable 1..%d sequence, every 'runs_after' edge "
        "respected" % EXPECTED_CASE_COUNT,
}


def main():
    rep = Report()
    rep.detail("UAT pack validation report")
    rep.detail("  workbook : %s" % XLSX)
    rep.detail("  markdown : %s" % MD)
    rep.detail("  cases    : %s" % CASES_DIR)
    rep.detail("")

    if not XLSX.exists():
        print("FATAL: workbook not found: %s" % XLSX)
        return 1

    wb = load_workbook(XLSX)   # formulas, not cached values

    T = CHECK_TITLES
    run_check(rep, 1, T[1], check_1_sheets, rep, wb, default=False)
    run_check(rep, 2, T[2], check_2_header, rep, wb, default=False)
    _ok3, last_data_row = run_check(rep, 3, T[3], check_3_rows, rep, wb,
                                    default=(False, None))
    if last_data_row is None:
        last_data_row = HEADER_ROW + EXPECTED_CASE_COUNT
    run_check(rep, 4, T[4], check_4_summary_ranges, rep, wb, last_data_row,
              default=False)
    run_check(rep, 5, T[5], check_5_modules, rep, wb, last_data_row,
              default=False)
    run_check(rep, 6, T[6], check_6_dv, rep, wb, last_data_row, default=False)
    run_check(rep, 7, T[7], check_7_urls, rep, wb, default=False)
    _ok8, trace_rows = run_check(rep, 8, T[8], check_8_markdown, rep,
                                 default=(False, None))
    run_check(rep, 9, T[9], check_9_coverage, rep, trace_rows, default=False)
    run_check(rep, 10, T[10], check_10_run_order, rep, wb, default=False)

    # every check must have produced exactly one verdict — a check that
    # neither passed, failed nor crashed would be a silent hole in the run
    recorded = [n for n, _t, _ok, _r in rep.results]
    for n in sorted(CHECK_TITLES):
        if recorded.count(n) != 1:
            rep.check(n, CHECK_TITLES[n], False,
                      "check produced %d verdicts, expected exactly 1"
                      % recorded.count(n))

    n_pass = sum(1 for _n, _t, ok, _r in rep.results if ok)
    n_fail = len(rep.results) - n_pass
    rep.detail("")
    rep.detail("=" * 78)
    rep.detail("SUMMARY: %d PASS  %d FAIL  (of %d checks)"
               % (n_pass, n_fail, len(rep.results)))
    rep.detail("=" * 78)

    with open(REPORT, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(rep.lines) + "\n")

    for n, title, ok, reason in rep.results:
        line = "CHECK %d  %s  %s" % (n, "PASS" if ok else "FAIL", title)
        if not ok and reason:
            line += "\n         reason: %s" % reason
        print(line)
    print("SUMMARY: %d PASS  %d FAIL   (full report: %s)"
          % (n_pass, n_fail, REPORT.name))
    return 0 if n_fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
