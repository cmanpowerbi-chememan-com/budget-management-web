#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Generate the UAT test-case workbook from the eight wave JSON files.

    python -X utf8 requirement_spec/5_uat/_build/build_uat_test_cases.py

Inputs  : .scratch/uat-prep/cases/W1.json ... W8.json   (54 cases total)
Outputs : requirement_spec/5_uat/Budget_UAT_Test_Cases_07.09.2026.xlsx
          requirement_spec/5_uat/UAT_Test_Script.md   (section 8 table,
          rewritten between the GENERATED TRACEABILITY markers, plus every
          quoted UAT-40 run-order position in the prose — see
          _sync_uat40_order: that number is the irreversible-Approve guard
          and must never drift from the computed sequence)

The workbook copies the shape of the SIT workbook
(requirement_spec/4_sit/Budget_SIT_Test_Cases_10.08.2026.xlsx) with six changes:

  1. new column ``Department`` after ``Module``
  2. new column ``Wave`` as the FIRST column, before ``No.``
  3. new column ``Run Order / ลำดับการรัน`` immediately after ``Wave`` — the
     topologically-sorted TRUE execution sequence 1..54 (see assign_run_order);
     Wave is only the grouping label, Run Order is the order to work in
  4. new column ``Defect ID`` between ``Severity`` and ``Tester``
  5. new sheet ``4. Defects``
  6. sheet ``3. Summary`` computes every range from the REAL last data row
     (the SIT workbook hard-coded ``$4:$64``), lists all 8 modules, and adds a
     breakdown by Wave and by Department

The SIT example row ``TC-EX`` is dropped; the worked example lives in the
instructions on sheet ``1. Info & Instructions`` instead.

Idempotent: running it twice produces a byte-identical .xlsx (zip entry
timestamps and document properties are pinned).

House rules honoured here: every open() passes encoding="utf-8"; run under
``python -X utf8`` because the Windows console is cp1252.
"""

from __future__ import annotations

import datetime as _dt
import json
import os
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------
# paths
# --------------------------------------------------------------------------
HERE = Path(__file__).resolve().parent           # requirement_spec/5_uat/_build
UAT_DIR = HERE.parent                            # requirement_spec/5_uat
REPO_ROOT = UAT_DIR.parent.parent                # repo root
CASES_DIR = REPO_ROOT / ".scratch" / "uat-prep" / "cases"
XLSX_OUT = UAT_DIR / "Budget_UAT_Test_Cases_07.09.2026.xlsx"
MD_OUT = UAT_DIR / "UAT_Test_Script.md"

WAVE_FILES = ["W1", "W2", "W3", "W4", "W5", "W6", "W7", "W8"]
EXPECTED_CASE_COUNT = 54

# --------------------------------------------------------------------------
# fixed facts (brief "Fixed header facts for sheet 1. Info & Instructions")
# --------------------------------------------------------------------------
PROD_URL = (
    "https://cman-budget-web-prd.kindstone-f34836dd.southeastasia."
    "azurecontainerapps.io/"
)

PROJECT_INFO = [
    ("Application", "Budget Web Application"),
    ("Environment", PROD_URL),
    ("Test Phase", "User Acceptance Test (UAT)"),
    ("Departments under test", "Data & Analytic ; Solution Delivery"),
    (
        "Fiscal year under test",
        "FY2027 — หน้าจอแสดงเป็น `Year 2026` (ปีวางแผน) · เป็นปีเดียวที่ระบบเปิดรับข้อมูล",
    ),
    ("Role — Filler (ผู้กรอกงบ) · ทั้งสองฝ่าย", "Pornthip (pornthipp@chememan.com)"),
    ("Role — Approver 1 (ผู้อนุมัติขั้นที่ 1) · ทั้งสองฝ่าย", "Laddawan (laddawank@chememan.com)"),
    ("Role — Approver 2 (ผู้อนุมัติขั้นที่ 2)", "Nipaporn (nipapornt@chememan.com)"),
    ("Role — Approver 3 (ผู้อนุมัติขั้นที่ 3)", "Waraporn (warapornt@chememan.com)"),
    (
        "Role — เจ้าของฝ่าย Solution Delivery ตัวจริง (ไม่ได้เป็นผู้ทดสอบเคสใดใน 54 เคส)",
        "Suchanya (suchanyay@chememan.com) — เธอเป็น “ผู้กรอกงบ (Filler)” ตัวจริงของฝ่าย Solution Delivery "
        "ในไฟล์ master `cc dept.xlsx` (10IT012000) และยังเป็นอยู่ · รอบนี้ขอ “ยืมฝ่าย” ของเธอมาทดสอบ "
        "โดยเพิ่ม Pornthip เข้าไปเป็นผู้กรอกอีกคน ไม่ได้ถอดสิทธิ์ของ Suchanya ออก\n"
        "รอบนี้ฝ่าย Solution Delivery ถูกส่งเข้าอนุมัติจริง (เคส UAT-36) และอนุมัติถึงขั้นที่ 1 ก่อนถูกตีกลับคืน "
        "(เคส UAT-43) ระหว่างนั้น **Suchanya แก้ไขฝ่ายตัวเองไม่ได้ชั่วคราว — ต้องแจ้งเธอล่วงหน้าก่อนเริ่มรอบ**\n"
        "หน้าที่ของเธอในรอบนี้คือ **ยืนยันว่างานของเธอไม่เสียหาย** — ทำ **หลังจบเคส UAT-43** (ตอนที่ฝ่ายถูกตีกลับคืน "
        "สถานะ `Rejected` ให้แก้ไขได้แล้ว) ล็อกอินตามปกติ เลือกฝ่าย Solution Delivery ปี Year 2026 แล้วยืนยัน 3 ข้อ: "
        "(1) ยังเห็นพื้นที่อัปโหลดไฟล์และปุ่มลบไฟล์ตามเดิม (2) กลับมากรอกยอดในตารางได้แล้ว "
        "(3) ตัวเลขเดิมของฝ่ายไม่เปลี่ยนไปเอง — แล้วแจ้งผลสั้น ๆ กลับให้ Test Lead "
        "(ไม่มีแถวหรือช่อง Status ของเธอเองในแท็บ “2. Test Cases”)\n"
        "ห้ามใช้บัญชีของเธอทดสอบเคส UAT-31 เพราะเคสนั้นต้องการคนที่ “เห็นฝ่ายได้แต่ไม่ใช่ผู้กรอก” "
        "ส่วน Suchanya เป็นผู้กรอกจริง จึงต้องเห็นปุ่มอัปโหลดและปุ่มลบ ซึ่งถูกต้องแล้ว",
    ),
    ("Test Lead", "Data and Analytics team"),
    ("Test Cycle / Round", "UAT Round 1"),
    ("Start Date", None),
    ("Target End Date", None),
]

WARNINGS = [
    "ข้อ 1 — UAT รอบนี้รันบน **ระบบจริง (production)** ใช้ฐานข้อมูลจริง และส่งอีเมลจริงถึงเพื่อนร่วมงานจริง "
    "ไม่มีระบบทดลองแยกต่างหาก ทุกอย่างที่กรอก ส่ง อนุมัติ หรือลบ คือข้อมูลจริงของปีงบประมาณ 2027",
    "ข้อ 2 — สถานะ `Approved` **ย้อนกลับไม่ได้ในระบบ** ต้องแก้ที่ฐานข้อมูลเท่านั้น "
    "อย่ากด Approve ขั้นที่ 3 (เคส UAT-40) จนกว่าทุกเคสที่มีเลข `ลำดับการรัน` น้อยกว่า {UAT40_ORDER} จะเสร็จครบ "
    "— UAT-40 อยู่ที่ลำดับการรันที่ {UAT40_ORDER} และเป็นเคสสุดท้ายของสายอนุมัติฝ่ายนั้น "
    "ถ้ากดก่อน เคส UAT-41 · UAT-42 (ข้อ 5) · UAT-44 · UAT-47 · UAT-52 (ข้อ 1) จะทดสอบไม่ได้อีกเลยตลอดรอบ",
    "ข้อ 3 — อายุ session ระหว่าง UAT ตั้งไว้ที่ **1 ชั่วโมง** (ค่าจริงคือ 14 ชั่วโมง และจะคืนค่าเดิมหลังจบรอบ) "
    "ถ้าระบบให้ล็อกอินใหม่ระหว่างทดสอบ ถือว่าเป็นเรื่องปกติของรอบนี้",
    "ข้อ 4 — **KNOWN ISSUE ที่ไม่ต้องแจ้งเป็น defect:** กล่อง “หมดเวลาการเข้าใช้งาน” ยังเขียนว่า 14 ชั่วโมง "
    "ทั้งที่ตั้งไว้ 1 ชั่วโมง — ตัดสินแล้วว่ายอมรับได้ (jakkaritw 2026-09-06)",
    "ข้อ 5 — **เตรียมไฟล์แนบให้ฝ่าย Solution Delivery ไว้ก่อนเริ่มรอบ อย่างน้อย 1 ไฟล์** "
    "ในโฟลเดอร์ `เอกสาร ฝ่าย/Solution Delivery/2027/` (ให้ Pornthip อัปโหลดด้วยขั้นตอนเดียวกับเคส UAT-27 "
    "ซึ่งอัปโหลดให้ฝ่าย Data & Analytic เท่านั้น ไม่มีเคสใดในรอบนี้อัปโหลดให้ฝ่าย Solution Delivery) "
    "— เคส UAT-31 ต้องมีไฟล์อยู่แล้วในโฟลเดอร์นี้ก่อนจึงจะทดสอบได้ ถ้าไม่เตรียมไว้ล่วงหน้า "
    "เคสนี้จะทำไม่ได้เมื่อถึงลำดับการรัน",
    "ข้อ 6 — **ฝ่าย Solution Delivery จะถูกล็อกชั่วคราวระหว่างรอบ (ตั้งแต่เคส UAT-36 ถึง UAT-43):** "
    "ทันทีที่ Pornthip กด Submit ฝ่ายนี้ (เคส UAT-36) คุณ Suchanya ผู้กรอกตัวจริงของฝ่าย Solution Delivery "
    "จะแก้ไขงบของฝ่ายตัวเองไม่ได้เลยจนกว่าจะถูกตีกลับ — **ต้องแจ้ง Suchanya ล่วงหน้าก่อนเริ่มรอบ** "
    "ว่าจะมีช่วงล็อกสั้น ๆ นี้ · เคส UAT-43 จะอนุมัติฝ่ายนี้ที่ขั้นที่ 1 แล้วตีกลับคืนทันทีในเคสเดียวกัน "
    "เพื่อให้เธอกลับมาแก้ไขได้ก่อนจบรอบ ไม่ต้องรอจนครบทั้ง 54 เคส",
]

HOW_TO = [
    "1.  ไปที่แท็บ “2. Test Cases” แล้วทดสอบ **เรียงตามคอลัมน์ `ลำดับการรัน` (Run Order) จาก 1 ถึง {N_CASES}** "
    "— **ไม่ใช่เรียงตามคอลัมน์ Wave** · คอลัมน์ Wave ยังเป็นชื่อกลุ่มของเคสและใช้กรองดูได้เหมือนเดิม "
    "แต่ลำดับที่ต้องลงมือทำจริงคือคอลัมน์ `ลำดับการรัน` เพราะบางเคสต้องทำข้ามระลอก "
    "(เช่น UAT-53 และ UAT-54 อยู่ระลอก W8 แต่ต้องทำก่อนระลอก W5 · UAT-03 อยู่ระลอก W1 แต่ต้องทำหลัง UAT-32)",
    "2.  วิธีเรียง: คลิกลูกศรกรองบนหัวคอลัมน์ `ลำดับการรัน` แล้วเลือก “Sort Smallest to Largest” "
    "จากนั้นไล่ทำจากบนลงล่าง",
    "3.  ⚠ **อย่ากด `Approve` ขั้นที่ 3 ในเคส UAT-40 (ลำดับการรันที่ {UAT40_ORDER}) "
    "จนกว่าทุกเคสที่มีลำดับการรันน้อยกว่า {UAT40_ORDER} จะเสร็จครบ** "
    "เพราะสถานะ `Approved` ย้อนกลับไม่ได้ ถ้ากดก่อน เคส UAT-41 · UAT-42 (ข้อ 5) · UAT-44 · UAT-47 · UAT-52 (ข้อ 1) "
    "จะทดสอบไม่ได้อีกเลยตลอดรอบ และต้องให้ทีม IT แก้ที่ฐานข้อมูลเท่านั้น",
    "4.  ทดสอบทีละเคสตาม “Test Steps” แล้วเทียบผลกับ “Expected Result”",
    "5.  **อ่าน “Expected Result” ให้จบทุกข้อก่อนตัดสินผล** — หลายเคสเป็นเคสรวม (merged) คือ 1 แถว "
    "แต่มีหลายข้อที่ต้องตรวจ ถ้าข้อใดข้อหนึ่งไม่ผ่าน ให้ทั้งแถวเป็น Fail แล้วระบุในช่อง Actual Result ว่าข้อไหนไม่ผ่าน",
    "6.  **ผู้ทดสอบกรอกเฉพาะช่องพื้นหลังสีเหลือง 7 ช่องนี้:** Actual Result, Status, Severity, Defect ID, "
    "Tester, Test Date, Remarks",
    "7.  เลือก Status จาก dropdown: Pass / Fail / Blocked / N/A · **ช่องที่ยังว่าง = `Not Run` "
    "(ยังไม่ได้ทดสอบ) ซึ่งเป็นค่าเริ่มต้น** ไม่ต้องพิมพ์คำว่า Not Run เอง "
    "— แท็บ “3. Summary” นับช่องว่างเป็น Not Run ให้อัตโนมัติ",
    "8.  ถ้า Fail ให้ระบุ Severity บันทึกรายละเอียดใน “Actual Result” แล้ว **เปิดแถวใหม่ในแท็บ “4. Defects”** "
    "จากนั้นนำเลข Defect ID (D-001, D-002, …) มาใส่ในคอลัมน์ Defect ID ของแถวนี้ "
    "· ช่อง Severity ของแถว defect ใช้ **ชุดค่าเดียวกัน** (Critical / High / Medium / Low) "
    "ยกค่าจากแถวเคสไปใส่ได้ตรง ๆ ไม่ต้องแปลงเป็นคำอื่น",
    "9.  แท็บ “3. Summary” สรุปผลอัตโนมัติ (จำนวน / อัตราผ่าน / แยกตาม Module, Wave และฝ่าย) ไม่ต้องกรอกเอง",
    "10. **ข้อยกเว้นเดียวของกติกา “กรอกเฉพาะช่องสีเหลือง” คือคอลัมน์ `PIC`** ซึ่งเป็นช่องสีขาว "
    "และคนที่กรอกไม่ใช่ผู้ทดสอบ แต่เป็น **Test Lead** ที่กรอกชื่อผู้รับผิดชอบรายเคสไว้ **ก่อน** เริ่มรอบทดสอบ "
    "· ระหว่างทดสอบ ผู้ทดสอบไม่ต้องแตะคอลัมน์นี้",
    "11. ตัวหนังสือ **สีม่วง** คือข้อความที่ AI ร่างไว้ให้ ยังไม่ผ่านการตรวจจากผู้ทดสอบ "
    "ถ้าพบว่าไม่ตรงกับระบบจริง แก้ได้เลยแล้วบันทึกไว้ในช่อง Remarks "
    "· **ยกเว้นคอลัมน์ `Wave` และ `ลำดับการรัน` ซึ่งพิมพ์ด้วยตัวสีดำ** เพราะเป็นค่าที่โปรแกรมคำนวณมาให้ "
    "ไม่ใช่ข้อความที่ AI ร่าง — ห้ามแก้ ถ้าเห็นว่าลำดับผิดให้แจ้ง Test Lead แทน",
    "12. อ่านเอกสารประกอบ `UAT_Test_Script.md` (ภาษาไทย) ก่อนเริ่ม — อธิบายลำดับการรัน บทบาทของแต่ละคน "
    "และสิ่งที่ตัดออกจากรอบนี้โดยตั้งใจ",
]

WORKED_EXAMPLE = [
    "เคสตัวอย่าง UAT-09 (ระลอก W2) — “กรอกยอดแล้วคลิกออกจากช่อง = บันทึกอัตโนมัติ · refresh หน้าแล้วค่ายังอยู่”",
    "ถ้าผ่าน → Actual Result = “กรอก 100 ที่เดือน ม.ค. คลิกออกจากช่อง แล้ว refresh หน้า ค่ายังเป็น 100” · "
    "Status = `Pass` · Severity = `-` · Defect ID = (เว้นว่าง) · Tester = `Pornthip` · Test Date = วันที่ทดสอบ · Remarks = (เว้นว่าง)",
    "ถ้าไม่ผ่าน → Actual Result = “refresh แล้วค่าหาย กลับเป็นช่องว่าง” · Status = `Fail` · Severity = `High` · "
    "Defect ID = `D-001` แล้วไปเปิดแถว D-001 ในแท็บ “4. Defects” กรอกหัวข้อ ผู้แจ้ง วันที่ และสถานะ `Open`",
    "ถ้าทดสอบไม่ได้เพราะติดเคสอื่นหรือข้อมูลยังไม่พร้อม → Status = `Blocked` และเขียนเหตุผลใน Remarks "
    "(อย่าปล่อยเป็น `Not Run` ถ้าพยายามทดสอบแล้ว)",
]

# 8 modules — the Summary breakdown must list exactly these
MODULES = [
    "Login / Authentication",
    "Budget Entry (Filler)",
    "Special GL Subform",
    "Attachments",
    "Submission / Workflow",
    "Approval",
    "Email Alert / Notification",
    "Business Rules & UI",
]

WAVE_TITLES = {
    "W1": "เข้าระบบและสิทธิ์ที่มองเห็น",
    "W2": "กรอกงบในตารางหลัก",
    "W3": "ฟอร์มย่อยและทริปเดินทาง",
    "W4": "แนบเอกสาร",
    "W5": "ส่งขออนุมัติ",
    "W6": "อนุมัติ 3 ขั้น",
    "W7": "อีเมลแจ้งเตือน",
    "W8": "หน้าตาและกติกาที่ต้องยอมรับ",
}

DEPARTMENTS = ["Data & Analytic", "Solution Delivery", "ทั้งสองฝ่าย", "-"]

# data-validation lists (SIT shapes kept verbatim where they existed)
DV_STATUS = '"Not Run,Pass,Fail,Blocked,N/A"'
DV_SEVERITY = '"Critical,High,Medium,Low,-"'
DV_PRIORITY = '"High,Medium,Low"'
DV_DEPARTMENT = '"Data & Analytic,Solution Delivery,ทั้งสองฝ่าย,-"'
# ONE severity vocabulary for the whole pack.  The SIT workbook used
# Critical/High/Medium/Low on the case sheet and Blocker/Major/Minor on the
# defect log; companion §5.4 tells the tester to carry a case's severity onto
# the defect row, where the second dropdown then rejected it.  The defect log
# now uses the same four values as the case sheet — nothing to translate.
# ('-' exists only on the case sheet, for a case that passed; a defect row
# exists only because something failed, so it never needs '-'.)
DV_DEFECT_SEVERITY = '"Critical,High,Medium,Low"'
DV_DEFECT_STATUS = '"Open,In Progress,Fixed,Retested,Closed,Wont Fix"'
DV_DEFECT_RETEST = '"Pass,Fail"'

# --------------------------------------------------------------------------
# palette (copied from the SIT workbook)
# --------------------------------------------------------------------------
C_TITLE = "FF1F3864"
C_HEADER = "FF2E5496"
C_LABEL = "FFD9E1F2"
C_WHITE = "FFFFFFFF"
C_YELLOW = "FFFFF2CC"
C_GREEN = "FFC6EFCE"
C_RED = "FFFFC7CE"
C_GREY = "FFF2F2F2"
C_BORDER = "FFBFBFBF"
C_AI_PURPLE = "FF7030A0"  # AI-written cells, per the SIT workbook editing protocol

FILL_TITLE = PatternFill("solid", fgColor=C_TITLE)
FILL_HEADER = PatternFill("solid", fgColor=C_HEADER)
FILL_LABEL = PatternFill("solid", fgColor=C_LABEL)
FILL_WHITE = PatternFill("solid", fgColor=C_WHITE)
FILL_YELLOW = PatternFill("solid", fgColor=C_YELLOW)
FILL_GREEN = PatternFill("solid", fgColor=C_GREEN)
FILL_RED = PatternFill("solid", fgColor=C_RED)
FILL_GREY = PatternFill("solid", fgColor=C_GREY)

THIN = Side(style="thin", color=C_BORDER)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BAND = Border(left=THIN, top=THIN, bottom=THIN)

AL_LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_LEFT_CENTER = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_CENTER_TOP = Alignment(horizontal="center", vertical="top", wrap_text=True)
AL_CENTER_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# fixed timestamp so a re-run produces a byte-identical file
PINNED_DT = _dt.datetime(2026, 9, 7, 0, 0, 0)
PINNED_ZIP_TS = (2026, 9, 7, 0, 0, 0)


# --------------------------------------------------------------------------
# load + validate the case data
# --------------------------------------------------------------------------
def load_cases():
    """Read the eight wave files, sort by ``no``, and sanity-check the set."""
    cases = []
    for wave in WAVE_FILES:
        path = CASES_DIR / f"{wave}.json"
        if not path.exists():
            raise SystemExit(f"FATAL: missing wave file {path}")
        with open(path, encoding="utf-8") as fh:
            chunk = json.load(fh)
        if not isinstance(chunk, list):
            raise SystemExit(f"FATAL: {path} is not a JSON list")
        cases.extend(chunk)

    cases.sort(key=lambda c: c["no"])

    if len(cases) != EXPECTED_CASE_COUNT:
        raise SystemExit(
            f"FATAL: expected {EXPECTED_CASE_COUNT} cases, found {len(cases)}"
        )
    ids = [c["id"] for c in cases]
    if len(set(ids)) != len(ids):
        dupes = sorted({i for i in ids if ids.count(i) > 1})
        raise SystemExit(f"FATAL: duplicate case ids {dupes}")
    expected_ids = [f"UAT-{n:02d}" for n in range(1, EXPECTED_CASE_COUNT + 1)]
    if ids != expected_ids:
        missing = sorted(set(expected_ids) - set(ids))
        extra = sorted(set(ids) - set(expected_ids))
        raise SystemExit(
            f"FATAL: case ids are not UAT-01..UAT-{EXPECTED_CASE_COUNT} "
            f"contiguous (missing={missing} extra={extra})"
        )
    used_modules = {c["module"] for c in cases}
    unknown = sorted(used_modules - set(MODULES))
    if unknown:
        raise SystemExit(f"FATAL: case module(s) not in the 8-module list: {unknown}")
    used_depts = {c["department"] for c in cases}
    unknown_d = sorted(used_depts - set(DEPARTMENTS))
    if unknown_d:
        raise SystemExit(f"FATAL: case department(s) not in the DV list: {unknown_d}")
    assign_run_order(cases)
    return cases


def assign_run_order(cases):
    """Stamp each case with ``run_order`` — the TRUE execution sequence 1..N.

    ``no`` (and therefore ``Wave``) is the grouping order, not the execution
    order: nine cases have to run outside their wave.  A tester who follows
    Wave order reaches UAT-40, lands the department in ``Approved`` — which is
    irreversible without a direct database edit — and permanently loses
    UAT-41, UAT-42 step 5 and UAT-44, on production, inside the round.

    The sequence is a topological sort of the ``runs_after`` graph (an edge
    ``d -> c`` for every ``d`` in ``c["runs_after"]``), with the existing
    ``no`` order as the tie-break so the result stays as close to wave order
    as the dependencies allow.  A cycle or an unknown id is FATAL: emitting a
    wrong order would be worse than emitting none.
    """
    import heapq

    by_id = {c["id"]: c for c in cases}
    deps = {}
    for case in cases:
        raw = case.get("runs_after")
        if raw is None:
            raise SystemExit(
                f"FATAL: {case['id']} has no 'runs_after' array — every case "
                "must declare one (use [] when it depends on nothing)"
            )
        if not isinstance(raw, list):
            raise SystemExit(
                f"FATAL: {case['id']} 'runs_after' is {type(raw).__name__}, "
                "expected a list"
            )
        unknown = [d for d in raw if d not in by_id]
        if unknown:
            raise SystemExit(
                f"FATAL: {case['id']} 'runs_after' names unknown case id(s): "
                f"{', '.join(map(str, unknown))}"
            )
        if case["id"] in raw:
            raise SystemExit(f"FATAL: {case['id']} 'runs_after' lists itself")
        deps[case["id"]] = set(raw)

    dependents = {c["id"]: [] for c in cases}
    for cid, ds in deps.items():
        for d in ds:
            dependents[d].append(cid)

    indegree = {cid: len(ds) for cid, ds in deps.items()}
    ready = [(by_id[cid]["no"], cid) for cid, n in indegree.items() if n == 0]
    heapq.heapify(ready)

    order = []
    while ready:
        _no, cid = heapq.heappop(ready)
        order.append(cid)
        for dep in dependents[cid]:
            indegree[dep] -= 1
            if indegree[dep] == 0:
                heapq.heappush(ready, (by_id[dep]["no"], dep))

    if len(order) != len(cases):
        stuck = sorted(set(deps) - set(order))
        raise SystemExit(
            "FATAL: 'runs_after' contains a cycle — these cases can never be "
            f"ordered: {', '.join(stuck)}"
        )

    for position, cid in enumerate(order, start=1):
        by_id[cid]["run_order"] = position
    return order


# --------------------------------------------------------------------------
# small helpers
# --------------------------------------------------------------------------
def _put(ws, coord, value, *, font=None, fill=None, align=None,
         border=BORDER_ALL, numfmt=None):
    cell = ws[coord]
    if not isinstance(cell, MergedCell):
        # a MergedCell only carries style; its value lives in the anchor cell
        cell.value = value
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if align is not None:
        cell.alignment = align
    if border is not None:
        cell.border = border
    if numfmt is not None:
        cell.number_format = numfmt
    return cell


def _set_widths(ws, widths):
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _estimate_row_height(pairs, floor=48.0, ceiling=240.0, line_pt=12.0):
    """Rough auto-fit: ``pairs`` is [(text, column_width_in_chars), ...]."""
    lines = 1
    for text, width in pairs:
        if not text:
            continue
        chars = max(8, int(width * 0.95))
        count = 0
        for para in str(text).split("\n"):
            count += max(1, -(-len(para) // chars))
        lines = max(lines, count)
    return min(ceiling, max(floor, lines * line_pt))


# --------------------------------------------------------------------------
# sheet 1 — Info & Instructions
# --------------------------------------------------------------------------
def _subst(text, n_cases, uat40_order):
    """Fill the two numbers that must never be typed by hand into sheet-1 prose.

    ``{UAT40_ORDER}`` is the run-order position of UAT-40 — the irreversible
    Approve step — so the warning can never drift away from the real sequence.
    """
    return (text
            .replace("{UAT40_ORDER}", str(uat40_order))
            .replace("{N_CASES}", str(n_cases)))


def build_info_sheet(ws, n_cases, last_row, uat40_order):
    ws.sheet_view.showGridLines = False
    _set_widths(ws, {"A": 34.0, "B": 92.0, "C": 3.0})

    f_title = Font(name="Aptos", size=12, bold=True, color=C_WHITE)
    f_sub = Font(name="Aptos", size=8, bold=False, color=C_WHITE)
    f_section = Font(name="Aptos", size=11, bold=True, color=C_WHITE)
    f_label = Font(name="Aptos", size=10, bold=True, color="FF000000")
    f_value = Font(name="Aptos", size=10, bold=False, color="FF000000")
    f_ai = Font(name="Aptos", size=10, bold=False, color=C_AI_PURPLE)
    f_small = Font(name="Aptos", size=9, bold=False, color="FF808080")

    def band(row, text, font, fill, height=None):
        ws.merge_cells(f"A{row}:B{row}")
        _put(ws, f"A{row}", text, font=font, fill=fill,
             align=AL_LEFT_CENTER if fill is FILL_HEADER or fill is FILL_TITLE
             else AL_LEFT_TOP, border=BORDER_BAND)
        for col in ("B",):
            ws[f"{col}{row}"].fill = fill
            ws[f"{col}{row}"].border = BORDER_BAND
        if height:
            ws.row_dimensions[row].height = height

    band(1, "Budget Web Application — UAT Test Cases", f_title, FILL_TITLE, 21.0)
    band(2,
         "Chememan PCL (CMAN)  |  User Acceptance Testing (UAT)  |  Data & Analytics  "
         "|  Round 1  |  FY2027",
         f_sub, FILL_HEADER, 13.5)
    ws.row_dimensions[3].height = 6.0

    band(4, "Project Information", f_section, FILL_HEADER, 19.5)
    row = 5
    for label, value in PROJECT_INFO:
        _put(ws, f"A{row}", label, font=f_label, fill=FILL_LABEL, align=AL_LEFT_TOP)
        _put(ws, f"B{row}", value, font=f_value, fill=FILL_WHITE, align=AL_LEFT_TOP)
        # most rows are one line; a long role description must still be readable.
        # line_pt is 15.0 (not 13.5): Thai vowels/tone marks stack above and
        # below the baseline, so a wrapped Thai line needs more vertical room.
        ws.row_dimensions[row].height = _estimate_row_height(
            [(label, 34.0), (value, 92.0)], floor=16.5, ceiling=200.0, line_pt=15.0)
        row += 1

    _put(ws, f"A{row}", "Total test cases", font=f_label, fill=FILL_LABEL,
         align=AL_LEFT_TOP)
    _put(ws, f"B{row}",
         f"{n_cases} เคส (UAT-01 ถึง UAT-{n_cases:02d}) · จัดกลุ่มเป็น 8 ระลอก (Wave) "
         f"แต่ **ลำดับการทดสอบจริงให้ดูคอลัมน์ `ลำดับการรัน` 1 → {n_cases}** ไม่ใช่คอลัมน์ Wave · "
         f"ข้อมูลจริงอยู่แถวที่ 3 ถึง {last_row} ของแท็บ “2. Test Cases”",
         font=f_value, fill=FILL_WHITE, align=AL_LEFT_TOP)
    ws.row_dimensions[row].height = _estimate_row_height(
        [(ws[f"B{row}"].value, 92.0)], floor=16.5, ceiling=48.0, line_pt=13.5)
    row += 2

    band(row, "ข้อควรระวังก่อนเริ่มทดสอบ", f_section, FILL_HEADER, 19.5)
    row += 1
    for text in WARNINGS:
        text = _subst(text, n_cases, uat40_order)
        band(row, text, f_ai, FILL_WHITE)
        ws.row_dimensions[row].height = _estimate_row_height(
            [(text, 120)], floor=30.0, ceiling=75.0, line_pt=14.0)
        row += 1
    row += 1

    band(row, "วิธีใช้ไฟล์นี้ / How to use this workbook", f_section, FILL_HEADER, 19.5)
    row += 1
    for text in HOW_TO:
        text = _subst(text, n_cases, uat40_order)
        band(row, text, f_ai, FILL_WHITE)
        ws.row_dimensions[row].height = _estimate_row_height(
            [(text, 120)], floor=16.5, ceiling=75.0, line_pt=14.0)
        row += 1
    row += 1

    band(row, "ตัวอย่างการกรอก 1 แถว (worked example)", f_section, FILL_HEADER, 19.5)
    row += 1
    for text in WORKED_EXAMPLE:
        band(row, text, f_ai, FILL_WHITE)
        ws.row_dimensions[row].height = _estimate_row_height(
            [(text, 120)], floor=16.5, ceiling=60.0, line_pt=14.0)
        row += 1
    row += 1

    band(row, "Legend / คำอธิบายสี & สถานะ", f_section, FILL_HEADER, 19.5)
    row += 1
    legend = [
        (FILL_YELLOW,
         "ช่องพื้นหลังสีเหลือง = ช่องที่ผู้ทดสอบกรอกระหว่างรอบทดสอบ (7 ช่อง: Actual Result, Status, "
         "Severity, Defect ID, Tester, Test Date, Remarks)", None),
        (FILL_WHITE,
         "ช่อง `PIC` เป็นช่องสีขาว แต่เป็นข้อยกเว้นเดียวที่ต้องกรอก — Test Lead กรอกก่อนเริ่มรอบ "
         "ผู้ทดสอบไม่ต้องแตะ · ช่องสีขาวอื่นทั้งหมดคือเนื้อหาเคส ไม่ต้องแก้", None),
        (FILL_WHITE,
         "คอลัมน์ `ลำดับการรัน` (Run Order) = ลำดับที่ต้องลงมือทำจริง 1 → {N_CASES} · "
         "คอลัมน์ `Wave` = ชื่อกลุ่มของเคส ใช้กรองดูเท่านั้น **ไม่ใช่ลำดับการทำ** · "
         "สองคอลัมน์นี้พิมพ์ด้วย **ตัวอักษรสีดำ ไม่ใช่สีม่วง** เพราะโปรแกรมคำนวณมาให้ "
         "ไม่ใช่ข้อความที่ AI ร่าง — **ห้ามแก้ทั้งสองคอลัมน์**", None),
        (FILL_GREEN, "Status = Pass  (ผลตรงตามที่คาดหวังทุกข้อ)", None),
        (FILL_RED, "Status = Fail  (ผลไม่ตรงอย่างน้อย 1 ข้อ — ต้องเปิดแถวใน “4. Defects” และใส่เลข Defect ID)", None),
        (FILL_GREY, "Status = Blocked (ทดสอบไม่ได้ ติดเคสอื่นหรือข้อมูลยังไม่พร้อม) · N/A (ไม่เกี่ยวกับรอบนี้)", None),
        (FILL_WHITE,
         "ตัวหนังสือสีม่วง = ข้อความที่ AI ร่างไว้ ยังไม่ผ่านการตรวจจากคน — แก้ได้ถ้าไม่ตรงกับระบบจริง "
         "· ตัวหนังสือสีดำในคอลัมน์ `Wave` และ `ลำดับการรัน` = ค่าที่โปรแกรมคำนวณ ห้ามแก้",
         C_AI_PURPLE),
    ]
    for fill, text, text_color in legend:
        text = _subst(text, n_cases, uat40_order)
        _put(ws, f"A{row}", None, font=f_value, fill=fill, align=AL_LEFT_TOP)
        _put(ws, f"B{row}", text,
             font=Font(name="Aptos", size=10, color=text_color or "FF000000"),
             fill=FILL_WHITE, align=AL_LEFT_TOP)
        ws.row_dimensions[row].height = _estimate_row_height(
            [(text, 92.0)], floor=16.5, ceiling=48.0, line_pt=13.5)
        row += 1

    band(row,
         "Severity: Critical = ระบบใช้งานไม่ได้ / ข้อมูลผิดพลาดร้ายแรง  |  High = ฟังก์ชันหลักเสีย  |  "
         "Medium = มีทางเลี่ยง  |  Low = ปัญหาเล็กน้อย / UI  ·  "
         "ชุดค่านี้ใช้เหมือนกันทั้งแท็บ “2. Test Cases” และแท็บ “4. Defects” "
         "จึงยกค่าจากแถวเคสไปใส่แถว defect ได้ตรง ๆ ไม่ต้องแปลง "
         "(ค่า `-` มีเฉพาะแท็บ “2. Test Cases” ใช้เมื่อเคสนั้นไม่ Fail)",
         f_small, FILL_WHITE)
    ws.row_dimensions[row].height = 45.0


# --------------------------------------------------------------------------
# sheet 2 — Test Cases
# --------------------------------------------------------------------------
# A Wave | B Run Order | C No. | D Test Case ID | E Module | F Department |
# G PIC | H Test Scenario | I Priority | J Preconditions | K Test Steps |
# L Test Data | M Expected Result | N Actual Result | O Status | P Severity |
# Q Defect ID | R Tester | S Test Date | T Remarks
#
# `Run Order` sits immediately after `Wave` on purpose: Wave is the GROUPING
# label and the filter, Run Order is the SEQUENCE the tester must follow.
CASE_COLUMNS = [
    ("A", "Wave", 7.0, "center"),
    ("B", "Run Order / ลำดับการรัน", 10.0, "center"),
    ("C", "No.", 8.29, "center"),
    ("D", "Test Case ID", 12.0, "left"),
    ("E", "Module", 21.57, "left"),
    ("F", "Department", 18.0, "left"),
    ("G", "PIC", 12.0, "center"),
    ("H", "Test Scenario", 37.71, "left"),
    ("I", "Priority", 12.0, "center"),
    ("J", "Preconditions / เงื่อนไขก่อนทดสอบ", 34.71, "left"),
    ("K", "Test Steps", 47.29, "left"),
    ("L", "Test Data", 45.57, "left"),
    ("M", "Expected Result", 53.71, "left"),
    ("N", "Actual Result", 18.57, "left"),
    ("O", "Status", 11.0, "left"),
    ("P", "Severity", 12.0, "center"),
    ("Q", "Defect ID", 11.0, "center"),
    ("R", "Tester", 12.0, "left"),
    ("S", "Test Date", 12.0, "center"),
    ("T", "Remarks / หมายเหตุ", 25.43, "left"),
]
LAST_COL = CASE_COLUMNS[-1][0]
YELLOW_COLS = ["N", "O", "P", "Q", "R", "S", "T"]   # tester-filled
# AI-written = every case-content column, printed in purple, which sheet 1's
# legend defines as "AI drafted it, correct it if the real system disagrees".
# Three columns are deliberately absent:
#   'G' (PIC)                  — white, filled by the Test Lead before the round
#   'A' (Wave) + 'B' (Run Order) — NOT prose and NOT correctable: both are
#       computed by this generator (Run Order is the topological sort of the
#       'runs_after' graph), and companion §5.1 forbids editing them.  Painting
#       them purple invited the tester to overwrite the one column the whole
#       round's safety depends on, so they print in plain black instead.
AI_COLS = ["C", "D", "E", "F", "H", "I", "J", "K", "L", "M"]


def build_cases_sheet(ws, cases):
    ws.sheet_view.showGridLines = False
    _set_widths(ws, {c[0]: c[2] for c in CASE_COLUMNS})

    f_title = Font(name="Arial", size=13, bold=True, color=C_WHITE)
    f_header = Font(name="Arial", size=10, bold=True, color=C_WHITE)
    f_ai = Font(name="Arial", size=9, bold=False, color=C_AI_PURPLE)
    f_plain = Font(name="Arial", size=9, bold=False, color="FF000000")

    # title band, row 1
    ws.merge_cells(f"A1:{LAST_COL}1")
    _put(ws, "A1", "Budget Web App — UAT Test Cases (Round 1 · FY2027 · production)",
         font=f_title, fill=FILL_TITLE, align=AL_LEFT_CENTER, border=BORDER_BAND)
    for col, *_ in CASE_COLUMNS[1:]:
        ws[f"{col}1"].fill = FILL_TITLE
        ws[f"{col}1"].border = BORDER_BAND
    ws.row_dimensions[1].height = 24.0

    # header row 2
    for col, header, _w, _al in CASE_COLUMNS:
        _put(ws, f"{col}2", header, font=f_header, fill=FILL_HEADER,
             align=AL_CENTER_CENTER)
    ws.row_dimensions[2].height = 33.75

    widths = {c[0]: c[2] for c in CASE_COLUMNS}
    aligns = {c[0]: c[3] for c in CASE_COLUMNS}

    first_row = 3
    for offset, case in enumerate(cases):
        row = first_row + offset
        values = {
            "A": case["wave"],
            "B": case["run_order"],
            "C": case["no"],
            "D": case["id"],
            "E": case["module"],
            "F": case["department"],
            "G": None,                      # PIC — Test Lead fills before the round
            "H": case["scenario"],
            "I": case["priority"],
            "J": case["preconditions"],
            "K": case["steps"],
            "L": case["test_data"],
            "M": _expected_with_note(case),
            "N": None, "O": None, "P": None, "Q": None,
            "R": None, "S": None, "T": None,
        }
        for col, _hdr, _w, al in CASE_COLUMNS:
            fill = FILL_YELLOW if col in YELLOW_COLS else FILL_WHITE
            font = f_ai if col in AI_COLS else f_plain
            align = AL_CENTER_TOP if al == "center" else AL_LEFT_TOP
            cell = _put(ws, f"{col}{row}", values[col], font=font, fill=fill,
                        align=align)
            if col == "S":
                cell.number_format = "dd/mm/yyyy"
        # Excel does not auto-fit rows that carry no stored height, so the
        # height is computed here.  The ceiling is deliberately generous: a
        # tester must be able to read the WHOLE Expected Result before marking
        # Pass (many cases are merged: one row, several checks).
        ws.row_dimensions[row].height = _estimate_row_height(
            [(values[c], widths[c]) for c in ("H", "J", "K", "L", "M")],
            floor=48.0, ceiling=620.0,
        )

    last_row = first_row + len(cases) - 1

    # data validation — every list covers the REAL used range
    dvs = [
        (DV_PRIORITY, f"I{first_row}:I{last_row}"),
        (DV_DEPARTMENT, f"F{first_row}:F{last_row}"),
        (DV_STATUS, f"O{first_row}:O{last_row}"),
        (DV_SEVERITY, f"P{first_row}:P{last_row}"),
    ]
    for formula, sqref in dvs:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            showDropDown=False, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(sqref)

    ws.auto_filter.ref = f"A2:{LAST_COL}{last_row}"
    # keep Wave / Run Order / No. / Test Case ID / Module on screen while
    # scrolling right — the tester works down Run Order, so it must stay visible
    ws.freeze_panes = "F3"
    return last_row


def _expected_with_note(case):
    """Expected Result + the case's own note, so no assertion is lost."""
    expected = (case.get("expected") or "").rstrip()
    note = (case.get("note") or "").strip()
    if case.get("merged"):
        expected = (
            "⚠ เคสรวม (merged) — 1 แถวแต่มีหลายข้อที่ต้องตรวจ "
            "อ่านให้ครบทุกข้อก่อนตัดสิน Pass\n" + expected
        )
    if note:
        expected += "\n\n— หมายเหตุสำคัญ —\n" + note
    return expected


# --------------------------------------------------------------------------
# sheet 3 — Summary
# --------------------------------------------------------------------------
def build_summary_sheet(ws, cases, last_row):
    ws.sheet_view.showGridLines = False
    _set_widths(ws, {
        "A": 2.29, "B": 26.0, "C": 8.0, "D": 8.0, "E": 8.0,
        "F": 2.29, "G": 8.0, "H": 30.0, "I": 8.0, "J": 8.0, "K": 8.0,
        "L": 2.29, "M": 20.0, "N": 8.0, "O": 8.0, "P": 8.0,
    })

    # column letters follow CASE_COLUMNS — they moved one to the right when the
    # 'Run Order' column was inserted after 'Wave'
    tc = "'2. Test Cases'"
    first = 3
    _col = {hdr.split(" / ")[0]: letter for letter, hdr, _w, _a in CASE_COLUMNS}
    rng_id = f"{tc}!${_col['Test Case ID']}${first}:${_col['Test Case ID']}${last_row}"
    rng_module = f"{tc}!${_col['Module']}${first}:${_col['Module']}${last_row}"
    rng_wave = f"{tc}!${_col['Wave']}${first}:${_col['Wave']}${last_row}"
    rng_dept = f"{tc}!${_col['Department']}${first}:${_col['Department']}${last_row}"
    rng_status = f"{tc}!${_col['Status']}${first}:${_col['Status']}${last_row}"

    f_title = Font(name="Arial", size=15, bold=True, color=C_WHITE)
    f_sub = Font(name="Arial", size=10, bold=False, color=C_WHITE)
    f_section = Font(name="Arial", size=11, bold=True, color=C_WHITE)
    f_colhdr = Font(name="Arial", size=10, bold=True, color=C_WHITE)
    f_label = Font(name="Arial", size=10, bold=False, color="FF000000")
    f_label_b = Font(name="Arial", size=10, bold=True, color="FF000000")
    f_num = Font(name="Arial", size=11, bold=True, color="FF000000")
    f_small_lbl = Font(name="Arial", size=9, bold=False, color="FF000000")
    f_small_num = Font(name="Arial", size=10, bold=False, color="FF000000")
    f_note = Font(name="Arial", size=9, bold=False, color="FF808080")

    def band(rng, text, font, fill, height=None):
        ws.merge_cells(rng)
        first_coord = rng.split(":")[0]
        _put(ws, first_coord, text, font=font, fill=fill, align=AL_LEFT_CENTER,
             border=BORDER_BAND)
        col_a, row_a = first_coord[0], int(first_coord[1:])
        col_b = rng.split(":")[1][0]
        for oc in range(ord(col_a) + 1, ord(col_b) + 1):
            c = ws[f"{chr(oc)}{row_a}"]
            c.fill = fill
            c.border = BORDER_BAND
        if height:
            ws.row_dimensions[row_a].height = height

    band("B2:P2", "UAT Test Execution — Summary  ·  Round 1  ·  FY2027", f_title,
         FILL_TITLE, 27.75)
    band("B3:P3",
         "สรุปผลอัตโนมัติจากแท็บ “2. Test Cases” (ไม่ต้องกรอกเอง) · "
         f"ทุกสูตรนับช่วงแถวจริง {first}–{last_row}",
         f_sub, FILL_HEADER, 15.75)

    # --- headline block -------------------------------------------------
    _put(ws, "B5", "Result", font=f_colhdr, fill=FILL_HEADER, align=AL_LEFT_TOP)
    ws.merge_cells("C5:E5")
    for col in ("C", "D", "E"):
        _put(ws, f"{col}5", "Count" if col == "C" else None, font=f_colhdr,
             fill=FILL_HEADER, align=AL_CENTER_TOP)

    headline = [
        (6, "Total test cases", f"=COUNTA({rng_id})", FILL_WHITE, f_label_b, None),
        (7, "Pass", f'=COUNTIF({rng_status},"Pass")', FILL_GREEN, f_label, None),
        (8, "Fail", f'=COUNTIF({rng_status},"Fail")', FILL_RED, f_label, None),
        (9, "Blocked", f'=COUNTIF({rng_status},"Blocked")', FILL_YELLOW, f_label, None),
        # 'Not Run' is the REMAINDER, not a COUNTIF.  The Status column ships
        # empty and sheet 1 calls Not Run the default, so a COUNTIF("Not Run")
        # read 0 on day one — the first number a Test Lead looks at.  Computing
        # it as Total minus the four executed statuses makes the five result
        # rows always reconcile to Total, counts an empty cell (the real
        # default) and an explicit "Not Run" alike, and leaves the yellow
        # tester-owned column genuinely empty so "blank = ยังไม่ทำ" still filters.
        (10, "Not Run (ยังไม่ได้ทดสอบ)", "=C6-C7-C8-C9-C11",
         FILL_GREY, f_label, None),
        (11, "N/A", f'=COUNTIF({rng_status},"N/A")', FILL_GREY, f_label, None),
        (12, "Pass Rate (Pass / Executed)",
         f'=IFERROR(COUNTIF({rng_status},"Pass")/'
         f'(COUNTIF({rng_status},"Pass")+COUNTIF({rng_status},"Fail")'
         f'+COUNTIF({rng_status},"Blocked")),"-")',
         FILL_LABEL, f_label_b, "0.0%"),
    ]
    for row, label, formula, fill, font, numfmt in headline:
        _put(ws, f"B{row}", label, font=font, fill=fill, align=AL_LEFT_TOP)
        ws.merge_cells(f"C{row}:E{row}")
        _put(ws, f"C{row}", formula, font=f_num, fill=fill, align=AL_CENTER_TOP,
             numfmt=numfmt)
        for col in ("D", "E"):
            _put(ws, f"{col}{row}", None, font=f_num, fill=fill,
                 align=AL_CENTER_TOP)
        ws.row_dimensions[row].height = 16.5
    ws.row_dimensions[12].height = 24.75

    # --- three breakdowns, side by side ---------------------------------
    band("B14:E14", "Breakdown by Module", f_section, FILL_HEADER, 18.0)
    band("G14:K14", "Breakdown by Wave", f_section, FILL_HEADER, 18.0)
    band("M14:P14", "Breakdown by Department", f_section, FILL_HEADER, 18.0)

    def head(cells):
        for coord, text, align in cells:
            _put(ws, coord, text, font=f_colhdr, fill=FILL_HEADER, align=align)

    head([("B15", "Module", AL_CENTER_TOP), ("C15", "Total", AL_CENTER_TOP),
          ("D15", "Pass", AL_CENTER_TOP), ("E15", "Fail", AL_CENTER_TOP)])
    head([("G15", "Wave", AL_CENTER_TOP), ("H15", "ชื่อระลอก", AL_CENTER_TOP),
          ("I15", "Total", AL_CENTER_TOP), ("J15", "Pass", AL_CENTER_TOP),
          ("K15", "Fail", AL_CENTER_TOP)])
    head([("M15", "Department", AL_CENTER_TOP), ("N15", "Total", AL_CENTER_TOP),
          ("O15", "Pass", AL_CENTER_TOP), ("P15", "Fail", AL_CENTER_TOP)])
    ws.row_dimensions[15].height = 15.75

    def breakdown(start_row, label_col, cols, key_range, keys, extra_label=None):
        """cols = (total, pass, fail) column letters."""
        c_total, c_pass, c_fail = cols
        row = start_row
        for key in keys:
            _put(ws, f"{label_col}{row}", key, font=f_small_lbl, fill=FILL_WHITE,
                 align=AL_LEFT_TOP)
            if extra_label:
                _put(ws, f"{extra_label}{row}", WAVE_TITLES.get(key, ""),
                     font=f_small_lbl, fill=FILL_WHITE, align=AL_LEFT_TOP)
            _put(ws, f"{c_total}{row}", f'=COUNTIF({key_range},"{key}")',
                 font=f_small_num, fill=FILL_WHITE, align=AL_CENTER_TOP)
            _put(ws, f"{c_pass}{row}",
                 f'=COUNTIFS({key_range},"{key}",{rng_status},"Pass")',
                 font=f_small_num, fill=FILL_GREEN, align=AL_CENTER_TOP)
            _put(ws, f"{c_fail}{row}",
                 f'=COUNTIFS({key_range},"{key}",{rng_status},"Fail")',
                 font=f_small_num, fill=FILL_RED, align=AL_CENTER_TOP)
            ws.row_dimensions[row].height = 15.0
            row += 1
        # total row
        _put(ws, f"{label_col}{row}", "รวม / Total", font=f_label_b,
             fill=FILL_LABEL, align=AL_LEFT_TOP)
        if extra_label:
            _put(ws, f"{extra_label}{row}", None, font=f_label_b, fill=FILL_LABEL,
                 align=AL_LEFT_TOP)
        for col in (c_total, c_pass, c_fail):
            _put(ws, f"{col}{row}", f"=SUM({col}{start_row}:{col}{row - 1})",
                 font=f_label_b, fill=FILL_LABEL, align=AL_CENTER_TOP)
        ws.row_dimensions[row].height = 15.75
        return row

    breakdown(16, "B", ("C", "D", "E"), rng_module, MODULES)
    breakdown(16, "G", ("I", "J", "K"), rng_wave, WAVE_FILES, extra_label="H")
    breakdown(16, "M", ("N", "O", "P"), rng_dept, DEPARTMENTS)

    band("B26:P26",
         "หมายเหตุ: ตัวเลขอัปเดตอัตโนมัติเมื่อกรอกคอลัมน์ Status ในแท็บ “2. Test Cases” · "
         f"ช่วงข้อมูลจริง = แถวที่ {first} ถึง {last_row} · ไฟล์นี้ไม่มีแถวตัวอย่าง (ไม่มี TC-EX) · "
         "ช่อง Total ของทั้ง 3 ตารางต้องเท่ากับ Total test cases ด้านบน · "
         "แถว Not Run คำนวณจาก “Total ลบด้วย Pass + Fail + Blocked + N/A” "
         "จึงนับช่อง Status ที่ยังว่างให้เป็น Not Run โดยอัตโนมัติ (ไม่ต้องพิมพ์คำว่า Not Run เอง) "
         "และทำให้ห้าแถวบนรวมกันได้เท่ากับ Total เสมอ",
         f_note, FILL_WHITE)
    ws.row_dimensions[26].height = 45.0


# --------------------------------------------------------------------------
# sheet 4 — Defects
# --------------------------------------------------------------------------
DEFECT_COLUMNS = [
    ("A", "Defect ID", 11.0, "center"),
    ("B", "UAT Case ID", 13.0, "center"),
    ("C", "Department", 18.0, "left"),
    ("D", "Title", 45.0, "left"),
    ("E", "Severity", 11.0, "center"),
    ("F", "Reported by", 15.0, "left"),
    ("G", "Reported date", 14.0, "center"),
    ("H", "Owner", 15.0, "left"),
    ("I", "Status", 13.0, "center"),
    ("J", "Fixed in", 16.0, "left"),
    ("K", "Retest date", 13.0, "center"),
    ("L", "Retest result", 13.0, "center"),
    ("M", "Notes", 40.0, "left"),
]
DEFECT_ROWS = 40


def build_defects_sheet(ws):
    ws.sheet_view.showGridLines = False
    _set_widths(ws, {c[0]: c[2] for c in DEFECT_COLUMNS})
    last_col = DEFECT_COLUMNS[-1][0]

    f_title = Font(name="Arial", size=13, bold=True, color=C_WHITE)
    f_header = Font(name="Arial", size=10, bold=True, color=C_WHITE)
    f_plain = Font(name="Arial", size=9, bold=False, color="FF000000")
    f_note = Font(name="Arial", size=9, bold=False, color=C_AI_PURPLE)

    ws.merge_cells(f"A1:{last_col}1")
    _put(ws, "A1",
         "Budget Web App — UAT Defect Log  ·  เปิดแถวใหม่ทุกครั้งที่มีเคส Status = Fail "
         "แล้วนำเลข Defect ID ไปใส่ในคอลัมน์ Defect ID ของแท็บ “2. Test Cases”",
         font=f_title, fill=FILL_TITLE, align=AL_LEFT_CENTER, border=BORDER_BAND)
    for col, *_ in DEFECT_COLUMNS[1:]:
        ws[f"{col}1"].fill = FILL_TITLE
        ws[f"{col}1"].border = BORDER_BAND
    ws.row_dimensions[1].height = 24.0

    for col, header, _w, _al in DEFECT_COLUMNS:
        _put(ws, f"{col}2", header, font=f_header, fill=FILL_HEADER,
             align=AL_CENTER_CENTER)
    ws.row_dimensions[2].height = 28.0

    first_row = 3
    last_row = first_row + DEFECT_ROWS - 1
    for n in range(DEFECT_ROWS):
        row = first_row + n
        for col, _hdr, _w, al in DEFECT_COLUMNS:
            value = f"D-{n + 1:03d}" if col == "A" else None
            align = AL_CENTER_TOP if al == "center" else AL_LEFT_TOP
            cell = _put(ws, f"{col}{row}", value, font=f_plain, fill=FILL_YELLOW,
                        align=align)
            if col in ("G", "K"):
                cell.number_format = "dd/mm/yyyy"
        ws.row_dimensions[row].height = 16.5

    dvs = [
        (DV_DEFECT_SEVERITY, f"E{first_row}:E{last_row}"),
        (DV_DEFECT_STATUS, f"I{first_row}:I{last_row}"),
        (DV_DEFECT_RETEST, f"L{first_row}:L{last_row}"),
    ]
    for formula, sqref in dvs:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                            showDropDown=False, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(sqref)

    ws.auto_filter.ref = f"A2:{last_col}{last_row}"
    ws.freeze_panes = "C3"

    note_row = last_row + 2
    ws.merge_cells(f"A{note_row}:{last_col}{note_row}")
    _put(ws, f"A{note_row}",
         "Severity ของ defect ใช้ **ชุดค่าเดียวกับคอลัมน์ Severity ในแท็บ “2. Test Cases”** "
         "ยกมาใส่ได้ตรง ๆ ไม่ต้องแปลงค่า: Critical = ระบบใช้งานไม่ได้ / ข้อมูลเงินผิดพลาดร้ายแรง · "
         "High = ฟังก์ชันหลักเสีย · Medium = เสียแต่มีทางเลี่ยง · Low = ปัญหาเล็กน้อยหรือ UI "
         "(ไม่มีค่า `-` ในแท็บนี้ เพราะแถว defect เกิดขึ้นเฉพาะตอนมีเคส Fail)  ·  "
         "Status: Open → In Progress → Fixed → Retested → Closed "
         "(หรือ Wont Fix ถ้าตัดสินว่าไม่แก้)",
         font=f_note, fill=FILL_WHITE, align=AL_LEFT_TOP, border=BORDER_BAND)
    for col, *_ in DEFECT_COLUMNS[1:]:
        ws[f"{col}{note_row}"].fill = FILL_WHITE
        ws[f"{col}{note_row}"].border = BORDER_BAND
    ws.row_dimensions[note_row].height = 34.0


# --------------------------------------------------------------------------
# traceability block for UAT_Test_Script.md (section 8)
# --------------------------------------------------------------------------
BEGIN_MARK = "<!-- BEGIN GENERATED TRACEABILITY (build_uat_test_cases.py) -->"
END_MARK = "<!-- END GENERATED TRACEABILITY -->"


def _source_label(case):
    sit = list(case.get("sit_refs") or [])
    new = list(case.get("new_refs") or [])
    parts = []
    if sit:
        parts.append("SIT " + ", ".join(sit))
    if new:
        parts.append("ใหม่ (" + ", ".join(new) + ")")
    if not parts:
        parts.append("ใหม่ (ไม่เคยมีใน SIT)")
    return " + ".join(parts)


def render_traceability(cases):
    lines = [
        BEGIN_MARK,
        "",
        "> ตารางนี้ถูกสร้างอัตโนมัติจากไฟล์ JSON ชุดเดียวกับที่ใช้สร้าง Excel "
        "(`_build/build_uat_test_cases.py`) ห้ามพิมพ์แก้ด้วยมือ — แก้ที่ JSON แล้วรัน generator ใหม่",
        "",
        "| UAT ID | Wave | Module | ที่มา (SIT TC / ใหม่) | เคสรวม? |",
        "|---|---|---|---|---|",
    ]
    for case in cases:
        merged = "รวมหลายเคส" if case.get("merged") else "-"
        lines.append(
            "| {id} | {wave} | {module} | {src} | {merged} |".format(
                id=case["id"], wave=case["wave"], module=case["module"],
                src=_source_label(case), merged=merged,
            )
        )
    sit_ids = sorted({r for c in cases for r in (c.get("sit_refs") or [])})
    new_ids = sorted({r for c in cases for r in (c.get("new_refs") or [])
                      if r.startswith("U-")})
    lines += [
        "",
        f"**สรุปความครอบคลุม:** {len(cases)} เคส UAT ครอบคลุมเคส SIT เดิม {len(sit_ids)} เคส "
        f"({', '.join(sit_ids)}) และประเด็นใหม่หลัง SIT อีก {len(new_ids)} ข้อ "
        f"({', '.join(new_ids)}) · ระลอก W4 (แนบเอกสาร) ทั้ง 5 เคสเป็นของใหม่ทั้งหมด "
        "เพราะฟีเจอร์แนบไฟล์ไม่เคยถูกทดสอบใน SIT เลยแม้แต่เคสเดียว",
        "",
        END_MARK,
    ]
    return "\n".join(lines)


# The companion doc quotes UAT-40's run-order position five times, in the ONE
# instruction whose whole job is to stop the tester pressing an irreversible
# Approve too early.  That number is computed here, so a 'runs_after' edit in
# the case JSON can silently move it and leave the doc telling the tester the
# wrong position.  These anchored patterns keep the prose in step: the digits
# are replaced, the surrounding Thai is not touched.
UAT40_ORDER_PATTERNS = [
    r"(`ลำดับการรัน` ที่ )\d+( จาก \d+)",
    r"(ทุกเคสที่มีลำดับการรันน้อยกว่า )\d+",
    r"(\(ลำดับการรันที่ \*\*)\d+(\*\* จาก \d+\))",
    r"(\(ลำดับการรันที่ )\d+(\))",
]

# UAT-47 and UAT-40 share one Approve click-set, and §4.1 states their
# adjacent run-order positions in a single sentence ("UAT-47 อยู่ที่ N และ
# UAT-40 อยู่ที่ M ติดกัน").  Both digits must track the computed order the
# same way UAT40_ORDER_PATTERNS above do, or a 'runs_after' edit could move
# one or both positions and leave this sentence silently wrong — the exact
# failure the pattern list exists to prevent.
UAT47_UAT40_ADJACENT_PATTERN = r"(UAT-47 อยู่ที่ )\d+( และ UAT-40 อยู่ที่ )\d+( ติดกัน)"


def _sync_uat40_order(text, uat40_order, uat47_order=None):
    """Rewrite every quoted UAT-40 (and paired UAT-47) run-order position."""
    hits = 0
    for pattern in UAT40_ORDER_PATTERNS:
        text, n = re.subn(pattern, lambda m: m.group(1) + str(uat40_order)
                          + (m.group(2) if m.lastindex and m.lastindex > 1 else ""),
                          text)
        hits += n
    if uat47_order is not None:
        text, n = re.subn(
            UAT47_UAT40_ADJACENT_PATTERN,
            lambda m: m.group(1) + str(uat47_order) + m.group(2)
            + str(uat40_order) + m.group(3),
            text,
        )
        hits += n
    return text, hits


def update_markdown(cases, uat40_order):
    if not MD_OUT.exists():
        print(f"WARN: {MD_OUT.name} not found — traceability table not written")
        return False
    with open(MD_OUT, encoding="utf-8") as fh:
        text = fh.read()
    if BEGIN_MARK not in text or END_MARK not in text:
        print(f"WARN: markers missing in {MD_OUT.name} — traceability table not written")
        return False
    head, rest = text.split(BEGIN_MARK, 1)
    _mid, tail = rest.split(END_MARK, 1)
    new_text = head + render_traceability(cases) + tail
    order_by_id = {c["id"]: c["run_order"] for c in cases}
    uat47_order = order_by_id.get("UAT-47")
    new_text, hits = _sync_uat40_order(new_text, uat40_order, uat47_order)
    if hits == 0:
        print(f"WARN: no UAT-40 run-order reference found in {MD_OUT.name} — "
              "the doc's 'do not Approve before position N' warning can no "
              "longer be kept in step with the computed order")
    if new_text != text:
        with open(MD_OUT, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(new_text)
    return hits


# --------------------------------------------------------------------------
# deterministic .xlsx
# --------------------------------------------------------------------------
PINNED_ISO = PINNED_DT.strftime("%Y-%m-%dT%H:%M:%SZ")


def _pin_core_props(data):
    """openpyxl stamps <dcterms:modified> with utcnow() at save time — pin it."""
    text = data.decode("utf-8")
    for tag in ("created", "modified"):
        text = re.sub(
            r"(<dcterms:%s[^>]*>)[^<]*(</dcterms:%s>)" % (tag, tag),
            r"\g<1>" + PINNED_ISO + r"\g<2>",
            text,
        )
    return text.encode("utf-8")


def normalise_zip(path):
    """Rewrite the .xlsx with pinned entry timestamps so re-runs are identical."""
    with zipfile.ZipFile(path, "r") as zin:
        entries = [(info, zin.read(info.filename)) for info in zin.infolist()]
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in entries:
            if info.filename == "docProps/core.xml":
                data = _pin_core_props(data)
            new_info = zipfile.ZipInfo(info.filename, date_time=PINNED_ZIP_TS)
            new_info.compress_type = zipfile.ZIP_DEFLATED
            new_info.external_attr = info.external_attr
            new_info.create_system = info.create_system
            zout.writestr(new_info, data)
    os.replace(tmp, path)


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------
def main():
    cases = load_cases()
    n = len(cases)
    last_row = 2 + n

    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "1. Info & Instructions"
    ws_cases = wb.create_sheet("2. Test Cases")
    ws_summary = wb.create_sheet("3. Summary")
    ws_defects = wb.create_sheet("4. Defects")

    order_by_id = {c["id"]: c["run_order"] for c in cases}
    uat40_order = order_by_id.get("UAT-40")
    if uat40_order is None:
        raise SystemExit("FATAL: UAT-40 missing — sheet 1 quotes its run order")

    build_info_sheet(ws_info, n, last_row, uat40_order)
    real_last = build_cases_sheet(ws_cases, cases)
    assert real_last == last_row, (real_last, last_row)
    build_summary_sheet(ws_summary, cases, last_row)
    build_defects_sheet(ws_defects)

    wb.active = 0
    props = wb.properties
    props.creator = "Data and Analytics team — Chememan PCL"
    props.lastModifiedBy = "build_uat_test_cases.py"
    props.title = "Budget Web App — UAT Test Cases (Round 1, FY2027)"
    props.subject = "User Acceptance Test"
    props.created = PINNED_DT
    props.modified = PINNED_DT

    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_OUT)
    normalise_zip(XLSX_OUT)

    md_hits = update_markdown(cases, uat40_order)

    modules_used = sorted({c["module"] for c in cases})
    print(f"WROTE {XLSX_OUT}")
    print(f"  sheets      : {wb.sheetnames}")
    print(f"  data rows   : {n}  (rows 3-{last_row} on '2. Test Cases')")
    print(f"  ids         : {cases[0]['id']} .. {cases[-1]['id']}")
    print(f"  modules used: {len(modules_used)} / {len(MODULES)} declared")
    print(f"  waves       : {sorted({c['wave'] for c in cases})}")
    run_seq = sorted(cases, key=lambda c: c["run_order"])
    # run-length encode the waves along the run order.  8 blocks would mean
    # plain Wave order works; anything more is a case that must run out of wave.
    blocks = []
    for case in run_seq:
        if blocks and blocks[-1][0] == case["wave"]:
            blocks[-1][1] += 1
        else:
            blocks.append([case["wave"], 1])
    print(f"  run order   : 1..{len(run_seq)} · UAT-40 (irreversible Approve) "
          f"at position {uat40_order} · first={run_seq[0]['id']} "
          f"last={run_seq[-1]['id']}")
    print(f"  wave blocks : {len(blocks)} (8 = plain Wave order would work) -> "
          + " ".join(f"{w}x{n}" for w, n in blocks))
    print(f"  defect rows : {DEFECT_ROWS} (D-001..D-{DEFECT_ROWS:03d})")
    print(f"  traceability: {'updated ' + MD_OUT.name if md_hits else 'NOT written'}"
          f" · UAT-40 run-order references synced in the doc: {md_hits or 0}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
