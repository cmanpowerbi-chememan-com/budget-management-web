"""Merge the AI pre-run (runner/results.json) into the UAT workbook's yellow columns.

Never hand-edit the xlsx — this script is the only writer after build_uat_test_cases.py.
Rules (jakkaritw, 2026-09-08 "Actual Result is pass by ai, if not pass, note issue on remarks"):
  * Actual Result (N) <- AI actual_result text
  * Status (O)        <- Pass / Fail / Blocked only. "Evidence only", "Partial", "Deferred"
                         stay blank (= Not Run) so a human decides; the reason goes to Remarks.
  * Test Date (S)     <- run date for every executed case (blank when Deferred)
  * Remarks (T)       <- "[AI pre-run dd/mm/yyyy] " + AI remarks
  * Tester (R), Severity (P), Defect ID (Q) untouched — human-owned.
Also adds sheet "5. AI Pre-Run" with the per-case list + the human-required list.

    python -X utf8 requirement_spec/5_uat/_build/merge_ai_results.py
"""
from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[3]
XLSX = ROOT / "requirement_spec/5_uat/Budget_UAT_Test_Cases_07.09.2026.xlsx"
RESULTS = ROOT / ".scratch/uat-prep/runner/results.json"
RUN_DATE = date(2026, 9, 9)
TAG = f"[AI pre-run {RUN_DATE:%d/%m/%Y}] "
FINAL = {"Pass", "Fail", "Blocked"}

HUMAN_REQUIRED = [
    ("UAT-05", "รอ session หมดอายุ 1 ชั่วโมงแล้วกด 'เข้าสู่ระบบใหม่' ด้วยบัญชีจริง (AI เฝ้าช่วงหมดเวลาให้แล้ว — ดู Remarks; ข้อความในกล่องเขียน '14 ชั่วโมง' ตายตัว รับไว้แล้วใน UAT-04)"),
    ("UAT-49", "เปิดอีเมล 'ได้รับการอนุมัติ' (03:22) บนมือถือ 1 ครั้ง และยืนยันว่าอยู่ Inbox ของพรทิพย์ — ข้ออื่นตรวจแล้ว"),
    ("UAT-51", "ตาคนดูภาพที่ 1920×1080 และ 1366×768 (AI ถ่ายด้วย Edge + วัด contrast ให้) แล้วติ๊ก Pass ถ้าอ่านออกชัด"),
    ("UAT-29", "ยืนยันว่าไฟล์แนบเปิด/ดาวน์โหลดได้จริงบน browser ของคุณ (AI จับการดาวน์โหลดและชื่อไฟล์ให้แล้ว)"),
]


def main() -> int:
    d = json.load(io.open(RESULTS, encoding="utf-8"))
    by_id = {c["id"]: c for c in d["cases"]}
    wb = load_workbook(XLSX)
    ws = wb["2. Test Cases"]
    written = 0
    for r in range(3, ws.max_row + 1):
        cid = ws.cell(r, 4).value
        c = by_id.get(cid)
        if not c:
            continue
        status = c["status"]
        head = status.split(" ")[0]
        ws.cell(r, 14).value = c["actual_result"]
        ws.cell(r, 15).value = head if head in FINAL and status == head else None
        if not str(status).startswith("Deferred"):
            ws.cell(r, 19).value = RUN_DATE
            ws.cell(r, 19).number_format = "dd/mm/yyyy"
        extra = "" if ws.cell(r, 15).value else f"สถานะจาก AI: {status} · "
        ws.cell(r, 20).value = TAG + extra + (c.get("remarks") or "")
        for col in (14, 20):
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
        written += 1

    # ---- sheet 5: AI pre-run summary -------------------------------------------------
    if "5. AI Pre-Run" in wb.sheetnames:
        del wb["5. AI Pre-Run"]
    s = wb.create_sheet("5. AI Pre-Run")
    bold = Font(bold=True)
    s["A1"] = f"AI pre-run บน production ผ่าน /sit/impersonate — {RUN_DATE:%d/%m/%Y} 00:54–03:26"
    s["A1"].font = Font(bold=True, size=12)
    s["A2"] = ("AI สวมสิทธิ์เป็นผู้ทดสอบตามคอลัมน์ Tester แล้วรันตาม Run Order · ผลอยู่ในช่องสีเหลืองของแท็บ 2 · "
               "Status ใส่เฉพาะ Pass/Fail/Blocked ที่ AI มั่นใจ ช่องว่าง = ให้คนตัดสินหรือทำเอง · ภาพหลักฐาน: .scratch/uat-prep/runner/shots/")
    hdr = ["Run Order", "Test Case ID", "Tester (impersonated)", "AI Status", "Remarks", "Evidence"]
    for i, h in enumerate(hdr, 1):
        cell = s.cell(4, i, h); cell.font = bold; cell.fill = PatternFill("solid", fgColor="FFDDEBF7")
    row = 5
    for c in sorted(d["cases"], key=lambda x: x["run_order"]):
        vals = [c["run_order"], c["id"], c.get("tester_used", ""), c["status"], c.get("remarks", ""), c.get("evidence", "")]
        for i, v in enumerate(vals, 1):
            s.cell(row, i, v).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    counts: dict[str, int] = {}
    for c in d["cases"]:
        k = c["status"].split(" ")[0] if c["status"].split(" ")[0] in FINAL and c["status"] in FINAL else "Human decides / not run"
        counts[k] = counts.get(k, 0) + 1
    row += 1
    s.cell(row, 1, "สรุป").font = bold; row += 1
    for k, v in counts.items():
        s.cell(row, 1, k); s.cell(row, 2, v); row += 1
    s.cell(row, 1, "ไม่ได้รันโดย AI (ไม่มีแถว)"); s.cell(row, 2, 54 - len(d["cases"])); row += 2
    s.cell(row, 1, "งานที่คนต้องทำต่อ").font = bold; row += 1
    for ids, why in HUMAN_REQUIRED:
        s.cell(row, 1, ids); s.cell(row, 2, why)
        s.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top"); row += 1
    for col, w in zip("ABCDEF", (12, 14, 34, 34, 90, 60)):
        s.column_dimensions[col].width = w
    wb.save(XLSX)
    print(f"merged {written} cases into {XLSX.name}; counts={counts}; missing={sorted(set(f'UAT-{i:02d}' for i in range(1,55)) - set(by_id))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
