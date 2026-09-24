"""Officer-review FILE = WEB = FABRIC reconcile — PRD #34's D4 (the gate)
and D15 (the one live integration-test property, "data in the file = data
on the web = data in Fabric").

Three independent sources, compared at cents:

- **FILE** — the saved xlsx BYTES, read back with `openpyxl` (never the
  in-memory rows `app.officer_workbook` built). Row-3 SUBTOTAL cells have no
  cached value (`fullCalcOnLoad`), so the FILE total is always the sum of the
  DATA rows, never a formula read.
- **WEB** — `app.officer_workbook.build_officer_workbook`'s own
  `BuildResult.web_rows` / `detail_by_key`: the literal `get_budget_grid` /
  `fetch_detail_lines` calls the web page itself would make.
- **FABRIC** — hand-written SQL in THIS module, run again from scratch. It
  never calls `app.read_model.get_budget_grid` / `merge_budget_rows` or
  `app.sap.fetch_sap_actuals` — reusing those would make this a
  self-consistency check, not an independent one (the exact trap
  `test_integration_live.py`'s own ADR-0030 test warns about). Only the SAP
  SQL TEXT constants are imported (`app.sap.SAP_ACTUALS_SQL` /
  `HIDE_DOCUMENT_SQL` / `DOC_NUMBER_PATTERN`) — the Python
  fetch/validate/pivot/merge/scope logic below is written fresh.

Any FAIL means: no publish, no mail (the caller, `jobs.officer_review`,
enforces that — this module only reports).
"""
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO

import pyodbc
from openpyxl import load_workbook

from app.budget_xlsx import FIRST_NUM_COL
from app.officer_workbook import (
    CC_COL,
    GL_COL,
    TOPIC_GROUP_NAMES,
    TOPIC_SHEETS,
    BuildResult,
    q,
)
from app.sap import DOC_NUMBER_PATTERN, HIDE_DOCUMENT_SQL, SAP_ACTUALS_SQL

CENT = Decimal("0.01")
ZERO = Decimal("0.00")
_MAX_FAILURES_LOGGED = 50


def _dec(v) -> Decimal:
    return ZERO if v is None else Decimal(str(round(float(v), 2))).quantize(CENT)


class FabricReconcileError(RuntimeError):
    """The independent Fabric-side read hit malformed data it must not
    silently paper over (mirrors `app.sap.SapActualsFetchError`'s loud
    contract, kept as a SEPARATE error type — this module's SQL/validation
    is independent of `app.sap`, D3)."""


@dataclass
class ReconcileResult:
    ok: bool
    failures: list[str] = field(default_factory=list)
    attempts: int = 1


# ---------------------------------------------------------------------------
# FILE side
# ---------------------------------------------------------------------------
@dataclass
class _Sheet1Row:
    months: tuple[Decimal, ...]
    total_year: Decimal
    board_total_year: Decimal
    sap_total_year: Decimal


@dataclass
class _TopicRow:
    cost_center: str
    gl_account: str
    total_year: Decimal
    months: tuple[Decimal, ...]


def _load_file_side(xlsx_bytes: bytes) -> tuple[dict[tuple[str, str], _Sheet1Row], dict[str, list[_TopicRow]]]:
    wb = load_workbook(BytesIO(xlsx_bytes))
    sheet1_name = wb.sheetnames[0]
    ws1 = wb[sheet1_name]

    sheet1: dict[tuple[str, str], _Sheet1Row] = {}
    for r in range(5, ws1.max_row + 1):
        cc = ws1.cell(row=r, column=4).value
        gl = ws1.cell(row=r, column=6).value
        if not cc and not gl:
            continue
        months = tuple(_dec(ws1.cell(row=r, column=FIRST_NUM_COL + i).value) for i in range(12))
        total_year = _dec(ws1.cell(row=r, column=FIRST_NUM_COL + 12).value)
        board_total = _dec(ws1.cell(row=r, column=FIRST_NUM_COL + 13).value)
        sap_total = _dec(ws1.cell(row=r, column=FIRST_NUM_COL + 14).value)
        sheet1[(str(cc), str(gl))] = _Sheet1Row(months=months, total_year=total_year, board_total_year=board_total, sap_total_year=sap_total)

    topics: dict[str, list[_TopicRow]] = {}
    for group, _swatch in TOPIC_SHEETS:
        if group not in wb.sheetnames:
            continue
        ws = wb[group]
        headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
        try:
            total_col = headers.index(next(h for h in headers if isinstance(h, str) and h.startswith("รวมปี "))) + 1
        except StopIteration:
            topics[group] = []
            continue
        rows_: list[_TopicRow] = []
        for r in range(5, ws.max_row + 1):
            cc = ws.cell(row=r, column=CC_COL).value
            if not cc:
                continue
            gl = ws.cell(row=r, column=GL_COL).value
            total = _dec(ws.cell(row=r, column=total_col).value)
            months = tuple(_dec(ws.cell(row=r, column=total_col + 1 + i).value) for i in range(12))
            rows_.append(_TopicRow(cost_center=str(cc), gl_account=str(gl), total_year=total, months=months))
        topics[group] = rows_
    return sheet1, topics


# ---------------------------------------------------------------------------
# WEB side (from the already-built `BuildResult` — no re-fetch)
# ---------------------------------------------------------------------------
def _web_sheet1(build: BuildResult) -> dict[tuple[str, str], _Sheet1Row]:
    out = {}
    for key, br in build.web_rows.items():
        months = (
            br.pending.m01, br.pending.m02, br.pending.m03, br.pending.m04,
            br.pending.m05, br.pending.m06, br.pending.m07, br.pending.m08,
            br.pending.m09, br.pending.m10, br.pending.m11, br.pending.m12,
        )
        out[key] = _Sheet1Row(
            months=tuple(_dec(v) for v in months),
            total_year=_dec(br.pending.total_year),
            board_total_year=_dec(br.board.total_year),
            sap_total_year=_dec(br.sap.total_year),
        )
    return out


def _web_topics(build: BuildResult) -> dict[str, list[_TopicRow]]:
    topics: dict[str, list[_TopicRow]] = defaultdict(list)
    for key, lines in build.detail_by_key.items():
        group = build.gl_group_by_key.get(key)
        if group not in TOPIC_GROUP_NAMES:
            continue
        for d in lines:
            months = tuple(_dec(d.get(f"m{m:02d}")) for m in range(1, 13))
            topics[group].append(_TopicRow(cost_center=key[0], gl_account=key[1], total_year=_dec(d["total_year"]), months=months))
    return dict(topics)


# ---------------------------------------------------------------------------
# FABRIC side — fully independent SQL + Python (D3)
# ---------------------------------------------------------------------------
_FABRIC_PENDING_SQL = (
    "SELECT cost_center, gl_account, m01, m02, m03, m04, m05, m06, m07, m08, m09, m10, m11, m12, "
    "total_year, template, department FROM budget.pending_budget WHERE fiscal_year = ?"
)
_FABRIC_BOARD_SQL = "SELECT cost_center, gl_account, total_year, department FROM dbo.board_budget WHERE fiscal_year = ?"
_FABRIC_CC_DEPT_SQL = """
    SELECT cost_center, department FROM (
        SELECT cost_center, department,
               ROW_NUMBER() OVER (PARTITION BY cost_center ORDER BY filler_email) AS rn
        FROM dbo.cc_filler_map
    ) ranked WHERE rn = 1
"""
_FABRIC_STATUS_SQL = "SELECT department, status FROM budget.approval_status WHERE fiscal_year = ?"
_FABRIC_GL_GROUP_SQL = "SELECT gl_code, gl_group FROM dbo.gl_group"
_FABRIC_DETAIL_SQL = (
    "SELECT cost_center, gl_account, detail_id, total_year, m01, m02, m03, m04, m05, m06, m07, m08, m09, m10, m11, m12 "
    "FROM budget.pending_budget_detail WHERE fiscal_year = ?"
)


def _fabric_hidden_doc_periods(fabric_conn: pyodbc.Connection, board_year: int) -> list[tuple[str, int]]:
    hidden: list[tuple[str, int]] = []
    for doc_raw, month in q(fabric_conn, HIDE_DOCUMENT_SQL, board_year):
        doc = str(doc_raw).strip()
        if not DOC_NUMBER_PATTERN.match(doc):
            raise FabricReconcileError(f"malformed document_number in dbo.hide_document for board_year={board_year}")
        if not isinstance(month, int) or not 1 <= month <= 12:
            raise FabricReconcileError(f"invalid month in dbo.hide_document for board_year={board_year}")
        hidden.append((doc, month))
    return hidden


def _fabric_sap_actuals(gold_conn: pyodbc.Connection, board_year: int, hidden: list[tuple[str, int]]) -> dict[tuple[str, str], list[Decimal]]:
    sql = SAP_ACTUALS_SQL
    params: tuple = (board_year,)
    if hidden:
        if len(hidden) > 1000:
            raise FabricReconcileError(f"hide list has {len(hidden)} rows — over the 1000-row safety limit")
        values_sql = ",".join("(?,?)" for _ in hidden)
        predicate = (
            "  AND NOT EXISTS (SELECT 1 FROM (VALUES "
            f"{values_sql}) AS h(doc, mo) "
            "WHERE h.doc = accounting_doc_number AND h.mo = CAST(period_month AS INT))\n"
        )
        sql = SAP_ACTUALS_SQL.replace("GROUP BY", predicate + "GROUP BY", 1)
        params = (board_year, *[v for pair in hidden for v in pair])

    sap: dict[tuple[str, str], list[Decimal]] = {}
    for cc, gl, _fy, period, amount in q(gold_conn, sql, *params):
        months = sap.setdefault((cc, gl), [ZERO] * 12)
        months[int(period) - 1] = _dec(amount)
    return sap


def _fabric_reconcile_data(
    fabric_conn: pyodbc.Connection, gold_conn: pyodbc.Connection, planning_year: int,
) -> tuple[dict[tuple[str, str], _Sheet1Row], dict[str, list[_TopicRow]], list[str]]:
    board_year = planning_year - 1
    failures: list[str] = []

    pending: dict[tuple[str, str], dict] = {}
    for cc, gl, *m, total_year, template, dept in q(fabric_conn, _FABRIC_PENDING_SQL, planning_year):
        key = (cc, gl)
        if key in pending:
            failures.append(f"FABRIC duplicate pending_budget row {key}")
            continue
        pending[key] = {"months": [_dec(v) for v in m], "total": _dec(total_year), "template": template, "department": dept}

    board: dict[tuple[str, str], dict] = {}
    for cc, gl, total_year, dept in q(fabric_conn, _FABRIC_BOARD_SQL, board_year):
        key = (cc, gl)
        if key in board:
            failures.append(f"FABRIC duplicate board_budget row {key}")
            continue
        board[key] = {"total": _dec(total_year), "department": dept}

    cc_dept = {r[0]: r[1] for r in q(fabric_conn, _FABRIC_CC_DEPT_SQL)}
    status_by_dept = {r[0]: r[1] for r in q(fabric_conn, _FABRIC_STATUS_SQL, planning_year)}
    approved_depts = frozenset(d for d, s in status_by_dept.items() if s == "APPROVED")
    master_gl = {code: group for code, group in q(fabric_conn, _FABRIC_GL_GROUP_SQL)}

    def live_dept(key: tuple[str, str]) -> str | None:
        cc = key[0]
        p, b = pending.get(key), board.get(key)
        return cc_dept.get(cc) or (p and p["department"]) or (b and b["department"])

    hidden = _fabric_hidden_doc_periods(fabric_conn, board_year)
    sap = _fabric_sap_actuals(gold_conn, board_year, hidden)
    sap_nonzero = {k for k, ms in sap.items() if any(m != ZERO for m in ms)}

    def in_scope(key: tuple[str, str]) -> bool:
        dept = live_dept(key)
        if dept in approved_depts:
            return True
        p = pending.get(key)
        return p is not None and p["template"] == "ADMIN"

    all_keys = set(pending) | set(board) | sap_nonzero
    sheet1: dict[tuple[str, str], _Sheet1Row] = {}
    for key in all_keys:
        cc, gl = key
        if gl not in master_gl:
            continue
        if not in_scope(key):
            continue
        p, b, s = pending.get(key), board.get(key), sap.get(key)
        if p is None and b is None and key not in sap_nonzero:
            continue
        sheet1[key] = _Sheet1Row(
            months=tuple(p["months"]) if p else tuple([ZERO] * 12),
            total_year=p["total"] if p else ZERO,
            board_total_year=b["total"] if b else ZERO,
            sap_total_year=sum(s, ZERO) if s else ZERO,
        )

    topics: dict[str, list[_TopicRow]] = defaultdict(list)
    for cc, gl, _detail_id, total_year, *m in q(fabric_conn, _FABRIC_DETAIL_SQL, planning_year):
        key = (cc, gl)
        if key not in sheet1:
            continue
        group = master_gl.get(gl)
        if group not in TOPIC_GROUP_NAMES:
            continue
        topics[group].append(_TopicRow(cost_center=cc, gl_account=gl, total_year=_dec(total_year), months=tuple(_dec(v) for v in m)))

    return sheet1, dict(topics), failures


# ---------------------------------------------------------------------------
# Compare
# ---------------------------------------------------------------------------
def _compare_sheet1(file_side: dict, web_side: dict, fabric_side: dict) -> list[str]:
    failures: list[str] = []
    file_keys, web_keys, fabric_keys = set(file_side), set(web_side), set(fabric_side)
    if file_keys != web_keys:
        failures.append(f"sheet-1 row-key set FILE != WEB: file-only={sorted(file_keys - web_keys)[:10]} web-only={sorted(web_keys - file_keys)[:10]}")
    if web_keys != fabric_keys:
        failures.append(f"sheet-1 row-key set WEB != FABRIC: web-only={sorted(web_keys - fabric_keys)[:10]} fabric-only={sorted(fabric_keys - web_keys)[:10]}")

    for key in sorted(file_keys & web_keys & fabric_keys):
        f, w, fb = file_side[key], web_side[key], fabric_side[key]
        if not (f.months == w.months == fb.months):
            failures.append(f"{key}: months mismatch FILE={f.months} WEB={w.months} FABRIC={fb.months}")
        if not (f.total_year == w.total_year == fb.total_year):
            failures.append(f"{key}: total_year mismatch FILE={f.total_year} WEB={w.total_year} FABRIC={fb.total_year}")
        if not (f.board_total_year == w.board_total_year == fb.board_total_year):
            failures.append(f"{key}: board_total mismatch FILE={f.board_total_year} WEB={w.board_total_year} FABRIC={fb.board_total_year}")
        if not (f.sap_total_year == w.sap_total_year == fb.sap_total_year):
            failures.append(f"{key}: sap_total mismatch FILE={f.sap_total_year} WEB={w.sap_total_year} FABRIC={fb.sap_total_year}")
        if len(failures) > _MAX_FAILURES_LOGGED:
            failures.append("... (further sheet-1 mismatches truncated)")
            break
    return failures


def _compare_topics(file_topics: dict, web_topics: dict, fabric_topics: dict, sheet1_web: dict) -> list[str]:
    failures: list[str] = []
    for group, _swatch in TOPIC_SHEETS:
        f_rows = file_topics.get(group, [])
        w_rows = web_topics.get(group, [])
        fb_rows = fabric_topics.get(group, [])
        counts = (len(f_rows), len(w_rows), len(fb_rows))
        if len(set(counts)) != 1:
            failures.append(f"topic '{group}': line count mismatch FILE={counts[0]} WEB={counts[1]} FABRIC={counts[2]}")
        sums = (
            sum((r.total_year for r in f_rows), ZERO),
            sum((r.total_year for r in w_rows), ZERO),
            sum((r.total_year for r in fb_rows), ZERO),
        )
        if len(set(sums)) != 1:
            failures.append(f"topic '{group}': SUM(รวมปี) mismatch FILE={sums[0]} WEB={sums[1]} FABRIC={sums[2]}")

        # month parity detail-vs-parent (D4: promoted to FAIL, was WARN-only
        # in the prototype — "PRD says must be equal").
        by_parent: dict[tuple[str, str], list[Decimal]] = defaultdict(lambda: [ZERO] * 12)
        for r in w_rows:
            key = (r.cost_center, r.gl_account)
            for i, v in enumerate(r.months):
                by_parent[key][i] += v
        for key, summed in by_parent.items():
            parent = sheet1_web.get(key)
            if parent is None:
                continue
            if tuple(summed) != parent.months:
                failures.append(f"topic '{group}' {key}: detail month sum {tuple(summed)} != parent sheet-1 months {parent.months}")
    return failures


def reconcile(
    xlsx_bytes: bytes, build: BuildResult, fabric_conn: pyodbc.Connection, gold_conn: pyodbc.Connection, *, planning_year: int,
) -> ReconcileResult:
    """Run the D4 three-way compare once. Returns ok=False with named
    failures on any mismatch — never raises for a business mismatch (a
    FabricReconcileError from malformed source data, e.g. a bad
    `dbo.hide_document` row, still propagates loudly — same contract as
    `app.sap`)."""
    failures: list[str] = []

    file_sheet1, file_topics = _load_file_side(xlsx_bytes)
    web_sheet1 = _web_sheet1(build)
    web_topics = _web_topics(build)
    fabric_sheet1, fabric_topics, fabric_failures = _fabric_reconcile_data(fabric_conn, gold_conn, planning_year)
    failures.extend(fabric_failures)

    failures.extend(_compare_sheet1(file_sheet1, web_sheet1, fabric_sheet1))
    failures.extend(_compare_topics(file_topics, web_topics, fabric_topics, web_sheet1))

    grand = {
        "FILE": sum((r.total_year for r in file_sheet1.values()), ZERO),
        "WEB": sum((r.total_year for r in web_sheet1.values()), ZERO),
        "FABRIC": sum((r.total_year for r in fabric_sheet1.values()), ZERO),
    }
    if len(set(grand.values())) != 1:
        failures.append(f"grand total_year mismatch: {grand}")
    grand_board = {
        "FILE": sum((r.board_total_year for r in file_sheet1.values()), ZERO),
        "WEB": sum((r.board_total_year for r in web_sheet1.values()), ZERO),
        "FABRIC": sum((r.board_total_year for r in fabric_sheet1.values()), ZERO),
    }
    if len(set(grand_board.values())) != 1:
        failures.append(f"grand board_total mismatch: {grand_board}")
    grand_sap = {
        "FILE": sum((r.sap_total_year for r in file_sheet1.values()), ZERO),
        "WEB": sum((r.sap_total_year for r in web_sheet1.values()), ZERO),
        "FABRIC": sum((r.sap_total_year for r in fabric_sheet1.values()), ZERO),
    }
    if len(set(grand_sap.values())) != 1:
        failures.append(f"grand sap_total mismatch: {grand_sap}")

    return ReconcileResult(ok=not failures, failures=failures)
