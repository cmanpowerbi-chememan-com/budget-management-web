""""ดาวน์โหลด Excel" button (issue #35) — one department's grid, as the
approved 1-sheet workbook.

Reuses `app.budget_xlsx` (PRD #34 D17's shared Sheet-1 writer) for every
layout rule so this button and the officer-review robot can never render
the sheet differently — this module never re-implements a fill colour, a
number format, or the row-3 SUBTOTAL. It only:

1. Enriches an already-fetched, already-department-filtered `BudgetRow` list
   into `SummaryRow`s (CC name/division/C-Level, GL name/group, the ฝ่าย's
   status label, and special-GL detail text) — the DB-touching half
   (`build_export_summary_rows`).
2. Paints those rows onto a fresh workbook via `write_summary_sheet`, then
   overwrites row 1 with THIS feature's own title ("ฝ่าย: <name>", issue #35
   Implementation Decisions — distinct from the officer robot's "ขอบเขต: ...
   (n ฝ่าย)") — the pure half (`build_export_workbook`, no DB access).

`filter_rows_by_department` is the server-side mirror of the frontend's
`admitRows` (`frontend/src/grid/model.ts`) — defense-in-depth alongside
`get_budget_grid`'s own `department_filter`, per the PRD's own admission
that this is normally a no-op for server data.
"""
import dataclasses
from datetime import date, datetime
from urllib.parse import quote
from zoneinfo import ZoneInfo
import re

import pyodbc
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.budget_xlsx import SummaryRow, render_detail_line, resolve_status_label, workbook_bytes, write_summary_sheet
from app.read_model import BudgetRow, fetch_cc_dims
from app.reference_data import fetch_gl_accounts
from app.special_gl import SPECIAL_GL_GROUPS
from app.subform_read import fetch_detail_lines, fetch_trips

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BANGKOK_TZ = ZoneInfo("Asia/Bangkok")

_INVALID_FILENAME_CHARS_RE = re.compile(r'[\\/:*?"<>|.\x00-\x1f]')
# The SharePoint board-budget import trigger (ADR-0021) — a downloaded file
# must never be able to collide with it if dropped back onto SharePoint
# (issue #35 story 34).
_APPROVED_BUDGET_IMPORT_RE = re.compile(r"^approved_budget_(\d{4})\.xlsx$")

_CC_NAME_SQL = """
    SELECT cost_center, description FROM (
        SELECT cost_center, description,
               ROW_NUMBER() OVER (PARTITION BY cost_center ORDER BY filler_email) AS rn
        FROM dbo.cc_filler_map
        WHERE cost_center IN ({placeholders})
    ) ranked WHERE rn = 1
"""
_DEPARTMENT_STATUS_SQL = "SELECT status FROM budget.approval_status WHERE department = ? AND fiscal_year = ?"
_IS_AUTO_CALC_SQL = "SELECT detail_id, is_auto_calc FROM budget.pending_budget_detail WHERE fiscal_year = ? AND cost_center IN ({placeholders})"


class ExportFilenameCollisionError(RuntimeError):
    """The generated export filename matched the SharePoint board-budget
    import trigger pattern — structurally unreachable (this builder's name
    always starts with `budget_FY`), kept as a defense-in-depth guard
    (issue #35 story 34)."""


# ---------------------------------------------------------------------------
# Pure helpers — filename, Content-Disposition, admitRows mirror.
# ---------------------------------------------------------------------------
def sanitize_department_for_filename(name: str) -> str:
    """Replace filesystem-unsafe characters with '_' and collapse
    whitespace, so a ฝ่าย name safely becomes one filename segment."""
    cleaned = _INVALID_FILENAME_CHARS_RE.sub("_", name.strip())
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned or "unknown"


def build_export_filename(planning_year: int, department: str, as_of: datetime) -> str:
    """`budget_FY<year>_<ฝ่าย sanitised>_<yyyymmdd_hhmm>.xlsx` (issue #35
    story 26). Raises `ExportFilenameCollisionError` if the result ever
    matched the board-budget import trigger — never actually reachable
    given the fixed `budget_FY` prefix, but tested anyway (story 34)."""
    filename = f"budget_FY{planning_year}_{sanitize_department_for_filename(department)}_{as_of:%Y%m%d_%H%M}.xlsx"
    if _APPROVED_BUDGET_IMPORT_RE.match(filename):
        raise ExportFilenameCollisionError(filename)
    return filename


def content_disposition(filename: str) -> str:
    """ASCII fallback (non-ASCII bytes replaced with '_', repeats
    collapsed) plus an RFC 5987 `filename*=UTF-8''...` carrying the real
    (possibly non-ASCII) ฝ่าย name (issue #35 story 26)."""
    ascii_fallback = re.sub(r"[^\x20-\x7e]", "_", filename)
    ascii_fallback = re.sub(r"_+", "_", ascii_fallback).strip("_") or "budget_export.xlsx"
    return f'attachment; filename="{ascii_fallback}"; filename*=UTF-8\'\'{quote(filename)}'


def filter_rows_by_department(rows: list[BudgetRow], department: str | None) -> list[BudgetRow]:
    """Server-side mirror of `frontend/src/grid/model.ts`'s `admitRows`:
    keep only rows whose resolved `department` equals the selected ฝ่าย;
    `department=None` admits every row."""
    if department is None:
        return rows
    return [r for r in rows if r.department == department]


# ---------------------------------------------------------------------------
# Enrichment (DB-touching) — BudgetRow list -> SummaryRow list.
# ---------------------------------------------------------------------------
def _fetch_cc_names(conn: pyodbc.Connection, cost_centers: list[str]) -> dict[str, str | None]:
    if not cost_centers:
        return {}
    placeholders = ", ".join(["?"] * len(cost_centers))
    cursor = conn.cursor()
    try:
        cursor.execute(_CC_NAME_SQL.format(placeholders=placeholders), *cost_centers)
        rows = cursor.fetchall()
    finally:
        cursor.close()
    return {r[0]: r[1] for r in rows}


def _fetch_department_status(conn: pyodbc.Connection, department: str, planning_year: int) -> str | None:
    cursor = conn.cursor()
    try:
        cursor.execute(_DEPARTMENT_STATUS_SQL, department, planning_year)
        row = cursor.fetchone()
    finally:
        cursor.close()
    return row[0] if row else None


def _fetch_is_auto_calc(conn: pyodbc.Connection, planning_year: int, cost_centers: list[str]) -> dict[int, bool]:
    if not cost_centers:
        return {}
    placeholders = ", ".join(["?"] * len(cost_centers))
    cursor = conn.cursor()
    try:
        cursor.execute(_IS_AUTO_CALC_SQL.format(placeholders=placeholders), planning_year, *cost_centers)
        rows = cursor.fetchall()
    finally:
        cursor.close()
    return {r[0]: bool(r[1]) for r in rows}


def _attach_detail_lines(
    fabric_conn: pyodbc.Connection, rows: list[SummaryRow], *, planning_year: int,
) -> list[SummaryRow]:
    """Special-GL parent rows only (issue #35 Implementation Decisions —
    "only for CCs that have special-GL rows"): one `fetch_detail_lines` call
    per special-GL parent key, one `fetch_trips` call per distinct
    Travelling-Expense cost_center (trip numbering k=1..n BY COST CENTER,
    matching the prototype), and one bulk `is_auto_calc` read scoped to
    those same cost centers (not selected by `fetch_detail_lines`)."""
    special_rows = [r for r in rows if r.gl_group in SPECIAL_GL_GROUPS]
    if not special_rows:
        return rows

    detail_by_key: dict[tuple[str, str], list[dict]] = {}
    travel_ccs: set[str] = set()
    for r in special_rows:
        lines = fetch_detail_lines(fabric_conn, r.cost_center, r.gl_account, planning_year)
        if lines:
            detail_by_key[(r.cost_center, r.gl_account)] = lines
        if r.gl_group == "Travelling Expense":
            travel_ccs.add(r.cost_center)

    trips_by_id: dict[int, dict] = {}
    trip_no_by_id: dict[int, int] = {}
    for cc in sorted(travel_ccs):
        for i, t in enumerate(fetch_trips(fabric_conn, cc, planning_year), start=1):
            trips_by_id[t["trip_id"]] = t
            trip_no_by_id[t["trip_id"]] = i

    is_auto_calc_by_id = _fetch_is_auto_calc(fabric_conn, planning_year, sorted({r.cost_center for r in special_rows}))

    out: list[SummaryRow] = []
    for r in rows:
        lines = detail_by_key.get((r.cost_center, r.gl_account))
        if not lines:
            out.append(r)
            continue
        ordered = sorted(lines, key=lambda d: (d.get("trip_id") is None, d.get("trip_id") or 0, d["detail_id"]))
        texts = tuple(render_detail_line(d, trips_by_id, trip_no_by_id, is_auto_calc_by_id) for d in ordered)
        out.append(dataclasses.replace(r, detail_lines=texts))
    return out


def build_export_summary_rows(
    fabric_conn: pyodbc.Connection,
    rows: list[BudgetRow],
    *,
    planning_year: int,
    department: str,
) -> list[SummaryRow]:
    """Enrich an already-fetched, already-department-filtered `BudgetRow`
    list into display-ready `SummaryRow`s. GL name/group prefer the GL
    master (`dbo.gl_group`), falling back to the row's own pending/board
    layer only when the master has nothing (gate fix round, item 4 —
    DECIDED: master-first, aligning with the grid's own `glMetaFor`, the
    approved prototype, and `app.officer_workbook`; issue #35's original
    "falling back to the GL master" wording was an authoring error); CC
    name/division/C-Level come from
    `dbo.cc_filler_map` (division/C-Level via `read_model.fetch_cc_dims`,
    name via this module's own scoped lookup — `fetch_cc_dims` does not
    carry the CC name); the ฝ่าย status is ONE lookup for the whole file
    (every row shares the same selected ฝ่าย, unlike the officer robot's
    multi-department scope)."""
    if not rows:
        return []

    cost_centers = sorted({r.cost_center for r in rows})
    cc_dims = fetch_cc_dims(fabric_conn, cost_centers)
    cc_names = _fetch_cc_names(fabric_conn, cost_centers)
    gl_master = {g["gl_code"]: g for g in fetch_gl_accounts(fabric_conn)}
    status_label = resolve_status_label(_fetch_department_status(fabric_conn, department, planning_year))

    summary_rows: list[SummaryRow] = []
    for br in rows:
        dims = cc_dims.get(br.cost_center, {})
        glmeta = gl_master.get(br.gl_account, {})
        side = "COST" if br.gl_account.startswith("5") else ("SGA" if br.gl_account.startswith("6") else "")
        summary_rows.append(SummaryRow(
            cost_center=br.cost_center,
            gl_account=br.gl_account,
            department=br.department or department,
            division=dims.get("division") or br.pending.division or br.board.division,
            c_level=dims.get("c_level") or br.pending.c_level or br.board.c_level,
            cc_name=cc_names.get(br.cost_center),
            gl_name=glmeta.get("gl_name") or br.pending.gl_name or br.board.gl_name,
            gl_group=glmeta.get("gl_group") or br.pending.gl_group or br.board.gl_group,
            side=side,
            status_label=status_label,
            months=(
                br.pending.m01, br.pending.m02, br.pending.m03, br.pending.m04,
                br.pending.m05, br.pending.m06, br.pending.m07, br.pending.m08,
                br.pending.m09, br.pending.m10, br.pending.m11, br.pending.m12,
            ),
            total_year=br.pending.total_year,
            board_total_year=br.board.total_year,
            sap_total_year=br.sap.total_year,
            remark=br.pending.remark or "",
            detail_lines=(),
        ))

    # Sort per PRD: C-Level > สายงาน (division) > ฝ่าย (department) >
    # Cost Center > COST/SGA > กลุ่ม GL > GL — same key `app.officer_workbook`
    # uses, so the two "never diverge" outputs also never sort differently.
    summary_rows.sort(
        key=lambda r: (r.c_level or "", r.division or "", r.department or "", r.cost_center, r.side, r.gl_group or "", r.gl_account)
    )

    return _attach_detail_lines(fabric_conn, summary_rows, planning_year=planning_year)


# ---------------------------------------------------------------------------
# Pure workbook builder — no DB access.
# ---------------------------------------------------------------------------
def build_export_workbook(
    rows: list[SummaryRow],
    *,
    planning_year: int,
    department: str,
    as_of: datetime,
    sap_watermark: date | None,
) -> tuple[str, bytes]:
    """Prepared rows + display context -> (file name, xlsx bytes). No DB
    access (issue #35 story 36's "deep module"). Reuses
    `app.budget_xlsx.write_summary_sheet` (PRD #34's shared layout, which
    now requires a tz-aware `as_of` and converts it to Asia/Bangkok itself)
    for every layout rule, then overwrites row 1 with this feature's own
    title ('ฝ่าย: <name>', not the officer robot's 'ขอบเขต: ... (n ฝ่าย)').
    The filename and A1 override both use the SAME Bangkok-converted
    instant `write_summary_sheet` labels row 2 with, so a caller passing a
    UTC-aware `as_of` (e.g. a UTC container) still gets one consistent
    local time across the title, the filename, and row 2.

    `department` is stripped of XML-illegal control characters (`\\x0b`,
    `\\x01`, ...) ONCE here and the cleaned value reused for `scope_label`,
    the A1 title override, and the filename (gate fix round, item 2): a raw
    control character reaching `write_summary_sheet`'s own title-row
    assignment (a plain `ws["A1"] = ...`, not routed through its
    `_write_text_cell` sanitizer) crashes with `IllegalCharacterError`
    before this function ever gets a chance to overwrite it. Per-row cells
    (the ฝ่าย TEXT COLUMN, `SummaryRow.department`) already go through that
    sanitizer inside `write_summary_sheet` and are unaffected."""
    if as_of.tzinfo is None:
        raise ValueError("build_export_workbook: as_of must be tz-aware (matches write_summary_sheet's SPEC-9 guard)")
    as_of_bkk = as_of.astimezone(BANGKOK_TZ)
    department_clean = ILLEGAL_CHARACTERS_RE.sub("", department)
    wb = Workbook()
    ws = wb.active
    write_summary_sheet(
        ws, rows, planning_year=planning_year, as_of=as_of_bkk,
        scope_label=department_clean, n_departments=1, sap_watermark=sap_watermark,
    )
    ws["A1"] = f"งบประมาณ FY{planning_year} · ข้อมูล ณ {as_of_bkk:%d/%m/%Y %H:%M} น. · ฝ่าย: {department_clean}"
    wb.calculation.fullCalcOnLoad = True  # row-3 SUBTOTALs have no cached value -> Excel computes on open
    xlsx_bytes = workbook_bytes(wb)
    filename = build_export_filename(planning_year, department_clean, as_of_bkk)
    return filename, xlsx_bytes
