"""Shared Sheet-1 "summary" xlsx writer (PRD #34 D17) — the layout the users'
"ดาวน์โหลด Excel" button and the officer-review robot (`app.officer_workbook`)
must NEVER be allowed to diverge on (issue #34 story 33). Whichever of the two
efforts lands first owns this module; the other reuses it as-is.

Layout ported from the owner-approved prototype
`playdata/excel-export-preview/build_preview.py` (2026-09-24 gate,
APPROVE-WITH-SUGGESTIONS) — that file is the layout/control SPECIFICATION,
never imported by production code (issue #34 Implementation Decisions).

This module is deliberately free of officer-only concepts (scope, PDPA
read-back, topic sheets, SharePoint publish) — it only knows how to paint ONE
worksheet from a list of already-resolved `SummaryRow`s plus a few display
parameters. Officer-only code lives in `app.officer_workbook`.
"""
import math
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.approval import APPROVED, DRAFT, PENDING_APPROVER1, PENDING_APPROVER2, PENDING_APPROVER3, REJECTED

MONTHS_TH: tuple[str, ...] = ("ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค.")

FONT_NAME = "Tahoma"
NUM_FMT = "#,##0.00"

# Sentinel status key for a `pending.template == 'ADMIN'` row in a department
# that is NOT itself APPROVED — never a real `budget.approval_status` value,
# used only as a dict key into STATUS_LABEL_TH (PRD #34, "Template 2 อนุมัติทันที").
ADMIN_TEMPLATE_STATUS_KEY = "ADMIN"

# D8: Thai status labels, WITH an explicit DRAFT key — the prototype's known
# leftover bug was a missing DRAFT key that let an unmapped status leak out in
# raw English. `resolve_status_label` below raises instead of doing that.
STATUS_LABEL_TH: dict[str, str] = {
    DRAFT: "ร่าง — ยังไม่ส่ง",
    PENDING_APPROVER1: "รออนุมัติ ขั้นที่ 1",
    PENDING_APPROVER2: "รออนุมัติ ขั้นที่ 2",
    PENDING_APPROVER3: "รออนุมัติ ขั้นที่ 3",
    APPROVED: "อนุมัติแล้ว",
    REJECTED: "ถูกตีกลับ",
    ADMIN_TEMPLATE_STATUS_KEY: "Template 2 อนุมัติทันที",
}


class UnknownStatusError(ValueError):
    """A status key has no Thai label in `STATUS_LABEL_TH` — must FAIL the
    build (D8), never silently leak the raw English key onto the sheet."""


def resolve_status_label(status_key: str | None) -> str:
    """`None` (no `budget.approval_status` row at all) means DRAFT, the same
    synthesis `app.approval.get_approval_status` performs. Any other
    unrecognised key raises `UnknownStatusError`."""
    key = status_key if status_key is not None else DRAFT
    try:
        return STATUS_LABEL_TH[key]
    except KeyError as exc:
        raise UnknownStatusError(f"no Thai label for approval status {key!r}") from exc


@dataclass(frozen=True)
class SummaryRow:
    """One Sheet-1 row: a fully resolved, display-ready `(cost_center,
    gl_account)` key. All money fields are plain floats (matching
    `app.read_model.BudgetRow`'s own layer shape) — callers quantize to
    Decimal only when reconciling, never when writing the sheet."""

    cost_center: str
    gl_account: str
    department: str | None
    division: str | None
    c_level: str | None
    cc_name: str | None
    gl_name: str | None
    gl_group: str | None
    side: str  # "COST" | "SGA" | ""
    status_label: str
    months: tuple[float, ...]  # exactly 12, m01..m12
    total_year: float
    board_total_year: float
    sap_total_year: float
    remark: str
    detail_lines: tuple[str, ...] = ()


def _sheet1_headers(planning_year: int) -> list[str]:
    board_year = planning_year - 1
    return (
        ["ฝ่าย", "สายงาน", "C-Level", "Cost Center", "ชื่อ Cost Center", "GL", "ชื่อ GL", "กลุ่ม GL", "COST/SGA", "สถานะฝ่าย"]
        + list(MONTHS_TH)
        + [f"รวมปี {planning_year}", f"งบอนุมัติ {board_year}", f"ใช้จริง SAP {board_year} (YTD)", "Remark", "รายละเอียด"]
    )


FIRST_NUM_COL = 11  # K = ม.ค.
LAST_NUM_COL = 25  # Y = ใช้จริง SAP <board_year> (YTD)
WIDTHS = [24, 22, 20, 12, 28, 12, 30, 20, 9, 12] + [11] * 12 + [13, 14, 15, 30, 70]

# Prior-year reference columns get their own colour (jakkaritw 2026-09-24):
# header = the owner's swatch, values + the row-3 total = the same colour at a
# 50% white tint. column index -> (header hex, value hex). Positions are
# fixed by the 27-column layout, not year-dependent.
COL_TINT: dict[int, tuple[str, str]] = {
    24: ("FAE7EB", "FDF3F5"),  # X งบอนุมัติ <board_year> — pink
    25: ("E0D4E7", "F0EAF3"),  # Y ใช้จริง SAP <board_year> (YTD) — lavender
}


def tint50(hexcolor: str) -> str:
    """50% white tint, per channel `ROUND_HALF_UP((c+255)/2)` — Decimal-based
    so it never truncates (`int(Decimal('252.5'))` would under-round). Used by
    the topic-sheet writer (`app.officer_workbook`) for the per-topic swatch;
    kept here so both writers share ONE tint formula."""
    out = []
    for i in (0, 2, 4):
        c = int(hexcolor[i:i + 2], 16)
        half = int((Decimal(c + 255) / Decimal(2)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        out.append(f"{half:02X}")
    return "".join(out)


def _row_height(detail_lines: tuple[str, ...], remark: str) -> float | None:
    lines = sum(max(1, math.ceil(len(t) / 68)) for t in detail_lines) if detail_lines else 1
    if remark:
        lines = max(lines, max(1, math.ceil(len(remark) / 28)))
    return 13.5 * lines if lines > 1 else None


def write_summary_sheet(
    ws: Worksheet,
    rows: list[SummaryRow],
    *,
    planning_year: int,
    as_of: datetime,
    scope_label: str,
    n_departments: int,
    sap_watermark: date | None,
) -> None:
    """Paint the Sheet-1 summary layout onto `ws` (already created by the
    caller — `wb.active` for a fresh workbook). Sets `ws.title` itself.

    `scope_label` is the plain-Thai description of what is in scope (e.g.
    "ฝ่ายที่อนุมัติแล้ว") — the `({n_departments} ฝ่าย)` suffix is added here.

    Empty scope (0 rows, D6 edge): the sheet still gets its full header +
    formatting, row 2 reads "ยังไม่มีรายการ" (matching an empty topic sheet's
    own note), and the SUBTOTAL/autofilter range still resolves to row 5
    (`last = max(first, 4 + len(rows))`, never the "backwards" `4 + 0 = 4`
    range an earlier prototype revision left untested)."""
    board_year = planning_year - 1
    headers = _sheet1_headers(planning_year)
    ws.title = f"งบประมาณ FY{planning_year}"

    base = Font(name=FONT_NAME, size=10)
    bold = Font(name=FONT_NAME, size=10, bold=True)
    head_font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="00805E")
    total_fill = PatternFill("solid", fgColor="E6F2EE")
    tint_head_font = Font(name=FONT_NAME, size=10, bold=True, color="1F1F1F")
    tint_head_fill = {c: PatternFill("solid", fgColor=h) for c, (h, _v) in COL_TINT.items()}
    tint_value_fill = {c: PatternFill("solid", fgColor=v) for c, (_h, v) in COL_TINT.items()}
    thin = Side(style="thin", color="7F7F7F")
    head_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    top_wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")

    ws["A1"] = f"งบประมาณ FY{planning_year} · ข้อมูล ณ {as_of:%d/%m/%Y %H:%M} น. · ขอบเขต: {scope_label} ({n_departments} ฝ่าย)"
    ws["A1"].font = bold
    if rows:
        wm_text = sap_watermark.strftime("%d/%m/%Y") if sap_watermark else "ไม่ทราบ (อ่าน watermark ไม่ได้)"
        ws["A2"] = (
            "ยอดเงินรวมจากคอลัมน์ตัวเลขเท่านั้น · คอลัมน์ 'รายละเอียด' เป็นข้อความอธิบาย ไม่ต้องนำไปบวก · "
            f"SAP {board_year} = ใช้จริงสะสมถึง {wm_text}"
        )
    else:
        ws["A2"] = "ยังไม่มีรายการ"
    ws["A2"].font = base

    first = 5
    last = max(first, 4 + len(rows))
    label = ws.cell(row=3, column=FIRST_NUM_COL - 1, value="รวม (ตามตัวกรอง)")
    label.font = bold
    label.alignment = Alignment(horizontal="right")
    label.fill = total_fill
    for c in range(FIRST_NUM_COL, LAST_NUM_COL + 1):
        col = get_column_letter(c)
        cell = ws.cell(row=3, column=c, value=f"=SUBTOTAL(9,{col}{first}:{col}{last})")
        cell.font = bold
        cell.number_format = NUM_FMT
        cell.fill = tint_value_fill.get(c, total_fill)

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = tint_head_font if c in COL_TINT else head_font
        cell.fill = tint_head_fill.get(c, head_fill)
        cell.border = head_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 30

    for i, r in enumerate(rows):
        rn = first + i
        text_vals = [
            r.department, r.division, r.c_level, r.cost_center, r.cc_name,
            r.gl_account, r.gl_name, r.gl_group, r.side, r.status_label,
        ]
        for c, v in enumerate(text_vals, start=1):
            cell = ws.cell(row=rn, column=c, value=v if v is not None else "")
            cell.font = base
            cell.alignment = top
            if c in (4, 6):  # Cost Center / GL stored as text
                cell.number_format = "@"
        nums = list(r.months) + [r.total_year, r.board_total_year, r.sap_total_year]
        for j, v in enumerate(nums):
            cell = ws.cell(row=rn, column=FIRST_NUM_COL + j, value=float(v))
            cell.font = base
            cell.number_format = NUM_FMT
            cell.alignment = top
            if FIRST_NUM_COL + j in tint_value_fill:
                cell.fill = tint_value_fill[FIRST_NUM_COL + j]
        remark_cell = ws.cell(row=rn, column=LAST_NUM_COL + 1, value=r.remark)
        remark_cell.font = base
        remark_cell.alignment = top_wrap
        detail_cell = ws.cell(row=rn, column=LAST_NUM_COL + 2, value="\n".join(r.detail_lines))
        detail_cell.font = base
        detail_cell.alignment = top_wrap
        height = _row_height(r.detail_lines, r.remark)
        if height:
            ws.row_dimensions[rn].height = height

    for c, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = f"{get_column_letter(FIRST_NUM_COL)}{first}"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"


def workbook_bytes(wb: Workbook) -> bytes:
    """Serialize an in-memory workbook to bytes without touching disk."""
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
