"""Seam B — the pure "ดาวน์โหลด Excel" workbook builder (issue #35). No DB:
feeds in-memory `SummaryRow`s + display context straight into
`build_export_workbook` / the small pure helpers, then opens the bytes with
openpyxl and asserts (prior art: the pure merge tests of the grid read
model)."""
import re
from datetime import date, datetime, timezone
from io import BytesIO
from unittest.mock import MagicMock, patch
from urllib.parse import quote
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
import pytest

from app.budget_export import (
    build_export_filename,
    build_export_summary_rows,
    build_export_workbook,
    content_disposition,
    filter_rows_by_department,
    sanitize_department_for_filename,
)
from app.budget_xlsx import COL_TINT, SummaryRow, render_detail_line
from app.read_model import BoardLayer, BudgetRow, PendingLayer, SapLayer

HEADERS = (
    ["ฝ่าย", "สายงาน", "C-Level", "Cost Center", "ชื่อ Cost Center", "GL", "ชื่อ GL", "กลุ่ม GL", "COST/SGA", "สถานะฝ่าย"]
    + ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]
    + ["รวมปี 2027", "งบอนุมัติ 2026", "ใช้จริง SAP 2026 (YTD)", "Remark", "รายละเอียด"]
)
FIRST_NUM_COL = 11  # K = ม.ค.
LAST_NUM_COL = 25  # Y = ใช้จริง SAP 2026 (YTD)
BANGKOK = ZoneInfo("Asia/Bangkok")
AS_OF = datetime(2026, 9, 24, 9, 24, tzinfo=BANGKOK)


def _row(**overrides) -> SummaryRow:
    base = dict(
        cost_center="10CS010000",
        gl_account="5211800030",
        department="Corporate Strategy 2",
        division="Strategy Division",
        c_level="CEO",
        cc_name="Strategy CC",
        gl_name="Office expenses",
        gl_group="Office",
        side="COST",
        status_label="อนุมัติแล้ว",
        months=(100.0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        total_year=100.0,
        board_total_year=50.0,
        sap_total_year=25.0,
        remark="",
        detail_lines=(),
    )
    base.update(overrides)
    return SummaryRow(**base)


# ---------------------------------------------------------------------------
# build_export_workbook — layout (headers, row 3, freeze/autofilter, tints,
# CC/GL-as-text, number formats, the feature's own title row).
# ---------------------------------------------------------------------------
def test_title_row_states_year_asof_and_department():
    filename, xlsx_bytes = build_export_workbook(
        [_row()], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=date(2026, 9, 21),
    )
    wb = load_workbook(BytesIO(xlsx_bytes))
    ws = wb.active
    assert ws["A1"].value == "งบประมาณ FY2027 · ข้อมูล ณ 24/09/2026 09:24 น. · ฝ่าย: Corporate Strategy 2"
    assert ws["A1"].font.bold is True
    assert filename.startswith("budget_FY2027_Corporate_Strategy_2_")


def test_naive_as_of_is_rejected_loudly():
    with pytest.raises(ValueError, match="tz-aware"):
        build_export_workbook(
            [_row()], planning_year=2027, department="Corporate Strategy 2",
            as_of=datetime(2026, 9, 24, 9, 24), sap_watermark=None,
        )


def test_utc_as_of_is_shown_and_filed_as_bangkok_local_time():
    utc_as_of = datetime(2026, 9, 24, 2, 24, tzinfo=timezone.utc)  # 09:24 ICT
    filename, xlsx_bytes = build_export_workbook(
        [_row()], planning_year=2027, department="Corporate Strategy 2", as_of=utc_as_of, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    assert ws["A1"].value == "งบประมาณ FY2027 · ข้อมูล ณ 24/09/2026 09:24 น. · ฝ่าย: Corporate Strategy 2"
    assert "20260924_0924" in filename


def test_header_row_matches_the_approved_column_order():
    _filename, xlsx_bytes = build_export_workbook(
        [_row()], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    hdr = [ws.cell(row=4, column=c).value for c in range(1, len(HEADERS) + 1)]
    assert hdr == HEADERS


def test_row3_subtotal_covers_exact_data_range_and_is_the_only_formula():
    rows = [_row(), _row(cost_center="10CS010001")]
    _filename, xlsx_bytes = build_export_workbook(
        rows, planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    first, last = 5, 4 + len(rows)
    assert ws.max_row == last  # no bottom total row
    other_formulas = 0
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                if cell.row == 3 and FIRST_NUM_COL <= cell.column <= LAST_NUM_COL:
                    from openpyxl.utils import get_column_letter
                    col = get_column_letter(cell.column)
                    assert v == f"=SUBTOTAL(9,{col}{first}:{col}{last})"
                else:
                    other_formulas += 1
    assert other_formulas == 0
    n_sub = sum(1 for c in range(FIRST_NUM_COL, LAST_NUM_COL + 1) if str(ws.cell(row=3, column=c).value).startswith("=SUBTOTAL(9,"))
    assert n_sub == LAST_NUM_COL - FIRST_NUM_COL + 1


def test_freeze_panes_and_autofilter():
    from openpyxl.utils import get_column_letter

    rows = [_row(), _row(cost_center="10CS010001")]
    _filename, xlsx_bytes = build_export_workbook(
        rows, planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    assert ws.freeze_panes == "K5"
    last_col = get_column_letter(len(HEADERS))
    assert ws.auto_filter.ref == f"A4:{last_col}6"  # 2 data rows -> last = 4+2


def test_cost_center_and_gl_are_stored_as_text():
    _filename, xlsx_bytes = build_export_workbook(
        [_row(cost_center="0000012345", gl_account="0005211800")],
        planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    cc_cell, gl_cell = ws.cell(row=5, column=4), ws.cell(row=5, column=6)
    assert cc_cell.value == "0000012345" and cc_cell.number_format == "@"
    assert gl_cell.value == "0005211800" and gl_cell.number_format == "@"


def test_number_format_is_thai_thousands_two_decimals():
    _filename, xlsx_bytes = build_export_workbook(
        [_row()], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    assert ws.cell(row=5, column=FIRST_NUM_COL).number_format == "#,##0.00"
    assert ws.cell(row=5, column=FIRST_NUM_COL + 12).number_format == "#,##0.00"  # รวมปี


def test_prior_year_columns_carry_their_own_tint():
    _filename, xlsx_bytes = build_export_workbook(
        [_row()], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    for col, (head_hex, value_hex) in COL_TINT.items():
        assert ws.cell(row=4, column=col).fill.fgColor.rgb.upper().endswith(head_hex)
        assert ws.cell(row=5, column=col).fill.fgColor.rgb.upper().endswith(value_hex)
        assert ws.cell(row=3, column=col).fill.fgColor.rgb.upper().endswith(value_hex)


def test_font_is_tahoma_10_throughout():
    _filename, xlsx_bytes = build_export_workbook(
        [_row()], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    assert ws.cell(row=5, column=1).font.name == "Tahoma"
    assert ws.cell(row=5, column=1).font.size == 10


def test_empty_scope_yields_header_only_workbook():
    _filename, xlsx_bytes = build_export_workbook(
        [], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    assert ws.max_row == 4  # header only, no leaked rows
    assert ws["A2"].value == "ยังไม่มีรายการ"


# ---------------------------------------------------------------------------
# Formula-injection guard (issue #35 story 33).
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("prefix", ["=SUM(A1:A9)", "+1+1", "-1-1", "@SUM(A1)", "\tcmd", "\rcmd"])
def test_free_text_starting_with_injection_prefixes_stays_plain_text(prefix):
    row = _row(remark=prefix, detail_lines=(prefix,))
    _filename, xlsx_bytes = build_export_workbook(
        [row], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    remark_cell = ws.cell(row=5, column=LAST_NUM_COL + 1)
    detail_cell = ws.cell(row=5, column=LAST_NUM_COL + 2)
    # A lone '\r' round-trips through the saved XML as '\n' on some XML
    # parser backends (observed on the Linux CI runner's lxml build, not on
    # Windows) — XML 1.0's own end-of-line normalization (spec section
    # 2.11), not an injection failure. Normalize CR->LF before comparing so
    # this stays platform-independent; every OTHER prefix is compared
    # byte-for-byte, unnormalized, and the security assertions below
    # (still a plain string, never a formula) apply to every case as-is.
    normalize = lambda s: s.replace("\r", "\n")  # noqa: E731
    assert normalize(remark_cell.value) == normalize(prefix)
    assert normalize(detail_cell.value) == normalize(prefix)
    assert remark_cell.data_type == "s"
    assert detail_cell.data_type == "s"


def test_only_formulas_in_the_file_are_the_row3_subtotals():
    row = _row(remark="=cmd|'/C calc'!A0", detail_lines=("=HYPERLINK(\"http://evil\")",))
    _filename, xlsx_bytes = build_export_workbook(
        [row], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    formula_cells = [
        (cell.row, cell.column)
        for row_cells in ws.iter_rows()
        for cell in row_cells
        if isinstance(cell.value, str) and cell.value.startswith("=") and cell.data_type == "f"
    ]
    assert formula_cells and all(r == 3 and FIRST_NUM_COL <= c <= LAST_NUM_COL for r, c in formula_cells)


# ---------------------------------------------------------------------------
# render_detail_line — Entertainment / trip / per-diem auto-calc / generic.
# ---------------------------------------------------------------------------
def test_render_detail_line_entertainment():
    d = {"detail_id": 1, "trip_id": None, "gl_group": "Entertainment",
         "meta_json": '{"ประเภทการรับรอง": "Customer", "รายละเอียด": "เลี้ยงลูกค้า"}', "total_year": 5000.0}
    text = render_detail_line(d, {}, {}, {})
    assert text == "Customer · เลี้ยงลูกค้า · 5,000 บาท"


def test_render_detail_line_trip_with_per_diem_auto_calc_marker():
    trips_by_id = {7: {"destination": "Chiang Mai", "days": 3, "travel_months": ["1", "2"], "project": "Site visit"}}
    d = {"detail_id": 2, "trip_id": 7, "gl_group": "Travelling Expense", "meta_json": None, "total_year": 3000.0}
    text = render_detail_line(d, trips_by_id, {7: 1}, {2: True})
    assert text == "ทริป 1: Chiang Mai · 3 วัน · เดือน Jan, Feb · Site visit · 3,000 บาท (เบี้ยเลี้ยงคำนวณอัตโนมัติ)"


def test_render_detail_line_trip_without_auto_calc_has_no_marker():
    trips_by_id = {7: {"destination": "Chiang Mai", "days": 3, "travel_months": [], "project": ""}}
    d = {"detail_id": 3, "trip_id": 7, "gl_group": "Travelling Expense", "meta_json": None, "total_year": 1200.0}
    text = render_detail_line(d, trips_by_id, {7: 1}, {3: False})
    assert "เบี้ยเลี้ยงคำนวณอัตโนมัติ" not in text
    assert text == "ทริป 1: Chiang Mai · 3 วัน · 1,200 บาท"


def test_render_detail_line_generic_meta_excludes_pii_keys():
    d = {"detail_id": 4, "trip_id": None, "gl_group": "Professional & Legal Fee",
         "meta_json": '{"Project": "Audit", "traveler_email": "someone@chememan.com"}', "total_year": 900.0}
    text = render_detail_line(d, {}, {}, {})
    assert "someone@chememan.com" not in text
    assert text == "Project: Audit · 900 บาท"


# ---------------------------------------------------------------------------
# filter_rows_by_department — server-side admitRows mirror.
# ---------------------------------------------------------------------------
def _budget_row(cc: str, dept: str | None) -> BudgetRow:
    return BudgetRow(cost_center=cc, gl_account="5211800030", department=dept)


def test_filter_rows_by_department_keeps_only_matching_department():
    rows = [_budget_row("CC1", "Corporate Strategy 2"), _budget_row("CC2", "Other Dept")]
    result = filter_rows_by_department(rows, "Corporate Strategy 2")
    assert [r.cost_center for r in result] == ["CC1"]


def test_filter_rows_by_department_none_admits_all():
    rows = [_budget_row("CC1", "Corporate Strategy 2"), _budget_row("CC2", "Other Dept")]
    assert filter_rows_by_department(rows, None) == rows


# ---------------------------------------------------------------------------
# Filename / Content-Disposition.
# ---------------------------------------------------------------------------
def test_filename_matches_the_approved_pattern():
    filename = build_export_filename(2027, "Corporate Strategy 2", AS_OF)
    assert filename == "budget_FY2027_Corporate_Strategy_2_20260924_0924.xlsx"


@pytest.mark.parametrize(
    "department", ["Corporate Strategy 2", "ฝ่ายบัญชี", "approved_budget_2026", "A/B:C*D?E", "approved_budget_2026.xlsx"]
)
def test_filename_never_matches_the_board_budget_import_trigger(department):
    filename = build_export_filename(2027, department, AS_OF)
    assert not re.match(r"^approved_budget_(\d{4})\.xlsx$", filename)


def test_sanitize_department_replaces_filesystem_unsafe_characters():
    assert sanitize_department_for_filename('A/B\\C:D*E?F"G<H>I|J') == "A_B_C_D_E_F_G_H_I_J"


# Item 8 (gate fix round): '.' was NOT in the sanitizer's unsafe-char set, so
# a ฝ่าย name containing one (or literally ending in a fake ".xlsx") left a
# stray dot embedded inside the filename SEGMENT — defense-in-depth, even
# though the fixed `budget_FY<year>_..._<timestamp>.xlsx` prefix/suffix
# already keeps the whole name from ever equaling the import trigger exactly.
def test_sanitize_department_also_replaces_dots():
    assert sanitize_department_for_filename("A.B.C") == "A_B_C"


# Item E (gate fix round 2): \x7f (DEL) plus Unicode Cf (format — e.g.
# U+202E RIGHT-TO-LEFT OVERRIDE, a filename-spoofing trick that can make
# "cmd.exe" read as "exe.dmc" in a file listing) and Cc (control, incl. the
# C1 block \x80-\x9f the plain regex does not reach) categories are dropped.
def test_sanitize_department_replaces_del_and_drops_format_and_control_categories():
    assert sanitize_department_for_filename("A\x7fB") == "A_B"
    assert sanitize_department_for_filename("A\u202eB") == "A_B"  # Cf: right-to-left override
    assert sanitize_department_for_filename("A\x80B") == "A_B"  # Cc: C1 control block


def test_content_disposition_carries_ascii_fallback_and_rfc5987_utf8_name():
    header = content_disposition("budget_FY2027_ฝ่ายบัญชี_20260924_0924.xlsx")
    assert 'filename="budget_FY2027_' in header
    assert ".xlsx" in header.split(";")[1]
    assert "filename*=UTF-8''" in header
    assert "%E0%B8%9D" in header  # percent-encoded ฝ (Thai) somewhere in the UTF-8 segment


# Item 6 (gate fix round): pin the EXACT percent-encoded segment, not just a
# stray substring of it, so a future encoding regression (e.g. swapping
# `quote` for something that encodes differently) fails loudly here.
def test_content_disposition_filename_star_is_exactly_percent_encoded():
    filename = "budget_FY2027_ฝ่ายA_20260924_0924.xlsx"
    header = content_disposition(filename)
    assert f"filename*=UTF-8''{quote(filename)}" in header


# ---------------------------------------------------------------------------
# Item 2 (gate fix round, LOW): a ฝ่าย name containing an XML-illegal control
# character (\x0b, \x01, ...) used to crash openpyxl with a raw
# IllegalCharacterError before `_write_text_cell`'s SEC-F4 stripping (added
# in `app.budget_xlsx`, which this module must NOT edit — PRD #34 owns it
# now, commit c8b763f) ever got a chance to run: `write_summary_sheet`'s own
# title-row assignment (`ws["A1"] = ...scope_label...`) is a PLAIN cell
# assignment, not routed through `_write_text_cell`. Fixed entirely inside
# `build_export_workbook` by stripping once and reusing the cleaned value for
# scope_label / the A1 override / the filename.
# ---------------------------------------------------------------------------
def test_department_with_illegal_xml_chars_does_not_crash_and_is_stripped_everywhere():
    filename, xlsx_bytes = build_export_workbook(
        [], planning_year=2027, department="A\x0bB\x01C", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    assert ws["A1"].value == "งบประมาณ FY2027 · ข้อมูล ณ 24/09/2026 09:24 น. · ฝ่าย: ABC"
    assert "\x0b" not in filename and "\x01" not in filename
    assert "ABC" in filename


# Items D + H (gate fix round 2): the FULL C0 control-char range (minus
# \t\n\r, which stay legal/usable text) PLUS the two Unicode noncharacters
# and a lone surrogate — `department_clean` now delegates to
# `app.budget_xlsx.XLSX_ILLEGAL_TEXT_RE` (item J), so every one of these
# must build without raising and leave the character out of both the title
# and the filename.
_C0_MINUS_TAB_NEWLINE_CR = [chr(c) for c in range(0x00, 0x20) if chr(c) not in ("\t", "\n", "\r")]


@pytest.mark.parametrize("illegal_char", _C0_MINUS_TAB_NEWLINE_CR + ["\ufffe", "\uffff", "\ud800"])
def test_department_with_any_illegal_char_builds_without_crashing(illegal_char):
    department = f"A{illegal_char}B"
    filename, xlsx_bytes = build_export_workbook(
        [], planning_year=2027, department=department, as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    assert illegal_char not in ws["A1"].value
    assert "AB" in ws["A1"].value
    assert illegal_char not in filename


# ---------------------------------------------------------------------------
# Item 7 (gate fix round): formula-injection guard for columns A..J — the
# system-sourced text columns (ฝ่าย/สายงาน/C-Level/ชื่อ Cost Center/ชื่อ GL/
# กลุ่ม GL/COST-SGA/สถานะฝ่าย), not just Remark/รายละเอียด. Already covered by
# `app.budget_xlsx`'s shared `_write_text_cell` applied uniformly to every
# `text_vals` column — this locks that in from OUR call site so a future
# change to either module can't silently regress it.
# ---------------------------------------------------------------------------
def test_formula_injection_guard_covers_columns_a_through_j():
    row = _row(department='=HYPERLINK("http://evil")', cc_name="=cmd|' /C calc'!A0", gl_name="=SUM(A1:A9)")
    _filename, xlsx_bytes = build_export_workbook(
        [row], planning_year=2027, department="Corporate Strategy 2", as_of=AS_OF, sap_watermark=None,
    )
    ws = load_workbook(BytesIO(xlsx_bytes)).active
    dept_cell, cc_name_cell, gl_name_cell = ws.cell(row=5, column=1), ws.cell(row=5, column=5), ws.cell(row=5, column=7)
    assert dept_cell.value == '=HYPERLINK("http://evil")' and dept_cell.data_type == "s"
    assert cc_name_cell.value == "=cmd|' /C calc'!A0" and cc_name_cell.data_type == "s"
    assert gl_name_cell.value == "=SUM(A1:A9)" and gl_name_cell.data_type == "s"


# ---------------------------------------------------------------------------
# Item 4 (gate fix round, DECIDED): GL name/group precedence is MASTER-FIRST
# (`dbo.gl_group` / `glmeta`), falling back to the row layers only when the
# master has nothing — aligns with the grid (`model.ts` `glMetaFor`), the
# approved prototype, and `app.officer_workbook`. Issue #35's original
# wording ("falling back to the GL master") was an authoring error.
# `build_export_summary_rows` touches the DB via its own leaf reads
# (`fetch_cc_dims`, `_fetch_cc_names`, `fetch_gl_accounts`,
# `_fetch_department_status`) — patched here rather than given a live
# connection, `fabric_conn` itself is a bare MagicMock never asked to do
# anything (special-GL detection short-circuits before touching it, since
# this GL is not in SPECIAL_GL_GROUPS).
# ---------------------------------------------------------------------------
def test_gl_name_and_group_prefer_the_master_over_the_row_layers():
    row = BudgetRow(
        cost_center="10CS010000",
        gl_account="5211800030",
        department="Corporate Strategy 2",
        pending=PendingLayer(gl_name="Layer Name", gl_group="Layer Group", total_year=100.0),
        board=BoardLayer(gl_name="Board Layer Name", gl_group="Board Layer Group", total_year=50.0),
        sap=SapLayer(total_year=25.0),
    )
    with patch("app.budget_export.fetch_cc_dims", return_value={}), patch(
        "app.budget_export._fetch_cc_names", return_value={}
    ), patch("app.budget_export._fetch_department_status", return_value="APPROVED"), patch(
        "app.budget_export.fetch_gl_accounts",
        return_value=[{"gl_code": "5211800030", "gl_name": "Master Name", "gl_group": "Master Group"}],
    ):
        result = build_export_summary_rows(MagicMock(), [row], planning_year=2027, department="Corporate Strategy 2")

    assert result[0].gl_name == "Master Name"
    assert result[0].gl_group == "Master Group"


# ---------------------------------------------------------------------------
# Item B (gate fix round 2, MED): a fixture that can actually distinguish
# the PRD sort key (c_level, division, department, cost_center, side,
# gl_group, gl_account) from sorting by gl_account alone / dropping
# gl_group / gl_account-before-gl_group, and from sorting cost_center ahead
# of c_level/division — `test_budget.py`'s own parity-test fixture cannot:
# its two same-CC GLs happen to agree on BOTH gl_group order AND gl_account
# order, so it would pass even with a wrong key.
# ---------------------------------------------------------------------------
def test_build_export_summary_rows_sort_key_isolates_every_level():
    """Item 5 (gate fix round 3): the round-2 fixture let c_level and
    division AGREE with each other (both pointed the same CC first), and
    let side and gl_group AGREE too (both same-CC rows were COST) — so a
    mutant dropping/swapping either pair could still pass by accident. This
    fixture makes every adjacent pair in the key DISAGREE, killing:
    drop c_level only, swap c_level/division, drop side, side placed after
    gl_group, and gl_name substituted for gl_group.

    CC "ZZ_CC": c_level "A_Level" (sorts FIRST) but division "Z_Div"
    (sorts LAST) — must still sort FIRST overall (c_level outranks
    division). CC "AA_CC": the opposite pairing (c_level "Z_Level" /
    division "A_Div") — must sort SECOND despite its division sorting
    first. Within "AA_CC", 3 rows whose side/gl_group/gl_name all disagree
    with each other: P (COST, group "Alpha", name "Zulu Name"), Q (COST,
    group "Zulu", name "Alpha Name"), R (SGA, group "Beta", name "Beta
    Name") — correct order is P, Q, R (side groups COST before SGA; within
    COST, gl_group "Alpha" < "Zulu"); every named mutant produces some
    OTHER order.
    """
    zz_row = BudgetRow(
        cost_center="ZZ_CC", gl_account="5299999999", department="Dept",
        pending=PendingLayer(total_year=1.0), board=BoardLayer(), sap=SapLayer(),
    )
    p_row = BudgetRow(  # COST, group "Alpha", name "Zulu Name"
        cost_center="AA_CC", gl_account="5215000010", department="Dept",
        pending=PendingLayer(total_year=2.0), board=BoardLayer(), sap=SapLayer(),
    )
    q_row = BudgetRow(  # COST, group "Zulu", name "Alpha Name"
        cost_center="AA_CC", gl_account="5211800030", department="Dept",
        pending=PendingLayer(total_year=3.0), board=BoardLayer(), sap=SapLayer(),
    )
    r_row = BudgetRow(  # SGA, group "Beta", name "Beta Name"
        cost_center="AA_CC", gl_account="6100000000", department="Dept",
        pending=PendingLayer(total_year=4.0), board=BoardLayer(), sap=SapLayer(),
    )

    cc_dims = {
        "ZZ_CC": {"department": "Dept", "division": "Z_Div", "c_level": "A_Level"},
        "AA_CC": {"department": "Dept", "division": "A_Div", "c_level": "Z_Level"},
    }
    gl_master = [
        {"gl_code": "5299999999", "gl_name": "ZZ Name", "gl_group": "ZZGroup"},
        {"gl_code": "5215000010", "gl_name": "Zulu Name", "gl_group": "Alpha"},
        {"gl_code": "5211800030", "gl_name": "Alpha Name", "gl_group": "Zulu"},
        {"gl_code": "6100000000", "gl_name": "Beta Name", "gl_group": "Beta"},
    ]

    with patch("app.budget_export.fetch_cc_dims", return_value=cc_dims), patch(
        "app.budget_export._fetch_cc_names", return_value={}
    ), patch("app.budget_export._fetch_department_status", return_value="APPROVED"), patch(
        "app.budget_export.fetch_gl_accounts", return_value=gl_master
    ):
        result = build_export_summary_rows(
            MagicMock(), [zz_row, p_row, q_row, r_row], planning_year=2027, department="Dept",
        )

    order = [(r.cost_center, r.gl_account) for r in result]
    assert order == [
        ("ZZ_CC", "5299999999"),
        ("AA_CC", "5215000010"),  # P
        ("AA_CC", "5211800030"),  # Q
        ("AA_CC", "6100000000"),  # R
    ]


# ---------------------------------------------------------------------------
# Item I (gate fix round 2): special-GL detection (`_attach_detail_lines`)
# must key off the SAME master-first `gl_group` `build_export_summary_rows`
# resolves (item 4), never the row's own pending layer — in EITHER
# direction (master says special but the layer disagrees, and vice versa).
# ---------------------------------------------------------------------------
def test_special_gl_detection_follows_the_master_gl_group_both_ways():
    # Master says "Travelling Expense", the row's own layer disagrees
    # ("Office") -> the master wins -> detail lines ARE fetched.
    travel_row = BudgetRow(
        cost_center="CC1", gl_account="5000000001", department="Dept",
        pending=PendingLayer(gl_group="Office", total_year=1.0), board=BoardLayer(), sap=SapLayer(),
    )
    with patch("app.budget_export.fetch_cc_dims", return_value={}), patch(
        "app.budget_export._fetch_cc_names", return_value={}
    ), patch("app.budget_export._fetch_department_status", return_value="APPROVED"), patch(
        "app.budget_export.fetch_gl_accounts",
        return_value=[{"gl_code": "5000000001", "gl_name": "X", "gl_group": "Travelling Expense"}],
    ), patch("app.budget_export.fetch_detail_lines", return_value=[]) as mock_detail, patch(
        "app.budget_export.fetch_trips", return_value=[]
    ) as mock_trips, patch("app.budget_export._fetch_is_auto_calc", return_value={}):
        build_export_summary_rows(MagicMock(), [travel_row], planning_year=2027, department="Dept")

    assert mock_detail.called
    assert mock_trips.called

    # Reverse: master says "Office", the row's own layer disagrees
    # ("Travelling Expense") -> the master wins -> NOT special, detail lines
    # are NOT fetched.
    office_row = BudgetRow(
        cost_center="CC1", gl_account="5000000002", department="Dept",
        pending=PendingLayer(gl_group="Travelling Expense", total_year=1.0), board=BoardLayer(), sap=SapLayer(),
    )
    with patch("app.budget_export.fetch_cc_dims", return_value={}), patch(
        "app.budget_export._fetch_cc_names", return_value={}
    ), patch("app.budget_export._fetch_department_status", return_value="APPROVED"), patch(
        "app.budget_export.fetch_gl_accounts",
        return_value=[{"gl_code": "5000000002", "gl_name": "X", "gl_group": "Office"}],
    ), patch("app.budget_export.fetch_detail_lines", return_value=[]) as mock_detail2, patch(
        "app.budget_export.fetch_trips", return_value=[]
    ) as mock_trips2, patch("app.budget_export._fetch_is_auto_calc", return_value={}) as mock_auto2:
        build_export_summary_rows(MagicMock(), [office_row], planning_year=2027, department="Dept")

    assert not mock_detail2.called
    assert not mock_trips2.called
    assert not mock_auto2.called
