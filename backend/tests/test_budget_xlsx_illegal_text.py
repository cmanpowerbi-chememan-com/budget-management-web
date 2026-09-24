"""Self-contained unit tests for `app.budget_xlsx.XLSX_ILLEGAL_TEXT_RE` and
its use inside `_write_text_cell`/`write_summary_sheet` (gate fix round 2,
item J). Depends ONLY on `app.budget_xlsx` — no `app.budget_export` import —
so this file stays valid on its own if `budget_xlsx.py` is committed
separately from the rest of issue #35's work.

No real names or money: every fixture value is a generic placeholder
(the repo is public).
"""
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
import pytest

from app.budget_xlsx import SummaryRow, XLSX_ILLEGAL_TEXT_RE, write_summary_sheet

AS_OF = datetime(2000, 1, 1, tzinfo=ZoneInfo("Asia/Bangkok"))


def _placeholder_row(**overrides) -> SummaryRow:
    base = dict(
        cost_center="CC1",
        gl_account="5000000000",
        department="Dept",
        division="Div",
        c_level="C",
        cc_name="CC Name",
        gl_name="GL Name",
        gl_group="Group",
        side="COST",
        status_label="Status",
        months=(0.0,) * 12,
        total_year=0.0,
        board_total_year=0.0,
        sap_total_year=0.0,
        remark="",
        detail_lines=(),
    )
    base.update(overrides)
    return SummaryRow(**base)


def _build_and_reload(rows: list[SummaryRow], warnings: list[str]) -> Worksheet:
    wb = Workbook()
    ws = wb.active
    write_summary_sheet(
        ws, rows, planning_year=2000, as_of=AS_OF, scope_label="Scope", n_departments=1,
        sap_watermark=None, warnings=warnings,
    )
    buf = BytesIO()
    wb.save(buf)  # the real round-trip — must never raise IllegalCharacterError
    return load_workbook(buf).active


REMARK_COL = 26  # LAST_NUM_COL(25) + 1
DETAIL_COL = 27  # LAST_NUM_COL(25) + 2


@pytest.mark.parametrize(
    "label, illegal_char",
    [
        ("U+FFFF noncharacter", "\uffff"),
        ("U+FFFE noncharacter", "\ufffe"),
        ("lone high surrogate", "\ud800"),
        ("C0 control char", "\x0b"),
    ],
)
def test_illegal_text_in_remark_is_stripped_saved_and_warned(label, illegal_char):
    row = _placeholder_row(remark=f"before{illegal_char}after")
    warnings: list[str] = []

    ws = _build_and_reload([row], warnings)

    cell_value = ws.cell(row=5, column=REMARK_COL).value
    assert illegal_char not in cell_value, label
    assert cell_value == "beforeafter", label
    assert any("illegal character stripped" in w for w in warnings), label
    # never leaks the value itself into the warning
    assert all("before" not in w and "after" not in w for w in warnings), label


def test_illegal_text_in_detail_cell_is_stripped_saved_and_warned():
    row = _placeholder_row(detail_lines=("line one\ufffeline two",))
    warnings: list[str] = []

    ws = _build_and_reload([row], warnings)

    cell_value = ws.cell(row=5, column=DETAIL_COL).value
    assert "\ufffe" not in cell_value
    assert cell_value == "line oneline two"
    assert any("illegal character stripped" in w for w in warnings)


def test_xlsx_illegal_text_re_matches_the_documented_set():
    for ch in ("\x00", "\x08", "\x0b", "\x0c", "\x0e", "\x1f", "\ufffe", "\uffff", "\ud800", "\udfff"):
        assert XLSX_ILLEGAL_TEXT_RE.search(ch), repr(ch)
    for ch in ("\t", "\n", "\r", "a", " ", "ก"):
        assert not XLSX_ILLEGAL_TEXT_RE.search(ch), repr(ch)
