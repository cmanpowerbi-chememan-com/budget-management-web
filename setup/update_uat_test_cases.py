"""Rewrite UAT-08's test-case text on the LIVE UAT workbook (SharePoint).

Background: UAT-08 originally tested a dash "-" + hover-tooltip mask that hid
incomplete-SAP months. ADR-0030 (docs/adr/0030-show-sap-actuals-as-is.md,
shipped 2026-09-14) removed that mask -- every month now shows its real
number, no dash, no tooltip anywhere on the SAP cells or the grand total; the
only freshness signal left is the legend chip ("SAP * ใช้จริง (2025) *
ข้อมูลบันทึกถึงวันที่ 11 Sep 26", built by `sapFreshnessLine` in
frontend/src/grid/model.ts, rendered in frontend/src/grid/BudgetGrid.tsx
~line 560). Per jakkaritw's instruction this script does NOT delete the row --
it rewrites the row's text so the steps that still apply stay, the ones that
no longer apply are gone, plus a scope note (this case checks display only,
not SAP reconciliation, which is out of scope for this UAT round).

Two modes, plus a batch mode for column L only:
  python -X utf8 setup/update_uat_test_cases.py                 # DRY RUN: download, edit UAT-08 (5 fields), verify, write a local preview file, print BEFORE/AFTER. Nothing leaves this machine.
  python -X utf8 setup/update_uat_test_cases.py --upload         # + push the edited bytes back to SharePoint (aborts if the file changed under us) + mirror the same text into .scratch/uat-prep/cases/W2.json
  add --keep-status to leave the Status cell untouched (today "N/A") instead of blanking it back to "Not Run" (UAT-08 mode only)

  python -X utf8 setup/update_uat_test_cases.py --steps-dir DIR           # DRY RUN: rewrite Test Steps (column L) ONLY, for every <CASE>.txt in DIR (CASE matches UAT-DD, two digits). --steps-dir alone (no value) uses .scratch/uat-text-audit/approved/colL.
  python -X utf8 setup/update_uat_test_cases.py --steps-dir DIR --upload  # + push the batch edit to SharePoint
  --steps-dir is mutually exclusive with the UAT-08 hard-coded 5-field edit above (and with --keep-status, which is meaningless in this mode -- Status is never touched).

WARNING: the live workbook (Budget_UAT_Test_Cases_07.09.2026.xlsx on
SharePoint) was hand-restructured by a tester to 19 columns (A:S) and carries
REAL tester results across 26 of its 54 rows. This script locates columns by
HEADER TEXT (never a hard-coded index). The UAT-08 mode edits ONLY that row's
5 text cells + Status; --steps-dir mode edits ONLY the Test Steps cell of each
named case. Either way every other cell's value AND style is left untouched,
and the sheet is never rebuilt from the repo copy (requirement_spec/5_uat/...xlsx),
which has a different, stale 20-column layout.
"""
import argparse
import base64
import io
import json
import os
import re
import zipfile
from pathlib import Path

import msal
import openpyxl
import requests
from dotenv import load_dotenv
from openpyxl.worksheet.worksheet import Worksheet

load_dotenv()

UAT_URL = (
    "https://chememan.sharepoint.com/:x:/r/teams/CMANDigitalTechnology/_layouts/15/"
    "Doc.aspx?sourcedoc=%7B69334028-028E-4B73-B74C-42C2962E2198%7D"
    "&file=Budget_UAT_Test_Cases_07.09.2026.xlsx&action=default&mobileredirect=true"
)
SHEET_CASES = "2. Test Cases"
SHEET_SUMMARY = "3. Summary"
HEADER_ROW = 2
EXPECTED_LAST_ROW = 56
EXPECTED_CASE_COUNT = 54
TARGET_CASE_ID = "UAT-08"

TENANT_ID = os.getenv("ENTRA_TENANT_ID")
CLIENT_ID = os.getenv("ENTRA_CLIENT_ID")
CLIENT_SECRET = os.getenv("ENTRA_CLIENT_SECRET")

PREVIEW_DIR = Path(__file__).parent.parent / "requirement_spec" / "5_uat" / "_build"
PREVIEW_NAME = "Budget_UAT_Test_Cases_07.09.2026.uat08-trim.preview.xlsx"

W2_JSON = Path(__file__).parent.parent / ".scratch" / "uat-prep" / "cases" / "W2.json"

# ---- --steps-dir batch mode: rewrite Test Steps (column L) only, for many
# cases at once, from one <CASE>.txt per case ----
DEFAULT_STEPS_DIR = Path(__file__).parent.parent / ".scratch" / "uat-text-audit" / "approved" / "colL"
PREVIEW_NAME_COLL = "Budget_UAT_Test_Cases_07.09.2026.colL.preview.xlsx"
CASE_ID_RE = re.compile(r"^(UAT-\d\d)$")

# ---- new UAT-08 text (ADR-0030: SAP shown as-is, no dash/tooltip mask; the
# only signal left is the "ข้อมูลบันทึกถึงวันที่ <date>" freshness chip on the legend) ----
NEW_SCENARIO = (
    "ตรวจว่าชั้น SAP · ใช้จริง แสดงตัวเลขจริงครบทั้ง 12 เดือน และป้าย legend "
    "บอกวันที่ข้อมูล SAP ล่าสุด (`ข้อมูลบันทึกถึงวันที่ <วันที่>`)"
)

NEW_PRECONDITIONS = (
    "- ล็อกอินเป็น Pornthip เปิดตารางฝ่าย `Data & Analytic` ปี `Year 2026` ไว้แล้ว\n"
    "- ระบบแสดง SAP ตามจริงทุกเดือน ไม่ซ่อนเดือนใด (ADR-0030)"
)

NEW_STEPS = (
    "1. ในตารางชั้น SAP · ใช้จริง ไล่ดูทั้ง 12 เดือน — ทุกเดือนต้องแสดงเป็นตัวเลข "
    "ไม่มีช่องไหนแสดงเป็นขีด –\n"
    "2. ดูป้าย legend เหนือตาราง ตรง `SAP · ใช้จริง (2025)` ต้องมีข้อความต่อท้ายว่า "
    "`ข้อมูลบันทึกถึงวันที่ <วันที่>` จดวันที่ที่เห็นลงช่อง Actual Result\n"
    "3. ดูแถว `รวมทั้งหมด · SAP · ใช้จริง` — ยอดรวมต้องเท่ากับผลบวกของทั้ง 12 เดือนในแถวนั้น"
)

NEW_TEST_DATA = (
    "ชั้นตรวจ: `SAP · ใช้จริง` ฝ่าย Data & Analytic\n"
    "มองหา: ป้าย legend `ข้อมูลบันทึกถึงวันที่ <วันที่>` เหนือตาราง"
)

NEW_EXPECTED = (
    "1) ทุกเดือนของชั้น SAP แสดงเป็นตัวเลข ไม่มีช่องที่แสดงขีด –\n"
    "2) ป้าย legend แสดง `ข้อมูลบันทึกถึงวันที่ <วันที่>` — ถ้ามี `⚠` นำหน้า แปลว่าข้อมูลช้ากว่าปกติ "
    "ให้จดไว้ใน Remarks ไม่นับเป็น Fail\n"
    "3) ยอด `รวมทั้งหมด · SAP · ใช้จริง` เท่ากับผลบวก 12 เดือน\n"
    "\n"
    "— หมายเหตุสำคัญ —\n"
    "- เคสนี้ตรวจการแสดงผล ไม่ใช่กระทบยอด SAP ซึ่งนอกขอบเขต UAT รอบนี้"
)

# header text -> new value (header text is matched startswith, see _col)
NEW_TEXT: dict[str, str] = {
    "Test Scenario": NEW_SCENARIO,
    "Preconditions": NEW_PRECONDITIONS,
    "Test Steps": NEW_STEPS,
    "Test Data": NEW_TEST_DATA,
    "Expected Result": NEW_EXPECTED,
}

# header text -> W2.json field name, for the --upload mirror step
W2_FIELD_MAP: dict[str, str] = {
    "Test Scenario": "scenario",
    "Preconditions": "preconditions",
    "Test Steps": "steps",
    "Test Data": "test_data",
    "Expected Result": "expected",
}


# ── SharePoint / Graph helpers (pattern copied from setup/update_sit_test_cases.py,
# parameterised on `share_url` here since this script targets a different file) ──
def _get_token() -> str:
    app = msal.ConfidentialClientApplication(
        CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{TENANT_ID}",
        client_credential=CLIENT_SECRET,
    )
    result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    if "access_token" not in result:
        raise RuntimeError(f"Token error: {result.get('error_description', result)}")
    return result["access_token"]


def _share_id(url: str) -> str:
    b64 = base64.b64encode(url.encode("utf-8")).decode("ascii")
    return "u!" + b64.rstrip("=").replace("/", "_").replace("+", "-")


def resolve_item(token: str, share_url: str) -> dict:
    r = requests.get(
        f"https://graph.microsoft.com/v1.0/shares/{_share_id(share_url)}/driveItem",
        headers={"Authorization": f"Bearer {token}"},
    )
    r.raise_for_status()
    item = r.json()
    return {
        "id": item["id"],
        "name": item["name"],
        "driveId": item["parentReference"]["driveId"],
        "lastModifiedDateTime": item.get("lastModifiedDateTime"),
        "lastModifiedBy": (item.get("lastModifiedBy") or {}).get("user", {}).get("displayName"),
        "webUrl": item.get("webUrl"),
    }


def download(token: str, item: dict) -> bytes:
    r = requests.get(
        f"https://graph.microsoft.com/v1.0/drives/{item['driveId']}/items/{item['id']}/content",
        headers={"Authorization": f"Bearer {token}"},
    )
    r.raise_for_status()
    return r.content


def upload(token: str, item: dict, data: bytes) -> str:
    r = requests.put(
        f"https://graph.microsoft.com/v1.0/drives/{item['driveId']}/items/{item['id']}/content",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/octet-stream"},
        data=data,
    )
    if r.status_code not in (200, 201):
        raise RuntimeError(f"Upload failed {r.status_code}: {r.text[:400]}")
    return r.json().get("webUrl", "")


# ── column lookup by header text (never a hard-coded index) ────────────────
def _header_map(ws: Worksheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(row=HEADER_ROW, column=col).value
        if raw:
            headers[str(raw).strip()] = col
    return headers


def _col(headers: dict[str, int], name: str) -> int:
    """Exact match first, else startswith (handles bilingual headers like
    'Preconditions / เงื่อนไขก่อนทดสอบ')."""
    if name in headers:
        return headers[name]
    for key, idx in headers.items():
        if key.startswith(name):
            return idx
    raise RuntimeError(f"Header {name!r} not found in {SHEET_CASES!r} — column layout changed, aborting")


def _find_target_row(ws: Worksheet, id_col: int) -> int:
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        if ws.cell(row=r, column=id_col).value == TARGET_CASE_ID:
            return r
    raise RuntimeError(f"{TARGET_CASE_ID} not found in {SHEET_CASES!r} — aborting")


def _row_by_id(ws: Worksheet, id_col: int) -> dict[str, int]:
    """Map every non-blank Test Case ID in `ws` to its row number."""
    rows: dict[str, int] = {}
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        case_id = ws.cell(row=r, column=id_col).value
        if case_id:
            rows[str(case_id)] = r
    return rows


def load_steps_dir(steps_dir: Path) -> dict[str, str]:
    """Read every `<CASE>.txt` in `steps_dir` whose stem matches UAT-\\d\\d.
    Each file's entire content becomes that case's new Test Steps value:
    CRLF normalised to LF, trailing whitespace stripped. Returns
    {case_id: new_text}, ordered by filename."""
    cases: dict[str, str] = {}
    for path in sorted(steps_dir.glob("*.txt")):
        match = CASE_ID_RE.match(path.stem)
        if not match:
            continue
        text = path.read_text(encoding="utf-8").replace("\r\n", "\n").rstrip()
        cases[match.group(1)] = text
    return cases


# ---------------------------------------------------------------------------
# SharePoint-only package parts must survive the save. Copied verbatim from
# setup/fill_sit_test_case_gaps.py (_preserved_parts / _with_fresh_ids /
# _inject / _preserve_sharepoint_parts / _save_workbook, ~lines 668-755 there)
# rather than importing that module, since it also carries a large SIT-specific
# EDITS/RESULTS/IMAGE_RESULTS dataset this script has no business depending on.
# openpyxl rebuilds the .xlsx from its own object model and silently drops any
# package part it doesn't model -- on this SharePoint-managed workbook that is
# customXml/* (library metadata plumbing) and docMetadata/LabelInfo.xml (the
# Microsoft sensitivity label). These helpers splice those parts back verbatim
# and re-declare them in [Content_Types].xml + the package/workbook rels.
# ---------------------------------------------------------------------------
_PRESERVED_PART_PREFIXES = ("customXml/", "docMetadata/")
_CUSTOM_XML_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
)
_RELATIONSHIP_RE = re.compile(r"<Relationship\b[^>]*/>")
_OVERRIDE_RE = re.compile(r"<Override\b[^>]*/>")
_ID_ATTR_RE = re.compile(r'\sId="[^"]*"')


def _preserved_parts(raw: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        return {
            name: z.read(name)
            for name in z.namelist()
            if name.startswith(_PRESERVED_PART_PREFIXES)
        }


def _with_fresh_ids(elements: list[str], prefix: str) -> list[str]:
    """Re-stamp each Relationship's Id so it cannot collide with the ids
    openpyxl generated for its own parts (it always numbers rId1..rIdN)."""
    return [
        _ID_ATTR_RE.sub("", el).replace("<Relationship", f'<Relationship Id="{prefix}{i}"', 1)
        for i, el in enumerate(elements, start=1)
    ]


def _inject(xml: str, closing_tag: str, elements: list[str]) -> str:
    return xml.replace(closing_tag, "".join(elements) + closing_tag, 1) if elements else xml


def _preserve_sharepoint_parts(original: bytes, saved: bytes) -> tuple[bytes, list[str]]:
    """Return `saved` with `original`'s SharePoint-only parts put back, plus
    the names restored (empty list when the source had none)."""
    keep = _preserved_parts(original)
    if not keep:
        return saved, []

    with zipfile.ZipFile(io.BytesIO(original)) as z:
        orig_ct = z.read("[Content_Types].xml").decode("utf-8")
        orig_pkg_rels = z.read("_rels/.rels").decode("utf-8")
        orig_wb_rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")

    overrides = [
        el
        for el in _OVERRIDE_RE.findall(orig_ct)
        if any(f'PartName="/{prefix}' in el for prefix in _PRESERVED_PART_PREFIXES)
    ]
    pkg_rels = _with_fresh_ids(
        [el for el in _RELATIONSHIP_RE.findall(orig_pkg_rels) if "docMetadata/" in el],
        "rIdSPpkg",
    )
    wb_rels = _with_fresh_ids(
        [el for el in _RELATIONSHIP_RE.findall(orig_wb_rels) if _CUSTOM_XML_REL_TYPE in el],
        "rIdSPcx",
    )

    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(saved)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = _inject(data.decode("utf-8"), "</Types>", overrides).encode("utf-8")
            elif item.filename == "_rels/.rels":
                data = _inject(data.decode("utf-8"), "</Relationships>", pkg_rels).encode("utf-8")
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = _inject(data.decode("utf-8"), "</Relationships>", wb_rels).encode("utf-8")
            dst.writestr(item, data)
        for name, data in keep.items():
            dst.writestr(name, data)
    return out.getvalue(), sorted(keep)


def _save_workbook(wb: openpyxl.Workbook, original_raw: bytes) -> bytes:
    """The ONLY place this script serialises a workbook — openpyxl's save plus
    the package-part restore above."""
    buf = io.BytesIO()
    wb.save(buf)
    merged, restored = _preserve_sharepoint_parts(original_raw, buf.getvalue())
    if restored:
        print(
            f"[Save] restored {len(restored)} SharePoint-only package part(s) that openpyxl "
            f"drops, sensitivity label included: {', '.join(restored)}"
        )
    return merged


# ── workbook edit ────────────────────────────────────────────────────────────
def edit_workbook(raw: bytes, keep_status: bool) -> tuple[bytes, dict[str, object], dict[str, object]]:
    """Rewrite UAT-08's 5 text cells (+ Status, unless `keep_status`) in
    place. This function only ever assigns `.value` on cells that already
    exist -- it never touches `.font` / `.alignment` / `.fill` / `.border`,
    and never adds or removes a row. Returns (edited bytes, before, after)."""
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb[SHEET_CASES]
    headers = _header_map(ws)
    id_col = _col(headers, "Test Case ID")
    row = _find_target_row(ws, id_col)

    before: dict[str, object] = {}
    after: dict[str, object] = {}

    for name, new_value in NEW_TEXT.items():
        col = _col(headers, name)
        cell = ws.cell(row=row, column=col)
        before[name] = cell.value
        cell.value = new_value
        after[name] = new_value

    status_col = _col(headers, "Status")
    status_cell = ws.cell(row=row, column=status_col)
    before["Status"] = status_cell.value
    if keep_status:
        after["Status"] = status_cell.value
    else:
        status_cell.value = None
        after["Status"] = None

    return _save_workbook(wb, raw), before, after


def edit_workbook_batch(
    raw: bytes, case_steps: dict[str, str]
) -> tuple[bytes, dict[str, object], dict[str, object], set[tuple[int, int]]]:
    """Rewrite the Test Steps cell for every case in `case_steps`, one cell
    per case. Only `.value` is ever assigned -- never `.font` / `.alignment`
    / `.fill` / `.border`, never a row insert/delete, never any other column
    (Actual Result, Status, Tester, etc. are untouched). A case id absent
    from the sheet is a hard error naming the id. Returns (edited bytes,
    before-by-case-id, after-by-case-id, {(row, col)} of every cell touched)."""
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb[SHEET_CASES]
    headers = _header_map(ws)
    id_col = _col(headers, "Test Case ID")
    steps_col = _col(headers, "Test Steps")
    rows = _row_by_id(ws, id_col)

    before: dict[str, object] = {}
    after: dict[str, object] = {}
    edited_cells: set[tuple[int, int]] = set()

    for case_id, new_text in case_steps.items():
        if case_id not in rows:
            raise RuntimeError(f"{case_id} not found in {SHEET_CASES!r} — aborting")
        row = rows[case_id]
        cell = ws.cell(row=row, column=steps_col)
        before[case_id] = cell.value
        cell.value = new_text
        after[case_id] = new_text
        edited_cells.add((row, steps_col))

    return _save_workbook(wb, raw), before, after, edited_cells


# ── verify ───────────────────────────────────────────────────────────────────
def verify(original_raw: bytes, new_raw: bytes, keep_status: bool) -> list[str]:
    """Hard-fail (AssertionError) on any deviation from the intended edit.
    Returns the list of passing-check lines to print. Runs in BOTH dry-run
    and --upload mode."""
    lines: list[str] = []

    wb_o = openpyxl.load_workbook(io.BytesIO(original_raw))
    wb_n = openpyxl.load_workbook(io.BytesIO(new_raw))

    assert wb_o.sheetnames == wb_n.sheetnames, (
        f"sheet list changed: {wb_o.sheetnames} -> {wb_n.sheetnames}"
    )
    lines.append(f"[Verify] sheet list unchanged: {wb_o.sheetnames}")

    ws_o = wb_o[SHEET_CASES]
    ws_n = wb_n[SHEET_CASES]
    headers = _header_map(ws_n)
    id_col = _col(headers, "Test Case ID")
    target_row = _find_target_row(ws_n, id_col)
    status_col = _col(headers, "Status")
    edited_cols = {_col(headers, name) for name in NEW_TEXT} | {status_col}

    # (a) UAT-08's 5 text cells + Status == the intended new values
    for name, expected_value in NEW_TEXT.items():
        got = ws_n.cell(row=target_row, column=_col(headers, name)).value
        assert got == expected_value, f"UAT-08 {name!r} mismatch after edit: {got!r}"
    status_got = ws_n.cell(row=target_row, column=status_col).value
    if keep_status:
        status_before = ws_o.cell(row=target_row, column=status_col).value
        assert status_got == status_before, (
            f"UAT-08 Status changed despite --keep-status: {status_before!r} -> {status_got!r}"
        )
    else:
        assert status_got is None, f"UAT-08 Status not blanked: {status_got!r}"
    lines.append(
        f"[Verify] (a) UAT-08 row {target_row}: 5 text cells + Status match the intended edit "
        f"(keep_status={keep_status})."
    )

    # (b) every OTHER cell of every sheet is byte-for-byte identical
    #     (also tallies (c) Summary formula-cell counts in the same pass)
    formula_count_o = 0
    formula_count_n = 0
    mismatch = None
    for name in wb_o.sheetnames:
        so, sn = wb_o[name], wb_n[name]
        for row_o, row_n in zip(so.iter_rows(), sn.iter_rows()):
            for cell_o, cell_n in zip(row_o, row_n):
                if name == SHEET_CASES and cell_o.row == target_row and cell_o.column in edited_cols:
                    continue
                if cell_o.value != cell_n.value:
                    mismatch = (name, cell_o.coordinate, cell_o.value, cell_n.value)
                    break
                if name == SHEET_SUMMARY:
                    if isinstance(cell_o.value, str) and cell_o.value.startswith("="):
                        formula_count_o += 1
                    if isinstance(cell_n.value, str) and cell_n.value.startswith("="):
                        formula_count_n += 1
            if mismatch:
                break
        if mismatch:
            break
    assert mismatch is None, (
        f"cell value changed outside the intended edit: sheet {mismatch[0]!r} "
        f"{mismatch[1]} {mismatch[2]!r} -> {mismatch[3]!r}"
    )
    lines.append(
        f"[Verify] (b) every cell outside UAT-08's {len(edited_cols)} edited cells is "
        f"byte-for-byte unchanged (checked {len(wb_o.sheetnames)} sheets, lockstep iter_rows)."
    )

    # (c) Summary formulas unchanged (values already proven identical by (b))
    assert formula_count_o == formula_count_n, (
        f"formula-cell count in {SHEET_SUMMARY!r} changed: {formula_count_o} -> {formula_count_n}"
    )
    lines.append(
        f"[Verify] (c) {SHEET_SUMMARY!r}: {formula_count_o} formula cells, count and values unchanged."
    )

    # (d) autofilter ref / freeze pane / data-validation sqrefs unchanged
    assert str(ws_o.auto_filter.ref) == str(ws_n.auto_filter.ref), (
        f"autofilter ref changed: {ws_o.auto_filter.ref} -> {ws_n.auto_filter.ref}"
    )
    assert ws_o.freeze_panes == ws_n.freeze_panes, (
        f"freeze pane changed: {ws_o.freeze_panes} -> {ws_n.freeze_panes}"
    )
    sqrefs_o = sorted(str(dv.sqref) for dv in ws_o.data_validations.dataValidation)
    sqrefs_n = sorted(str(dv.sqref) for dv in ws_n.data_validations.dataValidation)
    assert sqrefs_o == sqrefs_n, f"data validation sqrefs changed: {sqrefs_o} -> {sqrefs_n}"
    lines.append(
        f"[Verify] (d) autofilter={ws_n.auto_filter.ref!r}, freeze_panes={ws_n.freeze_panes!r}, "
        f"{len(sqrefs_n)} data-validation sqref(s) unchanged."
    )

    # (e) SharePoint-only package parts preserved (sensitivity label included)
    kept_o = set(_preserved_parts(original_raw))
    kept_n = set(_preserved_parts(new_raw))
    missing = kept_o - kept_n
    assert not missing, f"SharePoint-only part(s) lost by this edit: {sorted(missing)}"
    assert "docMetadata/LabelInfo.xml" in kept_n, (
        "docMetadata/LabelInfo.xml (sensitivity label) missing from the saved file"
    )
    lines.append(
        f"[Verify] (e) {len(kept_o)} SharePoint-only package part(s) preserved, sensitivity "
        f"label included: {sorted(kept_o)}"
    )

    # (f) row count still 56 / 54 case ids UAT-01..UAT-54 contiguous
    assert ws_n.max_row == EXPECTED_LAST_ROW, (
        f"row count changed: expected {EXPECTED_LAST_ROW}, got {ws_n.max_row}"
    )
    ids = [
        v
        for r in range(HEADER_ROW + 1, ws_n.max_row + 1)
        if (v := ws_n.cell(row=r, column=id_col).value)
    ]
    expected_ids = [f"UAT-{i:02d}" for i in range(1, EXPECTED_CASE_COUNT + 1)]
    assert ids == expected_ids, f"case id sequence broken: {ids}"
    lines.append(
        f"[Verify] (f) row count {ws_n.max_row}, {len(ids)} case ids "
        f"UAT-01..UAT-{EXPECTED_CASE_COUNT:02d} contiguous."
    )

    return lines


def verify_batch(original_raw: bytes, new_raw: bytes, case_steps: dict[str, str]) -> list[str]:
    """Hard-fail (AssertionError) on any deviation from the intended
    Test-Steps-only batch edit. Same six checks as verify(), generalised over
    N edited cells (one Test Steps cell per case) instead of one row's 5
    fields + Status. Runs in BOTH dry-run and --upload mode."""
    lines: list[str] = []

    wb_o = openpyxl.load_workbook(io.BytesIO(original_raw))
    wb_n = openpyxl.load_workbook(io.BytesIO(new_raw))

    assert wb_o.sheetnames == wb_n.sheetnames, (
        f"sheet list changed: {wb_o.sheetnames} -> {wb_n.sheetnames}"
    )
    lines.append(f"[Verify] sheet list unchanged: {wb_o.sheetnames}")

    ws_o = wb_o[SHEET_CASES]
    ws_n = wb_n[SHEET_CASES]
    headers = _header_map(ws_n)
    id_col = _col(headers, "Test Case ID")
    steps_col = _col(headers, "Test Steps")
    rows = _row_by_id(ws_n, id_col)

    # (a) every targeted case's Test Steps cell == the intended new text
    edited_cells: set[tuple[int, int]] = set()
    for case_id, expected_text in case_steps.items():
        row = rows[case_id]
        got = ws_n.cell(row=row, column=steps_col).value
        assert got == expected_text, f"{case_id} Test Steps mismatch after edit: {got!r}"
        edited_cells.add((row, steps_col))
    lines.append(
        f"[Verify] (a) {len(case_steps)} case(s) Test Steps cell match the intended edit: "
        f"{', '.join(sorted(case_steps))}."
    )

    # (b) every OTHER cell of every sheet is byte-for-byte identical
    #     (also tallies (c) Summary formula-cell counts in the same pass)
    formula_count_o = 0
    formula_count_n = 0
    mismatch = None
    for name in wb_o.sheetnames:
        so, sn = wb_o[name], wb_n[name]
        for row_o, row_n in zip(so.iter_rows(), sn.iter_rows()):
            for cell_o, cell_n in zip(row_o, row_n):
                if name == SHEET_CASES and (cell_o.row, cell_o.column) in edited_cells:
                    continue
                if cell_o.value != cell_n.value:
                    mismatch = (name, cell_o.coordinate, cell_o.value, cell_n.value)
                    break
                if name == SHEET_SUMMARY:
                    if isinstance(cell_o.value, str) and cell_o.value.startswith("="):
                        formula_count_o += 1
                    if isinstance(cell_n.value, str) and cell_n.value.startswith("="):
                        formula_count_n += 1
            if mismatch:
                break
        if mismatch:
            break
    assert mismatch is None, (
        f"cell value changed outside the intended edit: sheet {mismatch[0]!r} "
        f"{mismatch[1]} {mismatch[2]!r} -> {mismatch[3]!r}"
    )
    lines.append(
        f"[Verify] (b) every cell outside the {len(edited_cells)} edited Test Steps cell(s) is "
        f"byte-for-byte unchanged (checked {len(wb_o.sheetnames)} sheets, lockstep iter_rows)."
    )

    # (c) Summary formulas unchanged (values already proven identical by (b))
    assert formula_count_o == formula_count_n, (
        f"formula-cell count in {SHEET_SUMMARY!r} changed: {formula_count_o} -> {formula_count_n}"
    )
    lines.append(
        f"[Verify] (c) {SHEET_SUMMARY!r}: {formula_count_o} formula cells, count and values unchanged."
    )

    # (d) autofilter ref / freeze pane / data-validation sqrefs unchanged
    assert str(ws_o.auto_filter.ref) == str(ws_n.auto_filter.ref), (
        f"autofilter ref changed: {ws_o.auto_filter.ref} -> {ws_n.auto_filter.ref}"
    )
    assert ws_o.freeze_panes == ws_n.freeze_panes, (
        f"freeze pane changed: {ws_o.freeze_panes} -> {ws_n.freeze_panes}"
    )
    sqrefs_o = sorted(str(dv.sqref) for dv in ws_o.data_validations.dataValidation)
    sqrefs_n = sorted(str(dv.sqref) for dv in ws_n.data_validations.dataValidation)
    assert sqrefs_o == sqrefs_n, f"data validation sqrefs changed: {sqrefs_o} -> {sqrefs_n}"
    lines.append(
        f"[Verify] (d) autofilter={ws_n.auto_filter.ref!r}, freeze_panes={ws_n.freeze_panes!r}, "
        f"{len(sqrefs_n)} data-validation sqref(s) unchanged."
    )

    # (e) SharePoint-only package parts preserved (sensitivity label included)
    kept_o = set(_preserved_parts(original_raw))
    kept_n = set(_preserved_parts(new_raw))
    missing = kept_o - kept_n
    assert not missing, f"SharePoint-only part(s) lost by this edit: {sorted(missing)}"
    assert "docMetadata/LabelInfo.xml" in kept_n, (
        "docMetadata/LabelInfo.xml (sensitivity label) missing from the saved file"
    )
    lines.append(
        f"[Verify] (e) {len(kept_o)} SharePoint-only package part(s) preserved, sensitivity "
        f"label included: {sorted(kept_o)}"
    )

    # (f) row count still 56 / 54 case ids UAT-01..UAT-54 contiguous
    assert ws_n.max_row == EXPECTED_LAST_ROW, (
        f"row count changed: expected {EXPECTED_LAST_ROW}, got {ws_n.max_row}"
    )
    ids = [
        v
        for r in range(HEADER_ROW + 1, ws_n.max_row + 1)
        if (v := ws_n.cell(row=r, column=id_col).value)
    ]
    expected_ids = [f"UAT-{i:02d}" for i in range(1, EXPECTED_CASE_COUNT + 1)]
    assert ids == expected_ids, f"case id sequence broken: {ids}"
    lines.append(
        f"[Verify] (f) row count {ws_n.max_row}, {len(ids)} case ids "
        f"UAT-01..UAT-{EXPECTED_CASE_COUNT:02d} contiguous."
    )

    return lines


def _mirror_to_w2_json(after: dict[str, object]) -> None:
    """--upload only: mirror the same 5 texts into the UAT-08 object of
    .scratch/uat-prep/cases/W2.json, preserving key order and every other
    field (including `note`) untouched."""
    with open(W2_JSON, "r", encoding="utf-8") as f:
        data = json.load(f)
    target = next((obj for obj in data if obj.get("id") == TARGET_CASE_ID), None)
    if target is None:
        raise RuntimeError(f"{TARGET_CASE_ID} not found in {W2_JSON}")
    for header_name, json_key in W2_FIELD_MAP.items():
        target[json_key] = after[header_name]
    with open(W2_JSON, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"Mirrored : {W2_JSON} (UAT-08 fields {list(W2_FIELD_MAP.values())})")


def _run_steps_dir_batch(token: str, item: dict, raw: bytes, steps_dir: Path, do_upload: bool) -> None:
    """--steps-dir dispatch target: batch-rewrite Test Steps (column L) only
    for every case file in `steps_dir`. Mirrors main()'s print/verify/preview
    shape for the UAT-08 path, but never touches Status and never mirrors
    W2.json (not requested for this mode)."""
    case_steps = load_steps_dir(steps_dir)
    if not case_steps:
        raise RuntimeError(f"no <CASE>.txt files found in {steps_dir}")

    new_raw, before, after, _edited_cells = edit_workbook_batch(raw, case_steps)
    verify_lines = verify_batch(raw, new_raw, case_steps)

    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    preview_path = PREVIEW_DIR / PREVIEW_NAME_COLL
    preview_path.write_bytes(new_raw)

    print(f"\n=== BEFORE / AFTER — Test Steps ({len(case_steps)} case(s)) ===")
    for case_id in sorted(case_steps):
        print(f"\n[{case_id}] BEFORE:")
        print(before[case_id])
        print(f"[{case_id}] AFTER:")
        print(after[case_id])

    print("\n=== VERIFY ===")
    for line in verify_lines:
        print(line)

    print(f"\nPreview  : {preview_path} ({len(new_raw):,} bytes)")
    print(f"Applied  : {len(case_steps)} case(s): {', '.join(sorted(case_steps))}")

    if not do_upload:
        print("\nDRY RUN — nothing uploaded")
        return

    fresh = resolve_item(token, UAT_URL)
    if fresh["lastModifiedDateTime"] != item["lastModifiedDateTime"]:
        raise RuntimeError(
            "ABORT: the SharePoint file changed while this script was running "
            f"({item['lastModifiedDateTime']} -> {fresh['lastModifiedDateTime']}). Re-run."
        )
    new_url = upload(token, item, new_raw)
    print(f"\nUploaded : {new_url}")
    fresh2 = resolve_item(token, UAT_URL)
    print(f"New lastModifiedDateTime: {fresh2['lastModifiedDateTime']}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--upload", action="store_true", help="push the edited file back to SharePoint")
    ap.add_argument(
        "--keep-status",
        action="store_true",
        help=(
            "leave Status untouched (today 'N/A') instead of blanking it back to Not Run "
            "(UAT-08 mode only; --steps-dir mode never touches Status)"
        ),
    )
    ap.add_argument(
        "--steps-dir",
        nargs="?",
        const=str(DEFAULT_STEPS_DIR),
        default=None,
        help=(
            "batch-rewrite Test Steps (column L) ONLY, for every <CASE>.txt in DIR "
            f"(CASE matches UAT-\\d\\d). Bare flag (no value) uses {DEFAULT_STEPS_DIR}. "
            "Mutually exclusive with the UAT-08 hard-coded 5-field edit."
        ),
    )
    args = ap.parse_args()

    token = _get_token()
    item = resolve_item(token, UAT_URL)
    print(f"File     : {item['name']}")
    print(f"Item id  : {item['id']}")
    print(f"Drive id : {item['driveId']}")
    print(f"Modified : {item['lastModifiedDateTime']} by {item['lastModifiedBy']}")

    raw = download(token, item)
    print(f"Downloaded: {len(raw):,} bytes")

    if args.steps_dir is not None:
        _run_steps_dir_batch(token, item, raw, Path(args.steps_dir), args.upload)
        return

    new_raw, before, after = edit_workbook(raw, args.keep_status)
    verify_lines = verify(raw, new_raw, args.keep_status)

    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)
    preview_path = PREVIEW_DIR / PREVIEW_NAME
    preview_path.write_bytes(new_raw)

    print("\n=== BEFORE / AFTER — UAT-08 ===")
    for name in (*NEW_TEXT, "Status"):
        print(f"\n[{name}] BEFORE:")
        print(before[name])
        print(f"[{name}] AFTER:")
        print(after[name])

    print("\n=== VERIFY ===")
    for line in verify_lines:
        print(line)

    print(f"\nPreview  : {preview_path} ({len(new_raw):,} bytes)")

    if not args.upload:
        print("\nDRY RUN — nothing uploaded")
        return

    fresh = resolve_item(token, UAT_URL)
    if fresh["lastModifiedDateTime"] != item["lastModifiedDateTime"]:
        raise RuntimeError(
            "ABORT: the SharePoint file changed while this script was running "
            f"({item['lastModifiedDateTime']} -> {fresh['lastModifiedDateTime']}). Re-run."
        )
    new_url = upload(token, item, new_raw)
    print(f"\nUploaded : {new_url}")
    fresh2 = resolve_item(token, UAT_URL)
    print(f"New lastModifiedDateTime: {fresh2['lastModifiedDateTime']}")

    _mirror_to_w2_json(after)


if __name__ == "__main__":
    main()
